import io
import json
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.http import require_POST, require_GET
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.db.models import Q

from .models import (
    TipoServicoMatrizaria,
    MatrizFisica,
    SolicitacaoServicoMatrizaria,
    CicloExecucaoMatrizaria,
    HistoricoTransicaoServicoMatrizaria,
)
from .services import MatrizariaService
from .forms import (
    SolicitacaoServicoForm,
    FinalizarExecucaoForm,
    ConferirServicoForm,
    TransferirResponsabilidadeForm,
    CancelarServicoForm,
    RelatorioFiltroForm,
    EditarSolicitacaoForm,
)
from .decorators import (
    matrizaria_access_required,
    matrizaria_solicitar_required,
    matrizaria_executar_required,
    matrizaria_conferir_required,
    matrizaria_cancelar_required,
    matrizaria_relatorios_required,
    matrizaria_tv_required,
    _user_can_solicitar,
    _user_can_executar,
    _user_can_conferir,
    _user_can_cancelar,
    _user_can_relatorios,
    _user_can_tv,
)


def sanitize_excel_cell(val):
    """
    Previne fórmula injection em planilhas Excel quando strings iniciam
    com caracteres '=', '+', '-', '@', '\t', '\r'.
    """
    if val is None:
        return ""
    s = str(val)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s}"
    return s


@login_required
@matrizaria_access_required
def kanban_view(request):
    """
    Painel Kanban operacional da Matrizaria organizado em 4 raias:
    1. Solicitados / Retrabalho
    2. Em Execução
    3. Aguardando Conferência
    4. Concluídos Recentes (últimas 48h)
    """
    user = request.user

    # Raias
    fila_solicitados = (
        SolicitacaoServicoMatrizaria.objects.filter(status__in=["SOLICITADO", "AGUARDANDO_RETRABALHO"])
        .select_related("prensa", "tipo_servico", "matriz_fisica", "solicitado_por")
        .order_by("-prioridade", "data_solicitacao")
    )

    fila_execucao = (
        SolicitacaoServicoMatrizaria.objects.filter(status="EM_EXECUCAO")
        .select_related("prensa", "tipo_servico", "matriz_fisica", "responsavel_atribuido")
        .order_by("-prioridade", "data_inicio_execucao")
    )

    fila_conferencia = (
        SolicitacaoServicoMatrizaria.objects.filter(status="AGUARDANDO_CONFERENCIA")
        .select_related("prensa", "tipo_servico", "matriz_fisica", "finalizado_por")
        .order_by("-prioridade", "data_fim_execucao")
    )

    hoje_inicio = timezone.now().replace(hour=0, minute=0, second=0, microsecond=0)
    limite_48h = timezone.now() - timedelta(hours=48)

    fila_concluidos = (
        SolicitacaoServicoMatrizaria.objects.filter(status="CONCLUIDO", data_conferencia__gte=limite_48h)
        .select_related("prensa", "tipo_servico", "matriz_fisica", "conferido_por")
        .order_by("-data_conferencia")[:15]
    )

    concluidos_hoje_count = SolicitacaoServicoMatrizaria.objects.filter(
        status="CONCLUIDO", data_conferencia__gte=hoje_inicio
    ).count()

    context = {
        "fila_solicitados": fila_solicitados,
        "fila_execucao": fila_execucao,
        "fila_conferencia": fila_conferencia,
        "fila_concluidos": fila_concluidos,
        "total_pendentes": fila_solicitados.count(),
        "total_em_execucao": fila_execucao.count(),
        "total_aguardando_conferencia": fila_conferencia.count(),
        "total_concluidos_hoje": concluidos_hoje_count,
        "can_solicitar": _user_can_solicitar(user),
        "can_executar": _user_can_executar(user),
        "can_conferir": _user_can_conferir(user),
        "can_cancelar": _user_can_cancelar(user),
        "can_relatorios": _user_can_relatorios(user),
        "can_tv": _user_can_tv(user),
    }
    return render(request, "matrizaria/kanban.html", context)


@login_required
@matrizaria_solicitar_required
def solicitar_servico_view(request):
    """
    Formulário de abertura de novo chamado de serviço da Matrizaria com proteção de idempotência.
    """
    if request.method == "POST":
        form = SolicitacaoServicoForm(request.POST)
        idempotency_key = request.POST.get("idempotency_key", "").strip()

        # Proteção contra reenvio duplo por duplo clique ou timeout
        session_key = f"matrizaria_req_done_{idempotency_key}" if idempotency_key else None
        if session_key and request.session.get(session_key):
            messages.info(request, "Esta solicitação já foi processada anteriormente.")
            return redirect("matrizaria:detalhe_servico", pk=request.session[session_key])

        if form.is_valid():
            try:
                solicitacao = MatrizariaService.criar_solicitacao(
                    prensa=form.cleaned_data["prensa"],
                    tipo_servico=form.cleaned_data["tipo_servico"],
                    descricao_solicitacao=form.cleaned_data["descricao_solicitacao"],
                    solicitado_por=request.user,
                    prioridade=form.cleaned_data["prioridade"],
                    matriz_fisica=form.cleaned_data["matriz_fisica"],
                    destino=form.cleaned_data.get("destino", "MAQUINA"),
                )
                if session_key:
                    request.session[session_key] = solicitacao.id

                messages.success(
                    request,
                    f"Solicitação SM #{solicitacao.id} para {solicitacao.equipamento_display} aberta com sucesso!",
                )
                return redirect("matrizaria:detalhe_servico", pk=solicitacao.id)
            except ValidationError as e:
                form.add_error(None, str(e.message if hasattr(e, "message") else e))
    else:
        form = SolicitacaoServicoForm()

    context = {
        "form": form,
        "titulo": "Abrir Chamado de Matrizaria",
    }
    return render(request, "matrizaria/form_solicitacao.html", context)


@login_required
@matrizaria_access_required
def detalhe_servico_view(request, pk):
    """
    Visualização detalhada da solicitação, ciclos de execução e histórico completo de transições.
    Apresenta dinamicamente os formulários de ação adequados ao perfil e estado do chamado.
    """
    solicitacao = get_object_or_404(
        SolicitacaoServicoMatrizaria.objects.select_related(
            "prensa",
            "tipo_servico",
            "matriz_fisica",
            "solicitado_por",
            "responsavel_atribuido",
            "iniciado_por",
            "finalizado_por",
            "conferido_por",
            "cancelado_por",
        ),
        pk=pk,
    )

    ciclos = solicitacao.ciclos_execucao.all().order_by("numero_ciclo")
    historico = solicitacao.historico_transicoes.all().order_by("-data_evento", "-id")

    user = request.user
    user_can_exec = _user_can_executar(user)
    user_can_conf = _user_can_conferir(user)
    user_can_canc = _user_can_cancelar(user)

    # Regra de autoconferência: usuário que participou da execução técnica não pode aprovar
    impedimento_autoconferencia = False
    if solicitacao.status == "AGUARDANDO_CONFERENCIA" and user_can_conf:
        participou = solicitacao.ciclos_execucao.filter(
            Q(usuario_inicio=user) | Q(usuario_fim=user)
        ).exists()
        if (
            participou
            or solicitacao.iniciado_por == user
            or solicitacao.finalizado_por == user
            or solicitacao.responsavel_atribuido == user
        ):
            impedimento_autoconferencia = True

    # Regra de permissão para edição e exclusão de solicitação pendente
    user_can_sol = _user_can_solicitar(user)
    is_solicitante_proprio = (solicitacao.solicitado_por_id == user.id and user_can_sol)
    is_gestao = (
        user.is_superuser
        or user.is_staff
        or user_can_conf
        or user.groups.filter(name__in=["Liderança de Produção", "Lideres", "Tecnicos_Lideres"]).exists()
        or user.has_perm("matrizaria.change_solicitacaoservicomatrizaria")
    )
    can_editar = (
        (is_solicitante_proprio or is_gestao)
        and solicitacao.status == "SOLICITADO"
        and solicitacao.ciclos_execucao.count() == 0
        and not user.username.startswith("tv")
    )
    can_excluir = (
        user.is_superuser
        and solicitacao.status == "SOLICITADO"
        and solicitacao.ciclos_execucao.count() == 0
        and not user.username.startswith("tv")
    )

    # Formulários pré-instanciados com a versão atual de concorrência
    form_finalizar = FinalizarExecucaoForm(initial={"versao": solicitacao.versao, "matriz_fisica": solicitacao.matriz_fisica})
    form_conferir = ConferirServicoForm(initial={"versao": solicitacao.versao})
    form_transferir = TransferirResponsabilidadeForm(solicitacao=solicitacao, initial={"versao": solicitacao.versao})
    form_cancelar = CancelarServicoForm(initial={"versao": solicitacao.versao})
    form_editar = EditarSolicitacaoForm(
        initial={
            "destino": solicitacao.destino,
            "prensa": "__MATRIZARIA__" if solicitacao.destino == "MATRIZARIA" else solicitacao.prensa_id,
            "tipo_servico": solicitacao.tipo_servico_id,
            "matriz_fisica": solicitacao.matriz_fisica_id,
            "prioridade": solicitacao.prioridade,
            "descricao_solicitacao": solicitacao.descricao_solicitacao,
            "versao": solicitacao.versao,
        }
    )

    context = {
        "solicitacao": solicitacao,
        "ciclos": ciclos,
        "ciclo_ativo": solicitacao.ciclos_execucao.filter(forma_encerramento="EM_ANDAMENTO").first(),
        "historico": historico,
        "can_executar": user_can_exec,
        "can_conferir": user_can_conf and not impedimento_autoconferencia,
        "impedimento_autoconferencia": impedimento_autoconferencia,
        "can_cancelar": user_can_canc,
        "can_editar": can_editar,
        "can_excluir": can_excluir,
        "form_finalizar": form_finalizar,
        "form_conferir": form_conferir,
        "form_transferir": form_transferir,
        "form_cancelar": form_cancelar,
        "form_editar": form_editar,
    }
    return render(request, "matrizaria/detalhe_solicitacao.html", context)


@login_required
@require_POST
def editar_solicitacao_view(request, pk):
    """
    Ação de editar a solicitação original antes do primeiro início técnico.
    """
    solicitacao = get_object_or_404(SolicitacaoServicoMatrizaria, pk=pk)
    user = request.user

    user_can_sol = _user_can_solicitar(user)
    user_can_conf = _user_can_conferir(user)
    is_solicitante_proprio = (solicitacao.solicitado_por_id == user.id and user_can_sol)
    is_gestao = (
        user.is_superuser
        or user.is_staff
        or user_can_conf
        or user.groups.filter(name__in=["Liderança de Produção", "Lideres", "Tecnicos_Lideres"]).exists()
        or user.has_perm("matrizaria.change_solicitacaoservicomatrizaria")
    )
    can_editar = (
        (is_solicitante_proprio or is_gestao)
        and solicitacao.status == "SOLICITADO"
        and solicitacao.ciclos_execucao.count() == 0
        and not user.username.startswith("tv")
    )

    if not can_editar:
        messages.error(request, "Acesso negado. Apenas o solicitante original ou líderes autorizados podem editar uma solicitação pendente.")
        return redirect("matrizaria:detalhe_servico", pk=pk)

    form = EditarSolicitacaoForm(request.POST)
    if form.is_valid():
        try:
            solicitacao = MatrizariaService.editar_solicitacao(
                solicitacao_id=pk,
                usuario=user,
                versao_esperada=form.cleaned_data["versao"],
                prensa=form.cleaned_data["prensa"],
                tipo_servico=form.cleaned_data["tipo_servico"],
                matriz_fisica=form.cleaned_data["matriz_fisica"],
                prioridade=form.cleaned_data["prioridade"],
                descricao_solicitacao=form.cleaned_data["descricao_solicitacao"],
                motivo_edicao=form.cleaned_data["motivo_edicao"],
                destino=form.cleaned_data.get("destino", "MAQUINA"),
            )
            messages.success(request, f"Solicitação SM #{solicitacao.id} corrigida com sucesso!")
        except ValidationError as e:
            messages.error(request, str(e.message if hasattr(e, "message") else e))
        except Exception as e:
            messages.error(request, f"Erro ao editar solicitação: {str(e)}")
    else:
        for err in form.errors.values():
            messages.error(request, err.as_text())

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
def excluir_solicitacao_view(request, pk):
    """
    Exclusão operacional definitiva de solicitação pendente antes de qualquer início técnico.
    """
    solicitacao = get_object_or_404(SolicitacaoServicoMatrizaria, pk=pk)
    user = request.user

    user_can_sol = _user_can_solicitar(user)
    user_can_conf = _user_can_conferir(user)
    is_solicitante_proprio = (solicitacao.solicitado_por_id == user.id and user_can_sol)
    is_gestao = (
        user.is_superuser
        or user.is_staff
        or user_can_conf
        or user.groups.filter(name__in=["Liderança de Produção", "Lideres", "Tecnicos_Lideres"]).exists()
        or user.has_perm("matrizaria.delete_solicitacaoservicomatrizaria")
    )
    can_excluir = (
        user.is_superuser
        and solicitacao.status == "SOLICITADO"
        and solicitacao.ciclos_execucao.count() == 0
        and not user.username.startswith("tv")
    )

    if not can_excluir:
        messages.error(request, "Acesso negado. Apenas o superusuário pode excluir solicitações do sistema.")
        return redirect("matrizaria:detalhe_servico", pk=pk)

    try:
        MatrizariaService.excluir_solicitacao_operacional(solicitacao_id=pk, usuario=user)
        messages.success(request, f"Solicitação SM #{pk} excluída com sucesso.")
        return redirect("matrizaria:kanban")
    except (ValidationError, PermissionDenied) as e:
        messages.error(request, str(e.message if hasattr(e, "message") else e))
        return redirect("matrizaria:detalhe_servico", pk=pk)
    except Exception as e:
        messages.error(request, f"Erro ao excluir solicitação: {str(e)}")
        return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
@matrizaria_executar_required
def iniciar_atendimento_view(request, pk):
    """
    Ação de iniciar atendimento técnico e abrir ciclo de execução.
    """
    versao = request.POST.get("versao")
    try:
        versao_int = int(versao)
    except (ValueError, TypeError):
        messages.error(request, "Versão inválida da solicitação.")
        return redirect("matrizaria:detalhe_servico", pk=pk)

    try:
        solicitacao = MatrizariaService.iniciar_atendimento(
            solicitacao_id=pk,
            usuario=request.user,
            versao_esperada=versao_int,
        )
        messages.success(request, f"Atendimento da SM #{solicitacao.id} iniciado com sucesso por você!")
    except ValidationError as e:
        messages.error(request, str(e.message if hasattr(e, "message") else e))
    except Exception as e:
        messages.error(request, f"Erro ao iniciar atendimento: {str(e)}")

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
@matrizaria_executar_required
def transferir_responsabilidade_view(request, pk):
    """
    Transfere a responsabilidade de atendimento em execução para outro técnico autorizado.
    """
    solicitacao = get_object_or_404(SolicitacaoServicoMatrizaria, pk=pk)
    form = TransferirResponsabilidadeForm(request.POST, solicitacao=solicitacao)
    if form.is_valid():
        try:
            solicitacao = MatrizariaService.transferir_responsabilidade(
                solicitacao_id=pk,
                usuario_origem=request.user,
                novo_responsavel=form.cleaned_data["novo_responsavel"],
                motivo=form.cleaned_data["motivo_transferencia"],
                versao_esperada=form.cleaned_data["versao"],
            )
            messages.success(
                request,
                f"Responsabilidade da SM #{solicitacao.id} transferida para {solicitacao.responsavel_atribuido_nome}.",
            )
        except ValidationError as e:
            messages.error(request, str(e.message if hasattr(e, "message") else e))
        except Exception as e:
            messages.error(request, f"Erro ao transferir: {str(e)}")
    else:
        for err in form.errors.values():
            messages.error(request, err.as_text())

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
@matrizaria_executar_required
def finalizar_execucao_view(request, pk):
    """
    Encerra o ciclo de execução técnica e encaminha para conferência/inspeção.
    """
    form = FinalizarExecucaoForm(request.POST)
    if form.is_valid():
        try:
            solicitacao = MatrizariaService.finalizar_execucao(
                solicitacao_id=pk,
                usuario=request.user,
                descricao_servico_realizado=form.cleaned_data["descricao_servico_realizado"],
                matriz_fisica=form.cleaned_data.get("matriz_fisica"),
                versao_esperada=form.cleaned_data["versao"],
            )
            messages.success(
                request,
                f"Execução técnica da SM #{solicitacao.id} finalizada com sucesso! Encaminhado para conferência.",
            )
        except ValidationError as e:
            messages.error(request, str(e.message if hasattr(e, "message") else e))
        except Exception as e:
            messages.error(request, f"Erro ao finalizar execução: {str(e)}")
    else:
        for err in form.errors.values():
            messages.error(request, err.as_text())

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
@matrizaria_conferir_required
def conferir_servico_view(request, pk):
    """
    Realiza a conferência do serviço: APROVAR (Conclui) ou DEVOLVER_RETRABALHO (Retorna à fila).
    Aplica validação rigorosa contra autoconferência.
    """
    form = ConferirServicoForm(request.POST)
    if form.is_valid():
        aprovado = form.cleaned_data["acao"] == "APROVAR"
        observacao = (
            form.cleaned_data["observacoes"]
            if aprovado
            else form.cleaned_data["motivo_devolucao"]
        )

        try:
            solicitacao = MatrizariaService.conferir_solicitacao(
                solicitacao_id=pk,
                usuario_conferente=request.user,
                aprovado=aprovado,
                versao_esperada=form.cleaned_data["versao"],
                observacao=observacao,
            )
            if aprovado:
                messages.success(request, f"SM #{solicitacao.id} aprovada e CONCLUÍDA com sucesso!")
            else:
                messages.warning(
                    request,
                    f"SM #{solicitacao.id} devolvida para RETRABALHO. Motivo registrado no histórico.",
                )
        except ValidationError as e:
            messages.error(request, str(e.message if hasattr(e, "message") else e))
        except Exception as e:
            messages.error(request, f"Erro na conferência: {str(e)}")
    else:
        for err in form.errors.values():
            messages.error(request, err.as_text())

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@require_POST
@matrizaria_cancelar_required
def cancelar_servico_view(request, pk):
    """
    Cancela justificadamente a solicitação e interrompe ciclo de execução ativo se houver.
    """
    form = CancelarServicoForm(request.POST)
    if form.is_valid():
        try:
            solicitacao = MatrizariaService.cancelar_solicitacao(
                solicitacao_id=pk,
                usuario=request.user,
                versao_esperada=form.cleaned_data["versao"],
                motivo_cancelamento=form.cleaned_data["motivo_cancelamento"],
            )
            messages.info(request, f"SM #{solicitacao.id} foi CANCELADA. Motivo registrado no histórico.")
        except ValidationError as e:
            messages.error(request, str(e.message if hasattr(e, "message") else e))
        except Exception as e:
            messages.error(request, f"Erro ao cancelar: {str(e)}")
    else:
        for err in form.errors.values():
            messages.error(request, err.as_text())

    return redirect("matrizaria:detalhe_servico", pk=pk)


@login_required
@matrizaria_relatorios_required
def relatorios_view(request):
    """
    Consulta e relatórios em tela com 4 critérios temporais canônicos.
    """
    hoje = timezone.localdate()
    sete_dias_atras = hoje - timedelta(days=7)

    # Parâmetros GET ou padrões
    criterio_temporal = request.GET.get("criterio_temporal", "COM_EXECUCAO")
    data_inicio_str = request.GET.get("data_inicio", sete_dias_atras.strftime("%Y-%m-%d"))
    data_fim_str = request.GET.get("data_fim", hoje.strftime("%Y-%m-%d"))
    prensa_id = request.GET.get("prensa") or None
    tipo_servico_id = request.GET.get("tipo_servico") or None
    status = request.GET.get("status") or None
    solicitante_id = request.GET.get("solicitante") or None
    executante_id = request.GET.get("executante") or None

    form = RelatorioFiltroForm(request.GET or {
        "criterio_temporal": criterio_temporal,
        "data_inicio": data_inicio_str,
        "data_fim": data_fim_str,
    })

    try:
        data_inicio = date.fromisoformat(data_inicio_str)
        data_fim = date.fromisoformat(data_fim_str)
    except (ValueError, TypeError):
        data_inicio = sete_dias_atras
        data_fim = hoje

    prensa_val = request.GET.get("prensa") or None
    destino = None
    prensa_id_int = None
    if prensa_val == "__MATRIZARIA__":
        destino = "MATRIZARIA"
    elif prensa_val and prensa_val.isdigit():
        prensa_id_int = int(prensa_val)

    tipo_servico_id_int = int(tipo_servico_id) if tipo_servico_id and tipo_servico_id.isdigit() else None
    solicitante_id_int = int(solicitante_id) if solicitante_id and solicitante_id.isdigit() else None
    executante_id_int = int(executante_id) if executante_id and executante_id.isdigit() else None

    todos_resultados = MatrizariaService.consultar_relatorio(
        criterio_temporal=criterio_temporal,
        data_inicio=data_inicio,
        data_fim=data_fim,
        prensa_id=prensa_id_int,
        tipo_servico_id=tipo_servico_id_int,
        status=status,
        solicitante_id=solicitante_id_int,
        executante_id=executante_id_int,
        destino=destino,
    )

    # Paginação em tela (20 por página)
    paginator = Paginator(todos_resultados, 20)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "form": form,
        "page_obj": page_obj,
        "total_registros": len(todos_resultados),
        "criterio_temporal": criterio_temporal,
        "data_inicio": data_inicio,
        "data_fim": data_fim,
    }
    return render(request, "matrizaria/relatorios.html", context)


@login_required
@matrizaria_relatorios_required
def exportar_excel_view(request):
    """
    Exporta os relatórios filtrados para arquivo Excel com 2 abas:
    1. 'Serviços': uma linha por solicitação com durações, status e snapshots.
    2. 'Histórico': cronologia completa de ciclos e transições das solicitações selecionadas.
    Protegido contra CSV/Excel formula injection.
    """
    hoje = timezone.localdate()
    sete_dias_atras = hoje - timedelta(days=7)

    criterio_temporal = request.GET.get("criterio_temporal", "COM_EXECUCAO")
    data_inicio_str = request.GET.get("data_inicio", sete_dias_atras.strftime("%Y-%m-%d"))
    data_fim_str = request.GET.get("data_fim", hoje.strftime("%Y-%m-%d"))
    prensa_id = request.GET.get("prensa") or None
    tipo_servico_id = request.GET.get("tipo_servico") or None
    status = request.GET.get("status") or None
    solicitante_id = request.GET.get("solicitante") or None
    executante_id = request.GET.get("executante") or None

    try:
        data_inicio = date.fromisoformat(data_inicio_str)
        data_fim = date.fromisoformat(data_fim_str)
    except (ValueError, TypeError):
        data_inicio = sete_dias_atras
        data_fim = hoje

    prensa_val = request.GET.get("prensa") or None
    destino = None
    prensa_id_int = None
    if prensa_val == "__MATRIZARIA__":
        destino = "MATRIZARIA"
    elif prensa_val and prensa_val.isdigit():
        prensa_id_int = int(prensa_val)

    tipo_servico_id_int = int(tipo_servico_id) if tipo_servico_id and tipo_servico_id.isdigit() else None
    solicitante_id_int = int(solicitante_id) if solicitante_id and solicitante_id.isdigit() else None
    executante_id_int = int(executante_id) if executante_id and executante_id.isdigit() else None

    # Consulta TODOS os registros filtrados (sem paginação!)
    resultados = MatrizariaService.consultar_relatorio(
        criterio_temporal=criterio_temporal,
        data_inicio=data_inicio,
        data_fim=data_fim,
        prensa_id=prensa_id_int,
        tipo_servico_id=tipo_servico_id_int,
        status=status,
        solicitante_id=solicitante_id_int,
        executante_id=executante_id_int,
        destino=destino,
    )

    wb = openpyxl.Workbook()

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # ----------------------------------------------------
    # ABA 1: SERVIÇOS
    # ----------------------------------------------------
    ws_servicos = wb.active
    ws_servicos.title = "Serviços"
    ws_servicos.views.sheetView[0].showGridLines = True

    cabecalhos_servicos = [
        "Protocolo",
        "Prensa",
        "Tipo de Serviço",
        "Matriz Física",
        "Status Atual",
        "Status Fim Período",
        "Solicitante",
        "Data Abertura",
        "Executantes",
        "Primeiro Início",
        "Último Término",
        "Duração Total Intervenção",
        "Duração no Período",
        "Ciclos",
        "Retrabalhos",
        "Conferente",
        "Data Conferência",
        "Cancelado Por",
        "Data Cancelamento",
        "Motivo Cancelamento",
        "Motivo Devolução Retrabalho",
        "Descrição Solicitada",
        "Último Serviço Executado",
    ]

    ws_servicos.append(cabecalhos_servicos)
    for col_idx in range(1, len(cabecalhos_servicos) + 1):
        cell = ws_servicos.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item in resultados:
        dt_abertura = item["data_solicitacao"].strftime("%d/%m/%Y %H:%M") if item["data_solicitacao"] else ""
        dt_inicio = item["data_inicio_execucao"].strftime("%d/%m/%Y %H:%M") if item["data_inicio_execucao"] else ""
        dt_fim = item["data_fim_execucao"].strftime("%d/%m/%Y %H:%M") if item["data_fim_execucao"] else ""
        dt_conf = item["data_conferencia"].strftime("%d/%m/%Y %H:%M") if item["data_conferencia"] else ""
        dt_canc = item["data_cancelamento"].strftime("%d/%m/%Y %H:%M") if item["data_cancelamento"] else ""

        linha = [
            sanitize_excel_cell(item["protocolo"]),
            sanitize_excel_cell(item["prensa_nome"]),
            sanitize_excel_cell(item["tipo_servico_nome"]),
            sanitize_excel_cell(item["matriz_identificador"]),
            sanitize_excel_cell(item["status_atual_label"]),
            sanitize_excel_cell(item["status_fim_periodo_label"]),
            sanitize_excel_cell(item["solicitante_nome"]),
            dt_abertura,
            sanitize_excel_cell(item["executantes_str"]),
            dt_inicio,
            dt_fim,
            sanitize_excel_cell(item["duracao_total_intervencao_str"]),
            sanitize_excel_cell(item["duracao_periodo_intervencao_str"]),
            item["quantidade_ciclos"],
            item["quantidade_retrabalhos"],
            sanitize_excel_cell(item["conferente_nome"] or ""),
            dt_conf,
            sanitize_excel_cell(item["cancelado_por_nome"] or ""),
            dt_canc,
            sanitize_excel_cell(item["motivo_cancelamento"] or ""),
            sanitize_excel_cell(item["motivo_devolucao_retrabalho"] or ""),
            sanitize_excel_cell(item["descricao_solicitacao"]),
            sanitize_excel_cell(item["descricao_servico_executado"]),
        ]
        ws_servicos.append(linha)

    # ----------------------------------------------------
    # ABA 2: HISTÓRICO
    # ----------------------------------------------------
    ws_hist = wb.create_sheet(title="Histórico")
    ws_hist.views.sheetView[0].showGridLines = True

    cabecalhos_hist = [
        "Protocolo",
        "Prensa",
        "Tipo Registro",
        "Número Ciclo",
        "Data/Hora Evento",
        "Usuário Responsável",
        "Status Anterior",
        "Status Novo",
        "Forma Encerramento",
        "Duração Ciclo",
        "Detalhes / Observações",
    ]

    ws_hist.append(cabecalhos_hist)
    for col_idx in range(1, len(cabecalhos_hist) + 1):
        cell = ws_hist.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in resultados:
        s = item["solicitacao"]

        # Ciclos de Execução
        for c in s.ciclos_execucao.all().order_by("numero_ciclo"):
            dt_c_inicio = c.data_inicio.strftime("%d/%m/%Y %H:%M") if c.data_inicio else ""
            linha_ciclo = [
                sanitize_excel_cell(f"SM #{s.id}"),
                sanitize_excel_cell(item["prensa_nome"]),
                "CICLO_EXECUCAO",
                c.numero_ciclo,
                dt_c_inicio,
                sanitize_excel_cell(c.usuario_inicio_nome),
                "EM_EXECUCAO",
                "AGUARDANDO_CONFERENCIA" if c.forma_encerramento == "FINALIZADO_TECNICO" else c.forma_encerramento,
                sanitize_excel_cell(c.get_forma_encerramento_display()),
                sanitize_excel_cell(c.duracao_str),
                sanitize_excel_cell(c.descricao_servico_executado or ""),
            ]
            ws_hist.append(linha_ciclo)

        # Transições de Estado
        for h in s.historico_transicoes.all().order_by("data_evento"):
            dt_h = h.data_evento.strftime("%d/%m/%Y %H:%M:%S") if h.data_evento else ""
            linha_hist = [
                sanitize_excel_cell(f"SM #{s.id}"),
                sanitize_excel_cell(item["prensa_nome"]),
                sanitize_excel_cell(h.tipo_evento),
                "-",
                dt_h,
                sanitize_excel_cell(h.usuario_nome_snapshot),
                sanitize_excel_cell(h.status_anterior),
                sanitize_excel_cell(h.status_novo),
                "-",
                "-",
                sanitize_excel_cell(h.observacao or ""),
            ]
            ws_hist.append(linha_hist)

    # Autoajuste de largura das colunas nas duas abas
    for ws in [ws_servicos, ws_hist]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # Congelar primeira linha
    ws_servicos.freeze_panes = "A2"
    ws_hist.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"Relatorio_Matrizaria_{criterio_temporal}_{data_inicio.strftime('%Y%m%d')}_{data_fim.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
@matrizaria_tv_required
def tv_view(request):
    """
    Tela de TV da Matrizaria Industrial:
    Alterna periodicamente entre a fila de serviços técnicos (20s) e
    os alertas preventivos de limite de bladder (15s) da Vulcanização.
    """
    context = MatrizariaService.get_tv_dashboard_context()
    return render(request, "matrizaria/tv.html", context)


@login_required
@require_GET
@matrizaria_tv_required
def api_tv_data_view(request):
    """
    Endpoint JSON assíncrono para atualização contínua da TV sem recarregar a página inteira.
    Falha na telemetria SCADA de bladder é isolada e não derruba a exibição de serviços.
    """
    context = MatrizariaService.get_tv_dashboard_context()

    def serialize_solicitacao(s):
        matriz_txt = (
            s.matriz_identificador_snapshot
            or (s.matriz_fisica.nome_exibicao if s.matriz_fisica else "Não informada")
        )
        return {
            "id": s.id,
            "protocolo": f"SM #{s.id}",
            "prensa_nome": s.prensa_nome_snapshot or s.prensa.nome,
            "tipo_servico": s.tipo_servico_nome_snapshot or s.tipo_servico.nome,
            "matriz_identificador": matriz_txt,
            "prioridade": s.prioridade,
            "badge_prioridade": s.badge_prioridade_classe,
            "status": s.get_status_display(),
            "status_code": s.status,
            "badge_status": s.badge_status_classe,
            "responsavel": s.responsavel_atribuido_nome or "Aguardando",
            "tempo_espera": s.tempo_espera_str,
            "descricao": s.descricao_solicitacao,
            "retrabalhos": s.quantidade_retrabalhos,
        }

    servicos_json = {
        "total_solicitados": context["servicos"]["total_solicitados"],
        "total_retrabalhos": context["servicos"]["total_retrabalhos"],
        "total_em_execucao": context["servicos"]["total_em_execucao"],
        "total_aguardando_conferencia": context["servicos"]["total_aguardando_conferencia"],
        "solicitados": [serialize_solicitacao(s) for s in context["servicos"]["solicitados"]],
        "retrabalhos": [serialize_solicitacao(s) for s in context["servicos"]["retrabalhos"]],
        "em_execucao": [serialize_solicitacao(s) for s in context["servicos"]["em_execucao"]],
        "aguardando_conferencia": [serialize_solicitacao(s) for s in context["servicos"]["aguardando_conferencia"]],
    }

    from django.template.loader import render_to_string

    html_solicitados = render_to_string(
        "matrizaria/includes/tv_col_solicitados.html",
        {"fila_solicitados": context["servicos"]["fila_solicitados"]},
        request=request,
    )
    html_execucao = render_to_string(
        "matrizaria/includes/tv_col_execucao.html",
        {"em_execucao": context["servicos"]["em_execucao"]},
        request=request,
    )
    html_conferencia = render_to_string(
        "matrizaria/includes/tv_col_conferencia.html",
        {"aguardando_conferencia": context["servicos"]["aguardando_conferencia"]},
        request=request,
    )
    html_bladders = render_to_string(
        "matrizaria/includes/tv_col_bladders.html",
        {"bladder_alerts": context["bladder_alerts"]},
        request=request,
    )

    return JsonResponse({
        "servicos": servicos_json,
        "html_solicitados": html_solicitados,
        "html_execucao": html_execucao,
        "html_conferencia": html_conferencia,
        "html_bladders": html_bladders,
        "bladder_alerts": context["bladder_alerts"],
        "total_bladder_alerts": context["total_bladder_alerts"],
        "scada_status": context["scada_status"],
        "scada_mensagem": context["scada_mensagem"],
        "timestamp": context["timestamp"],
        "hora_atual": context["hora_atual"],
    })
