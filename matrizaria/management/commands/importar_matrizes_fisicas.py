import hashlib
import json
import os
import re
import sys
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction, models
from django.utils import timezone
from django.contrib.auth.models import User

from production.models import ProductionMatrixCatalog
from matrizaria.models import MatrizFisica, LoteImportacaoMatrizFisica, get_canonical_tooling_model


DEFAULT_EXCEL_PATH = r"C:\Users\Unicompo\Documents\03_PYTHON1\07 - Manutencao_estudo\MATRIZES_PRE_CADASTRO_LOTE.xlsx"
DEFAULT_RECONCILIATION_PATH = r"C:\Users\Unicompo\Documents\03_PYTHON1\07 - Manutencao_estudo\relatorio_reconciliacao_inventario.xlsx"


def compute_file_hash(filepath: str) -> str:
    """Calcula hash SHA-256 do arquivo exclusivamente para auditoria e histórico."""
    sha256 = hashlib.sha256()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


def clean_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def normalize_header(header: str) -> str:
    """Normaliza nomes de cabeçalhos para busca flexível e tolerante a acentos."""
    h = clean_str(header).lower()
    replacements = {
        "á": "a", "à": "a", "ã": "a", "â": "a",
        "é": "e", "ê": "e",
        "í": "i",
        "ó": "o", "ô": "o", "õ": "o",
        "ú": "u",
        "ç": "c",
    }
    for orig, rep in replacements.items():
        h = h.replace(orig, rep)
    h = re.sub(r"[^a-z0-9]", "", h)
    return h


def clean_measure(val) -> str:
    """Remove espaços, barras, pontos e hífens para conferência semântica estrita de medida."""
    v = clean_str(val).upper().replace(" ", "")
    v = re.sub(r"-0(\d)", r"-\1", v)  # Normaliza aro com zero à esquerda (ex: -08 -> -8)
    return re.sub(r"[\/\.\-]", "", v)


def extract_measure_tokens(text: str) -> str:
    """Extrai os componentes de medida ignorando prefixos como PNEU, PNEUS ou marca."""
    t = clean_str(text).upper()
    t = re.sub(r"^PNEUS?\s+", "", t)
    for brand in ["HOPPER", "SPEEDY", "WINGS", "READY", "OPTION", "WINTER", "RIVER", "ROBOT"]:
        t = re.sub(r"^" + brand + r"\s+", "", t)
    t = re.sub(r"\s+S/C$", "", t)
    return clean_measure(t)


def check_catalog_compatibility(cat_obj: ProductionMatrixCatalog, modelo_lido: str, med_proposta: str, med_lida: str) -> bool:
    """
    Verifica se um item do catálogo é semanticamente compatível com a linha.
    Evita que o usuário associe acidentalmente um ID de outro modelo/medida.
    """
    if not cat_obj:
        return False
    cat_name = (cat_obj.nome_exibicao or cat_obj.nome_scada or cat_obj.produto or cat_obj.descricao or "").upper()
    
    # 1. Marca / Modelo
    mod_norm = clean_str(modelo_lido).upper()
    if mod_norm not in cat_name:
        return False
        
    # 2. Medida
    cat_measure = extract_measure_tokens(cat_name)
    prop_clean = clean_measure(med_proposta)
    lida_clean = clean_measure(med_lida)
    
    if prop_clean and prop_clean == cat_measure:
        return True
    if lida_clean and lida_clean == cat_measure:
        return True
    return False


def build_missing_model_code(modelo: str, medida: str) -> str:
    """Gera código interno alfanumérico limpo para modelos ausentes (sem inventar código SCADA)."""
    m_clean = clean_str(medida).replace("/", "").replace(".", "").replace("-", "").replace(" ", "").upper()
    mod_clean = clean_str(modelo).upper()
    return f"INT-{mod_clean}-{m_clean}"


class Command(BaseCommand):
    help = (
        "Importador e simulador em lote de matrizes físicas a partir da planilha de inventário. "
        "Separa correspondência de catálogo, conferência da fonte e autorização de carga. "
        "Não inventa códigos SCADA, não escolhe variantes S/C por conveniência e previne duplicações."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "arquivo",
            nargs="?",
            type=str,
            default=DEFAULT_EXCEL_PATH,
            help=f"Caminho do arquivo Excel (padrão: {DEFAULT_EXCEL_PATH})",
        )
        parser.add_argument(
            "--simulacao",
            action="store_true",
            default=True,
            help="Modo de simulação/prévia (padrão ativo: não grava alterações no banco)",
        )
        parser.add_argument(
            "--aplicar",
            action="store_true",
            default=False,
            help="Efetiva a gravação no banco de dados para as linhas aprovadas e mapeadas.",
        )
        parser.add_argument(
            "--responsavel",
            type=str,
            default="",
            help="Identificação do operador/responsável técnico pela aplicação do lote (obrigatório se --aplicar).",
        )
        parser.add_argument(
            "--aprovar-linhas-seguras",
            action="store_true",
            default=False,
            help="Aprova em lote para carga apenas as linhas com correspondência ÚNICA no catálogo E fonte 100%% conferida.",
        )
        parser.add_argument(
            "--carga-inicial",
            action="store_true",
            default=False,
            help="Autoriza a carga inicial do inventário (52 combinações / 78 unidades), "
                 "adotando premissas de normalização de medidas, cadastrando modelos ausentes "
                 "e preservando ressalvas técnicas da fonte.",
        )
        parser.add_argument(
            "--linhas",
            type=str,
            default="",
            help="Lista opcional de IDs de linhas a processar, separados por vírgula (ex: P1-22,P1-23,P2-06).",
        )
        parser.add_argument(
            "--planilha-reconciliacao",
            type=str,
            default=DEFAULT_RECONCILIATION_PATH,
            help=f"Caminho para gerar a planilha de reconciliação Excel (padrão: {DEFAULT_RECONCILIATION_PATH})",
        )

    def handle(self, *args, **options):
        caminho_arquivo = options["arquivo"]
        aplicar = options["aplicar"]
        simulacao = not aplicar
        responsavel = options["responsavel"].strip()
        aprovar_seguras = options["aprovar_linhas_seguras"]
        carga_inicial = options["carga_inicial"]
        linhas_filtro_raw = options["linhas"].strip()
        caminho_reconciliacao = options["planilha_reconciliacao"]

        linhas_filtro = set()
        if linhas_filtro_raw:
            linhas_filtro = {l.strip().upper() for l in linhas_filtro_raw.split(",") if l.strip()}

        if aplicar and not responsavel:
            raise CommandError(
                "O parâmetro --responsavel <NOME> é OBRIGATÓRIO ao aplicar a carga real no banco de dados."
            )

        responsavel_audit = responsavel
        if responsavel:
            user_match = User.objects.filter(
                models.Q(username__iexact=responsavel) | models.Q(first_name__icontains=responsavel)
            ).first()
            if user_match:
                responsavel_audit = user_match.username

        if not os.path.exists(caminho_arquivo):
            raise CommandError(f"Arquivo não encontrado no caminho: {caminho_arquivo}")

        file_hash = compute_file_hash(caminho_arquivo)
        file_name = os.path.basename(caminho_arquivo)

        self.stdout.write("=" * 80)
        self.stdout.write("  FERRAMENTA DE INVENTÁRIO E CADASTRO EM LOTE DE MATRIZES FÍSICAS")
        self.stdout.write("=" * 80)
        self.stdout.write(f"Arquivo: {caminho_arquivo}")
        self.stdout.write(f"Hash SHA-256: {file_hash} (Metadado de auditoria)")
        modo_str = "APLICAÇÃO REAL NO BANCO" if aplicar else "SIMULAÇÃO (DRY-RUN - SEM GRAVAÇÃO)"
        cor_modo = self.style.WARNING if aplicar else self.style.SUCCESS
        self.stdout.write(f"Modo de Execução: {cor_modo(modo_str)}")
        if carga_inicial:
            self.stdout.write(self.style.SUCCESS("Regime: CARGA INICIAL AUTORIZADA (52 combinações / 78 unidades)"))
        if aplicar:
            self.stdout.write(f"Responsável: {responsavel_audit}")
        self.stdout.write("-" * 80)

        # 1. Carregar planilha
        try:
            wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        except Exception as e:
            raise CommandError(f"Falha ao abrir planilha Excel: {e}")

        if "Inventario" not in wb.sheetnames:
            raise CommandError(
                f"A aba 'Inventario' é a fonte obrigatória e não foi encontrada. Abas disponíveis: {wb.sheetnames}"
            )

        ws = wb["Inventario"]

        # 2. Localizar linha de cabeçalho dinamicamente
        header_row_idx = None
        header_map = {}
        for r in range(1, 20):
            first_val = ws.cell(r, 1).value
            if first_val and "id da linha" in clean_str(first_val).lower():
                header_row_idx = r
                break

        if not header_row_idx:
            raise CommandError("Não foi possível localizar a linha de cabeçalho contendo 'ID da linha' na aba Inventario.")

        expected_cols = {
            "iddalinha": "id_linha",
            "medidalidanafoto": "medida_lida",
            "medidaproposta": "medida_proposta",
            "modelolido": "modelo_lido",
            "quantidadetotal": "quantidade_total",
            "comdote": "com_dote",
            "semdote": "sem_dote",
            "dotenaoinformado": "dote_nao_informado",
            "ultimacolunaoriginal": "ultima_coluna_orig",
            "revisaodaleitura": "revisao_leitura",
            "aprovardados": "aprovar_dados",
            "idnocatalogo": "id_catalogo",
            "observacoeslimitacoes": "observacoes",
            "fonte": "fonte",
        }

        for c in range(1, ws.max_column + 1):
            val = ws.cell(header_row_idx, c).value
            if val is not None:
                norm = normalize_header(str(val))
                for exp_norm, field_key in expected_cols.items():
                    if exp_norm in norm:
                        header_map[field_key] = c

        missing_fields = [f for f in ["id_linha", "medida_proposta", "modelo_lido", "quantidade_total", "com_dote", "sem_dote", "dote_nao_informado"] if f not in header_map]
        if missing_fields:
            raise CommandError(f"Colunas obrigatórias não identificadas no cabeçalho: {missing_fields}")

        # 3. Ler dados da aba Inventario
        linhas_lidas = []
        for r in range(header_row_idx + 1, ws.max_row + 1):
            id_val = clean_str(ws.cell(r, header_map["id_linha"]).value)
            if not id_val:
                continue

            def get_cell_val(key):
                if key in header_map:
                    return ws.cell(r, header_map[key]).value
                return None

            linhas_lidas.append({
                "row_num": r,
                "id_linha": id_val,
                "medida_lida": clean_str(get_cell_val("medida_lida")),
                "medida_proposta": clean_str(get_cell_val("medida_proposta")),
                "modelo_lido": clean_str(get_cell_val("modelo_lido")).upper(),
                "quantidade_total": get_cell_val("quantidade_total"),
                "com_dote": get_cell_val("com_dote"),
                "sem_dote": get_cell_val("sem_dote"),
                "dote_nao_informado": get_cell_val("dote_nao_informado"),
                "ultima_coluna_orig": clean_str(get_cell_val("ultima_coluna_orig")),
                "revisao_leitura": clean_str(get_cell_val("revisao_leitura")),
                "aprovar_dados": clean_str(get_cell_val("aprovar_dados")).upper(),
                "id_catalogo": get_cell_val("id_catalogo"),
                "observacoes": clean_str(get_cell_val("observacoes")),
                "fonte": clean_str(get_cell_val("fonte")),
            })

        # 4. Catálogo Canônico SCADA
        all_catalogs = list(ProductionMatrixCatalog.objects.all())
        catalog_by_id = {c.id: c for c in all_catalogs}
        catalog_by_scada = {c.codigo_scada: c for c in all_catalogs if c.codigo_scada is not None}

        # 5. Processamento e Confronto Detalhado
        relatorio_linhas = []
        propostas_novos_modelos = []
        totais_contagem = {
            "linhas_lidas": len(linhas_lidas),
            "unidades_lidas": 0,
            "linhas_unicas": 0,
            "unidades_unicas": 0,
            "linhas_unicas_conferidas": 0,
            "unidades_unicas_conferidas": 0,
            "linhas_unicas_fonte_pendente": 0,
            "unidades_unicas_fonte_pendente": 0,
            "linhas_ambiguas": 0,
            "unidades_ambiguas": 0,
            "linhas_ausentes": 0,
            "unidades_ausentes": 0,
            "linhas_inconsistentes": 0,
            "unidades_propostas_novas": 0,
            "unidades_existentes_preservadas": 0,
            "unidades_com_conflito": 0,
            "linhas_aprovadas_escopo": 0,
            "unidades_aprovadas_escopo": 0,
        }

        lote_id = f"LOTE-INV-{timezone.now().strftime('%Y%m%d-%H%M%S%f')[:17]}-{uuid.uuid4().hex[:6]}"

        for item in linhas_lidas:
            lid = item["id_linha"]

            # Filtro opcional de linhas
            if linhas_filtro and lid.upper() not in linhas_filtro:
                continue

            # Validação matemática de quantidades
            try:
                tot = int(item["quantidade_total"] or 0)
                c_dote = int(item["com_dote"] or 0)
                s_dote = int(item["sem_dote"] or 0)
                ni_dote = int(item["dote_nao_informado"] or 0)
            except (ValueError, TypeError):
                tot, c_dote, s_dote, ni_dote = -1, 0, 0, 0

            math_ok = (tot > 0) and (c_dote >= 0) and (s_dote >= 0) and (ni_dote >= 0) and (tot == c_dote + s_dote + ni_dote)

            if not math_ok:
                totais_contagem["linhas_inconsistentes"] += 1
                relatorio_linhas.append({
                    "item": item,
                    "status_cat": "INCONSISTENTE_MATEMATICA",
                    "status_fonte": "PENDENTE_CONFERENCIA",
                    "pendencias_fonte": ["Soma de dotes diverge do total"],
                    "candidatos": [],
                    "modelo_escolhido": None,
                    "motivo_bloqueio": f"Quantidade total ({tot}) não confere com soma de dotes ({c_dote}+{s_dote}+{ni_dote})",
                    "aprovado": False,
                    "unidades_planejadas": [],
                })
                continue

            totais_contagem["unidades_lidas"] += tot

            # Avaliação de Conferência da Fonte (dimensão separada do catálogo)
            pendencias_fonte = []
            rev_raw = item["revisao_leitura"].upper()
            if "CONFERIR" in rev_raw:
                pendencias_fonte.append(f"Revisão da anotação indica '{item['revisao_leitura']}'")
            if ni_dote > 0:
                pendencias_fonte.append(f"Possui {ni_dote} unidade(s) com dote não informado na foto")
            if lid == "P2-23":
                pendencias_fonte.append("Medida 4.80/4.00-8 é proposta de normalização de 4.80/400-8 pendente de confirmação humana")
            elif lid in ["P2-12", "P2-13", "P2-14"]:
                pendencias_fonte.append(f"Medida original escrita com barra ('{item['medida_lida']}') normalizada para ponto ('{item['medida_proposta']}')")
            elif item["medida_lida"] and item["medida_proposta"] and clean_measure(item["medida_lida"]) != clean_measure(item["medida_proposta"]):
                pendencias_fonte.append(f"Medida original '{item['medida_lida']}' difere da proposta '{item['medida_proposta']}'")

            if pendencias_fonte:
                status_fonte = "PENDENTE_CONFERENCIA"
            else:
                status_fonte = "CONFERIDO"

            # Matching no catálogo
            candidatos = []
            id_cat_raw = item["id_catalogo"]
            id_incompativel_detectado = False

            if id_cat_raw is not None and clean_str(id_cat_raw):
                try:
                    cid = int(id_cat_raw)
                    cat_especifico = catalog_by_id.get(cid) or catalog_by_scada.get(cid)
                    if cat_especifico:
                        cat_canon = get_canonical_tooling_model(cat_especifico)
                        if check_catalog_compatibility(cat_canon, item["modelo_lido"], item["medida_proposta"], item["medida_lida"]):
                            candidatos = [cat_canon]
                        else:
                            id_incompativel_detectado = True
                            candidatos = []
                    else:
                        id_incompativel_detectado = True
                        candidatos = []
                except ValueError:
                    id_incompativel_detectado = True

            if id_incompativel_detectado:
                status_cat = "ID_INCOMPATIVEL"
                motivo_id_incompativel = f"ID de catálogo '{id_cat_raw}' fornecido na planilha não existe ou é incompatível com '{item['modelo_lido']} {item['medida_proposta']}'."
            else:
                motivo_id_incompativel = ""

            # Se não tem ID explícito válido, confrontar modelo e medida
            if not candidatos and not id_incompativel_detectado:
                mod_lido = item["modelo_lido"]
                med_prop = item["medida_proposta"]
                med_lida = item["medida_lida"]
                prop_clean = clean_measure(med_prop)
                lida_clean = clean_measure(med_lida)

                raw_candidatos = []
                for cat in all_catalogs:
                    cat_name = (cat.nome_exibicao or cat.nome_scada or cat.produto or cat.descricao or "").upper()
                    if mod_lido not in cat_name:
                        continue

                    cat_measure = extract_measure_tokens(cat_name)
                    if (prop_clean and prop_clean == cat_measure) or (lida_clean and lida_clean == cat_measure):
                        raw_candidatos.append(cat)

                # Regra Definitiva de Ferramental Compartilhado:
                # Mapeia todas as variantes para o modelo canônico administrativo (sem S/C)
                # e consolida candidatos únicos por id do modelo canônico.
                candidatos = list({get_canonical_tooling_model(c).id: get_canonical_tooling_model(c) for c in raw_candidatos}.values())

            # Classificação de Catálogo
            if id_incompativel_detectado:
                status_cat = "ID_INCOMPATIVEL"
            elif len(candidatos) == 1:
                status_cat = "UNICA"
                totais_contagem["linhas_unicas"] += 1
                totais_contagem["unidades_unicas"] += tot
                if status_fonte == "CONFERIDO":
                    totais_contagem["linhas_unicas_conferidas"] += 1
                    totais_contagem["unidades_unicas_conferidas"] += tot
                else:
                    totais_contagem["linhas_unicas_fonte_pendente"] += 1
                    totais_contagem["unidades_unicas_fonte_pendente"] += tot
            elif len(candidatos) > 1:
                status_cat = "AMBIGUA"
                totais_contagem["linhas_ambiguas"] += 1
                totais_contagem["unidades_ambiguas"] += tot
            else:
                status_cat = "AUSENTE"
                totais_contagem["linhas_ausentes"] += 1
                totais_contagem["unidades_ausentes"] += tot
                # Proposta para cadastro do modelo ausente
                propostas_novos_modelos.append({
                    "linha_origem": lid,
                    "modelo": item["modelo_lido"],
                    "medida_proposta": item["medida_proposta"],
                    "medida_lida": item["medida_lida"],
                    "codigo_interno_proposto": build_missing_model_code(item["modelo_lido"], item["medida_proposta"]),
                    "nome_exibicao_proposto": f"PNEU {item['modelo_lido']} {item['medida_proposta']}",
                    "codigo_scada": "[Nenhum - Sem Telemetria SCADA]",
                    "variante_sc": False,
                    "ativo": True,
                    "tempo_producao_padrao": 0,
                    "tempo_vulcanizacao_padrao": 0,
                    "pendencias": "; ".join(pendencias_fonte) if pendencias_fonte else "Nenhuma pendência na fonte",
                    "observacoes": item["observacoes"],
                })

            # Determinação de Aprovação para Carga
            is_aprovada = False
            motivo_bloqueio = ""

            if id_incompativel_detectado:
                is_aprovada = False
                motivo_bloqueio = motivo_id_incompativel
            elif carga_inicial:
                # Carga inicial autorizada pelo usuário:
                # Normalização de medidas aceita, variantes unificadas na referência canônica sem S/C,
                # modelos ausentes aprovados para cadastro em lote e ressalvas preservadas.
                is_aprovada = True
                motivo_bloqueio = ""
            elif item["aprovar_dados"] == "APROVADO":
                if status_cat == "UNICA":
                    # Pendências impeditivas de leitura que mesmo aprovação manual preliminar não pode ignorar
                    pendencias_bloqueantes = [
                        p for p in pendencias_fonte
                        if "CONFERIR" in p.upper() or "P2-23" in p.upper()
                    ]
                    if pendencias_bloqueantes:
                        is_aprovada = False
                        motivo_bloqueio = f"Marcada APROVADA na planilha, mas bloqueada por pendência na leitura original: {'; '.join(pendencias_bloqueantes)}"
                    else:
                        is_aprovada = True
                elif status_cat == "AMBIGUA":
                    motivo_bloqueio = f"Marcada APROVADA na planilha, mas é ambígua ({len(candidatos)} variantes S/C). Preencher o ID exato na coluna 'ID no catálogo'."
                else:
                    motivo_bloqueio = "Marcada APROVADA na planilha, mas modelo ausente no catálogo SCADA. Use --carga-inicial para incluir."
            elif aprovar_seguras:
                if status_cat == "UNICA" and status_fonte == "CONFERIDO":
                    is_aprovada = True
                elif status_cat == "UNICA" and status_fonte != "CONFERIDO":
                    is_aprovada = False
                    motivo_bloqueio = f"Correspondência única no catálogo, mas bloqueada por pendência na fonte: {'; '.join(pendencias_fonte)}"
                elif status_cat == "AMBIGUA":
                    variantes_str = ", ".join([f"ID={c.id} SCADA={c.codigo_scada} ({c.nome_exibicao})" for c in candidatos])
                    motivo_bloqueio = f"Ambíguo no catálogo ({len(candidatos)} variantes): {variantes_str}. Aguarda confirmação se ferramental é compartilhado ou distinto."
                else:
                    motivo_bloqueio = f"Modelo '{item['modelo_lido']} {item['medida_proposta']}' ausente no catálogo de 55 códigos SCADA. Requer cadastro prévio no catálogo."
            else:
                if status_cat == "UNICA":
                    if status_fonte == "CONFERIDO":
                        motivo_bloqueio = "Pronta para carga, mas pendente de aprovação (use --carga-inicial ou --aprovar-linhas-seguras)."
                    else:
                        motivo_bloqueio = f"Pendente de conferência da fonte: {'; '.join(pendencias_fonte)}"
                elif status_cat == "AMBIGUA":
                    variantes_str = ", ".join([f"ID={c.id} SCADA={c.codigo_scada} ({c.nome_exibicao})" for c in candidatos])
                    motivo_bloqueio = f"Ambíguo no catálogo ({len(candidatos)} variantes): {variantes_str}. Não adivinhar variante S/C."
                else:
                    motivo_bloqueio = f"Modelo '{item['modelo_lido']} {item['medida_proposta']}' ausente no catálogo de 55 códigos SCADA. Use --carga-inicial para cadastrar."

            # Planejamento das unidades físicas (expansão)
            unidades_planejadas = []
            if status_cat == "UNICA" or (status_cat == "AUSENTE" and carga_inicial):
                if status_cat == "UNICA":
                    modelo_cat = candidatos[0]
                else:
                    cod_int = build_missing_model_code(item["modelo_lido"], item["medida_proposta"])
                    modelo_cat = ProductionMatrixCatalog.objects.filter(codigo=cod_int).first()

                if modelo_cat:
                    existentes_modelo = list(
                        MatrizFisica.objects.filter(modelo=modelo_cat).order_by("numero_sequencial")
                    )
                    existentes_desta_linha = [
                        m for m in existentes_modelo if m.linha_origem == lid or (m.chave_unidade_origem and m.chave_unidade_origem.startswith(f"{lid}-"))
                    ]
                else:
                    existentes_modelo = []
                    existentes_desta_linha = []

                # Proteção contra alteração de quantidade sem reconciliação:
                # Se a linha já foi importada anteriormente e a contagem diferir, bloquear!
                if existentes_desta_linha and len(existentes_desta_linha) != tot:
                    is_aprovada = False
                    motivo_bloqueio = (
                        f"Divergência de inventário: linha {lid} já possui {len(existentes_desta_linha)} unidades importadas no banco, "
                        f"mas o arquivo declara {tot}. Alteração de quantidade de lote anterior exige reconciliação explícita."
                    )
                    totais_contagem["unidades_com_conflito"] += tot

                dotes_distribuicao = []
                for _ in range(c_dote):
                    dotes_distribuicao.append("SIM")
                for _ in range(s_dote):
                    dotes_distribuicao.append("NAO")
                for _ in range(ni_dote):
                    dotes_distribuicao.append("NAO_INFORMADO")

                cod_ref = (
                    str(modelo_cat.codigo_scada)
                    if (modelo_cat and modelo_cat.codigo_scada is not None)
                    else (modelo_cat.codigo if modelo_cat else build_missing_model_code(item["modelo_lido"], item["medida_proposta"]))
                )

                for u_idx, dote_val in enumerate(dotes_distribuicao, start=1):
                    chave_u = f"{lid}-U{u_idx:02d}"

                    # Observação individual com ressalva técnica específica da linha
                    if lid == "P2-06":
                        if u_idx == 1:
                            obs_unidade = "Origem Linha P2-06. Ref 01 com dote para diferenciação interna (não lida de gravação física). Carga inicial autorizada."
                        else:
                            obs_unidade = "Origem Linha P2-06. Ref 02 sem dote para diferenciação interna (não lida de gravação física). Carga inicial autorizada."
                    elif lid == "P1-06":
                        obs_unidade = "Origem Linha P1-06. OPTION 90/90-18, 1 un sem dote. Carga inicial autorizada."
                    elif lid == "P1-07":
                        obs_unidade = "Origem Linha P1-07. Carga inicial autorizada (5 un adotadas com ressalva de rasura na folha física preservada)."
                    elif lid == "P1-08":
                        obs_unidade = "Origem Linha P1-08. Carga inicial autorizada (distribuição de dote transcrita com ressalva de leitura preservada)."
                    elif lid == "P1-25":
                        obs_unidade = "Origem Linha P1-25. Carga inicial autorizada (dote não informado na foto original)."
                    elif lid == "P2-23":
                        obs_unidade = "Origem Linha P2-23. Carga inicial autorizada (medida 4.80/4.00-8 adotada a partir da transcrição de 4.80/400-8)."
                    elif lid in ["P2-12", "P2-13", "P2-14"]:
                        obs_unidade = f"Origem Linha {lid}. Carga inicial autorizada (medida normalizada de {item['medida_lida']} para {item['medida_proposta']})."
                    else:
                        obs_orig = item["observacoes"].strip()
                        if obs_orig:
                            obs_unidade = f"Origem Linha {lid}. Carga inicial autorizada, com ressalvas preservadas. {obs_orig}"
                        else:
                            obs_unidade = f"Origem Linha {lid}. Carga inicial autorizada."

                    # Reconciliação com existentes
                    ja_existe = None
                    for ex in existentes_desta_linha:
                        if ex.chave_unidade_origem == chave_u or ex.numero_sequencial == u_idx:
                            ja_existe = ex
                            break

                    if ja_existe:
                        unidades_planejadas.append({
                            "chave_unidade": chave_u,
                            "sequencial": ja_existe.numero_sequencial,
                            "identificador_estavel": ja_existe.identificador_estavel,
                            "possui_dote": ja_existe.possui_dote,
                            "situacao_identificacao": ja_existe.situacao_identificacao,
                            "numero_fisico_confirmado": ja_existe.numero_fisico_confirmado,
                            "status_reconciliacao": "PRESERVADA_EXISTENTE",
                            "obj_existente": ja_existe,
                            "observacao_unidade": ja_existe.observacao_identificacao,
                        })
                        totais_contagem["unidades_existentes_preservadas"] += 1
                    else:
                        seq_proposto = u_idx
                        ident_proposto = f"MF-{cod_ref}-{seq_proposto:03d}"

                        conflito = any(ex.numero_sequencial == seq_proposto or ex.identificador_estavel == ident_proposto for ex in existentes_modelo)
                        if conflito:
                            totais_contagem["unidades_com_conflito"] += 1
                            status_rec = "CONFLITO_SEQUENCIAL"
                        else:
                            totais_contagem["unidades_propostas_novas"] += 1
                            status_rec = "PROPOSTA_NOVA"

                        unidades_planejadas.append({
                            "chave_unidade": chave_u,
                            "sequencial": seq_proposto,
                            "identificador_estavel": ident_proposto,
                            "possui_dote": dote_val,
                            "situacao_identificacao": "PENDENTE",
                            "numero_fisico_confirmado": None,
                            "status_reconciliacao": status_rec,
                            "obj_existente": None,
                            "observacao_unidade": obs_unidade,
                        })

                if any(u["status_reconciliacao"] == "CONFLITO_SEQUENCIAL" for u in unidades_planejadas):
                    is_aprovada = False
                    motivo_bloqueio = f"Conflito com {len(existentes_modelo)} matrizes físicas já cadastradas para este modelo. Bloqueado para evitar colisão."

            if is_aprovada:
                totais_contagem["linhas_aprovadas_escopo"] += 1
                totais_contagem["unidades_aprovadas_escopo"] += tot

            relatorio_linhas.append({
                "item": item,
                "status_cat": status_cat,
                "status_fonte": status_fonte,
                "pendencias_fonte": pendencias_fonte,
                "candidatos": candidatos,
                "modelo_escolhido": candidatos[0] if len(candidatos) == 1 else None,
                "motivo_bloqueio": motivo_bloqueio,
                "aprovado": is_aprovada,
                "unidades_planejadas": unidades_planejadas,
            })

        # 6. Apresentação do Relatório Detalhado
        self.stdout.write(f"\n{'ID':<7} | {'MODELO / MEDIDA PROPOSTA':<26} | {'TOT':<3} | {'FONTE':<10} | {'CATÁLOGO':<12} | {'STATUS / DETALHE'}")
        self.stdout.write("-" * 95)

        for rl in relatorio_linhas:
            it = rl["item"]
            lid = it["id_linha"]
            mod_med = f"{it['modelo_lido']} {it['medida_proposta']}"[:26]
            tot = it["quantidade_total"]
            st_fonte = "OK" if rl["status_fonte"] == "CONFERIDO" else "PENDENTE"
            st_cat = rl["status_cat"]

            if st_cat == "UNICA":
                cat = rl["modelo_escolhido"]
                cat_info = f"ID:{cat.id} SC:{cat.codigo_scada}"
                if rl["aprovado"]:
                    status_badge = self.style.SUCCESS("APROVADA")
                    detalhe = f"{status_badge} - {cat.nome_exibicao}"
                else:
                    status_badge = self.style.WARNING("BLOQUEADA")
                    detalhe = f"{status_badge} - {rl['motivo_bloqueio']}"
            elif st_cat == "AMBIGUA":
                cat_info = f"AMBÍGUO ({len(rl['candidatos'])})"
                detalhe = self.style.NOTICE(f"PENDENTE - Variantes: {[c.nome_exibicao for c in rl['candidatos']]}")
            elif st_cat == "ID_INCOMPATIVEL":
                cat_info = "INCOMPATÍVEL"
                detalhe = self.style.ERROR(f"BLOQUEADA - {rl['motivo_bloqueio']}")
            elif st_cat == "AUSENTE":
                if rl["aprovado"]:
                    cat_info = "NOVO MODELO"
                    status_badge = self.style.SUCCESS("APROVADA")
                    detalhe = f"{status_badge} - Modelo novo a cadastrar no catálogo: PNEU {it['modelo_lido']} {it['medida_proposta']}"
                else:
                    cat_info = "AUSENTE"
                    detalhe = self.style.ERROR("PENDENTE - Não cadastrado no catálogo SCADA")
            else:
                cat_info = "ERRO"
                detalhe = self.style.ERROR(f"BLOQUEADA - {rl['motivo_bloqueio']}")

            self.stdout.write(f"{lid:<7} | {mod_med:<26} | {tot:<3} | {st_fonte:<10} | {cat_info:<12} | {detalhe}")

        # Resumo Estatístico Executivo
        self.stdout.write("\n" + "=" * 80)
        self.stdout.write("  RESUMO ESTATÍSTICO DO INVENTÁRIO (REVISADO)")
        self.stdout.write("=" * 80)
        self.stdout.write(f"Total de Linhas Lidas:               {totais_contagem['linhas_lidas']:>3}  |  Unidades Totais:              {totais_contagem['unidades_lidas']:>3}")
        self.stdout.write(f"Correspondência ÚNICA no Catálogo:   {totais_contagem['linhas_unicas']:>3}  |  Unidades Únicas:              {totais_contagem['unidades_unicas']:>3}")
        self.stdout.write(f"  -> Fonte 100% Conferida (Seguras): {totais_contagem['linhas_unicas_conferidas']:>3}  |  Unidades Seguras:             {totais_contagem['unidades_unicas_conferidas']:>3}")
        self.stdout.write(f"  -> Fonte com Pendência de Leitura: {totais_contagem['linhas_unicas_fonte_pendente']:>3}  |  Unidades com Ressalva:        {totais_contagem['unidades_unicas_fonte_pendente']:>3}")
        self.stdout.write(f"Correspondência AMBÍGUA (Var. S/C):  {totais_contagem['linhas_ambiguas']:>3}  |  Unidades com Múltiplas Opções:  {totais_contagem['unidades_ambiguas']:>3}")
        self.stdout.write(f"Ausentes no Catálogo SCADA:          {totais_contagem['linhas_ausentes']:>3}  |  Unidades Sem Modelo SCADA:      {totais_contagem['unidades_ausentes']:>3}")
        self.stdout.write(f"Linhas Inconsistentes Matemáticas:   {totais_contagem['linhas_inconsistentes']:>3}")
        self.stdout.write("-" * 80)
        self.stdout.write(f"Linhas Aprovadas no Escopo:          {totais_contagem['linhas_aprovadas_escopo']:>3}  |  Unidades no Escopo:             {totais_contagem['unidades_aprovadas_escopo']:>3}")
        self.stdout.write(f"Unidades Novas a Criar:              {totais_contagem['unidades_propostas_novas']:>3}  |  Unidades Existentes Preservadas:{totais_contagem['unidades_existentes_preservadas']:>3}")
        self.stdout.write(f"Unidades com Conflito Cadastral:     {totais_contagem['unidades_com_conflito']:>3}")
        self.stdout.write("=" * 80)

        # Pendências e Alertas Específicos
        self.stdout.write("\n" + self.style.WARNING("[!] PENDENCIAS E ALERTAS TECNICOS DETECTADOS:"))
        alertas_especificos = [
            ("P1-07", "HOPPER 2/75-18", "Quantidade lida como 5 apos rasura na folha. Carga inicial com ressalva preservada."),
            ("Bloco 2/75-18", "READY riscado", "Nome riscado na foto 1 sem quantidade valida. Mantido fora do cadastro."),
            ("P1-08", "WINGS 2/75-18", "Anotacao 'Nao' proxima a linha riscada de READY. Carga inicial com ressalva preservada."),
            ("P1-25", "WINTER 120/100-18", "Ultima coluna em branco. Dote mantido como 'NAO INFORMADO'."),
            ("P2-06", "WINGS 80/100-14", "Caso misto (1 Sim, 1 Nao). U01=com dote, U02=sem dote definidas para distincao interna."),
            ("P2-23", "HOPPER 4.80/400-8", "Medida proposta 4.80/4.00-8 adotada para carga inicial a partir de 4.80/400-8."),
            ("P2-12/13/14", "Medidas com barra", "Medidas 2/75-17 e 2/50-17 normalizadas para 2.75 e 2.50."),
        ]
        for ref, tit, msg in alertas_especificos:
            self.stdout.write(f"  * [{ref}] {tit}: {msg}")

        # 7. Geração da Planilha de Reconciliação
        try:
            self.gerar_planilha_reconciliacao(relatorio_linhas, propostas_novos_modelos, caminho_reconciliacao)
            self.stdout.write("\n" + self.style.SUCCESS(f"[OK] Planilha de Reconciliacao gerada: {caminho_reconciliacao}"))
        except Exception as e:
            self.stdout.write(self.style.ERROR(f"Erro ao gerar planilha de reconciliação: {e}"))

        # 8. Execução ou Simulação
        unidades_criadas_ids = []
        unidades_criadas_count = 0
        unidades_preservadas_count = 0

        if simulacao:
            self.stdout.write("\n" + self.style.SUCCESS("[OK] SIMULACAO CONCLUIDA: NENHUM REGISTRO FOI GRAVADO NO BANCO DE DADOS."))
            self.stdout.write("Para efetivar a carga inicial autorizada no banco de dados, execute:")
            self.stdout.write(f"python manage.py importar_matrizes_fisicas \"{caminho_arquivo}\" --aplicar --responsavel \"{responsavel_audit or 'paulo'}\" --carga-inicial")
            return

        # 9. Aplicação Real (somente com --aplicar)
        linhas_para_gravar = [rl for rl in relatorio_linhas if rl["aprovado"]]
        if not linhas_para_gravar:
            self.stdout.write(self.style.WARNING("\nNenhuma linha aprovada para carga neste escopo. Operação concluída sem gravações."))
            return

        self.stdout.write(f"\nIniciando transação atômica para gravação de {len(linhas_para_gravar)} linha(s)...")

        with transaction.atomic():
            modelos_criados_count = 0
            # 9.1 Garantir cadastro em lote dos modelos ausentes em ProductionMatrixCatalog
            for rl in linhas_para_gravar:
                if (rl["status_cat"] == "AUSENTE" or not rl["modelo_escolhido"]) and rl["aprovado"]:
                    it = rl["item"]
                    cod_int = build_missing_model_code(it["modelo_lido"], it["medida_proposta"])
                    cat_criado, criado = ProductionMatrixCatalog.objects.get_or_create(
                        codigo=cod_int,
                        defaults={
                            "codigo_scada": None,
                            "nome_scada": None,
                            "nome_exibicao": f"PNEU {it['modelo_lido']} {it['medida_proposta']}",
                            "produto": f"PNEU {it['modelo_lido']} {it['medida_proposta']}",
                            "medida_str": it["medida_proposta"],
                            "descricao": "Inventário inicial Matrizaria (Não homologado PCP)",
                            "tempo_producao_segundos": 0,
                            "tempo_vulcanizacao_segundos": 0,
                            "variante_sc": False,
                            "ativo": True,
                        }
                    )
                    if criado:
                        modelos_criados_count += 1
                    rl["modelo_escolhido"] = cat_criado

            # 9.2 Gravação das matrizes físicas individuais
            for rl in linhas_para_gravar:
                cat = rl["modelo_escolhido"]
                lid = rl["item"]["id_linha"]

                for u in rl["unidades_planejadas"]:
                    if u["status_reconciliacao"] == "PRESERVADA_EXISTENTE":
                        unidades_preservadas_count += 1
                        continue

                    nova_matriz = MatrizFisica.objects.create(
                        modelo=cat,
                        identificador_estavel=u["identificador_estavel"],
                        numero_sequencial=u["sequencial"],
                        possui_dote=u["possui_dote"],
                        numero_fisico_confirmado=None,
                        situacao_identificacao="PENDENTE",
                        observacao_identificacao=u["observacao_unidade"],
                        origem_cadastro="INVENTARIO_LOTE",
                        lote_importacao=lote_id,
                        linha_origem=lid,
                        chave_unidade_origem=u["chave_unidade"],
                        ativo=True,
                    )
                    unidades_criadas_ids.append(nova_matriz.id)
                    unidades_criadas_count += 1

            # 9.3 Gravar log de auditoria do lote
            linhas_ids = [rl["item"]["id_linha"] for rl in linhas_para_gravar]
            LoteImportacaoMatrizFisica.objects.create(
                identificacao_lote=lote_id,
                arquivo_nome=file_name,
                arquivo_hash=file_hash,
                responsavel=responsavel_audit,
                simulacao=False,
                quantidade_linhas_lidas=len(linhas_para_gravar),
                quantidade_unidades_criadas=unidades_criadas_count,
                quantidade_unidades_preservadas=unidades_preservadas_count,
                status="SUCESSO",
                linhas_origem_json=json.dumps(linhas_ids),
                ids_criados_json=json.dumps(unidades_criadas_ids),
                relatorio_execucao=(
                    f"Importação de {unidades_criadas_count} matrizes físicas "
                    f"({modelos_criados_count} novos modelos no catálogo) concluída com sucesso por {responsavel_audit}."
                ),
            )

        self.stdout.write(self.style.SUCCESS(
            f"\n[OK] SUCESSO: Lote {lote_id} gravado no banco com sucesso!"
        ))
        self.stdout.write(f"Modelos novos criados no catálogo: {modelos_criados_count}")
        self.stdout.write(f"Matrizes físicas criadas: {unidades_criadas_count}")
        self.stdout.write(f"Matrizes físicas preservadas: {unidades_preservadas_count}")
        self.stdout.write(f"IDs criados: {unidades_criadas_ids}")

    def gerar_planilha_reconciliacao(self, relatorio_linhas, propostas_novos_modelos, output_path):
        """Gera planilha Excel profissional em duas abas para conferência humana detalhada."""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Reconciliacao_Inventario"

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        headers_ws1 = [
            "ID Linha",
            "Modelo Lido",
            "Medida Lida (Foto)",
            "Medida Proposta",
            "Qtd Total",
            "Com Dote",
            "Sem Dote",
            "Dote Não Inf.",
            "Situação da Fonte",
            "Pendências da Fonte",
            "Status Catálogo",
            "Candidatos Catálogo (Banco Local)",
            "ID Sugerido",
            "ID Confirmado Usuário",
            "Aprovação para Carga",
            "Motivo de Bloqueio / Ação Necessária",
            "Observações Originais",
        ]

        ws1.append(headers_ws1)
        for col_idx in range(1, len(headers_ws1) + 1):
            cell = ws1.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for rl in relatorio_linhas:
            it = rl["item"]
            cands = rl["candidatos"]
            cands_str = "; ".join([f"ID={c.id} SCADA={c.codigo_scada} ({c.nome_exibicao})" for c in cands]) if cands else "Nenhum correspondente"
            id_sugerido = cands[0].id if len(cands) == 1 else ""

            aprov_str = "APROVADA" if rl["aprovado"] else ("BLOQUEADA" if rl["motivo_bloqueio"] else "PENDENTE")

            row_vals = [
                it["id_linha"],
                it["modelo_lido"],
                it["medida_lida"],
                it["medida_proposta"],
                it["quantidade_total"],
                it["com_dote"],
                it["sem_dote"],
                it["dote_nao_informado"],
                rl["status_fonte"],
                "; ".join(rl["pendencias_fonte"]) if rl["pendencias_fonte"] else "Fonte 100% conferida",
                rl["status_cat"],
                cands_str,
                id_sugerido,
                "",  # Espaço para o usuário preencher na conferência
                aprov_str,
                rl["motivo_bloqueio"],
                it["observacoes"],
            ]
            ws1.append(row_vals)

        # Ajuste de largura das colunas
        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Aba 2: Proposta para Novos Modelos
        ws2 = wb.create_sheet(title="Proposta_Novos_Modelos")
        headers_ws2 = [
            "Linha Origem",
            "Modelo",
            "Medida Proposta",
            "Medida Lida (Foto)",
            "Código Interno Proposto (codigo)",
            "Nome de Exibição Proposto",
            "Código SCADA",
            "Variante S/C",
            "Tempo Produção Padrão (s)",
            "Tempo Vulcanização Padrão (s)",
            "Pendências da Fonte",
            "Observações",
        ]

        ws2.append(headers_ws2)
        header_fill_2 = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        for col_idx in range(1, len(headers_ws2) + 1):
            cell = ws2.cell(1, col_idx)
            cell.fill = header_fill_2
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for p in propostas_novos_modelos:
            row_vals_2 = [
                p["linha_origem"],
                p["modelo"],
                p["medida_proposta"],
                p["medida_lida"],
                p["codigo_interno_proposto"],
                p["nome_exibicao_proposto"],
                p["codigo_scada"],
                p["variante_sc"],
                p["tempo_producao_padrao"],
                p["tempo_vulcanizacao_padrao"],
                p["pendencias"],
                p["observacoes"],
            ]
            ws2.append(row_vals_2)

        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
