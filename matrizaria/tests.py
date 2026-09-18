from datetime import datetime, date, time, timedelta
from unittest.mock import patch
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.exceptions import ValidationError, PermissionDenied
from django.utils import timezone

from django.db.models import ProtectedError
from maintenance.models import Sector, Machine
from production.models import ProductionMatrixCatalog
from matrizaria.models import (
    TipoServicoMatrizaria,
    MatrizFisica,
    SolicitacaoServicoMatrizaria,
    CicloExecucaoMatrizaria,
    HistoricoTransicaoServicoMatrizaria,
)
from matrizaria.services import MatrizariaService, AdminCascadeDeletionService
from matrizaria.views import sanitize_excel_cell

User = get_user_model()


class MatrizariaBaseTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        # 1. Setor e Máquina
        cls.sector_vulc = Sector.objects.create(nome="VULCANIZAÇÃO")
        cls.prensa_01 = Machine.objects.create(
            nome="PRENSA BOM 01",
            setor=cls.sector_vulc,
            criticidade="ALTA",
        )
        cls.prensa_02 = Machine.objects.create(
            nome="PRENSA BOM 02",
            setor=cls.sector_vulc,
            criticidade="MEDIA",
        )

        # 2. Catálogo Canônico de Matrizes
        cls.cat_hopper, _ = ProductionMatrixCatalog.objects.get_or_create(
            codigo_scada=1,
            defaults={
                "nome_scada": "HOPPER 90/90-18",
                "nome_exibicao": "HOPPER 90/90-18",
                "ativo": True,
            },
        )

        # 3. Exemplar Físico de Matriz
        cls.matriz_fisica_01 = MatrizFisica.objects.create(
            modelo=cls.cat_hopper,
            identificador_estavel="MF-HOPPER-001",
            numero_sequencial=1,
            ativo=True,
        )

        # 4. Tipos de Serviço
        cls.tipo_troca = TipoServicoMatrizaria.objects.create(
            nome="Troca de Matriz Programada",
            descricao="Troca preventiva do ferramental",
            exige_matriz_fisica=True,
            ativo=True,
            ordem_exibicao=1,
        )
        cls.tipo_ajuste = TipoServicoMatrizaria.objects.create(
            nome="Ajuste Mecânico de Prensa",
            descricao="Alinhamento sem troca de matriz",
            exige_matriz_fisica=False,
            ativo=True,
            ordem_exibicao=2,
        )

        # 5. Grupos de Usuários
        cls.grp_matrizaria, _ = Group.objects.get_or_create(name="Matrizaria")
        cls.grp_vulc, _ = Group.objects.get_or_create(name="Operadores Vulcanização")
        cls.grp_lider, _ = Group.objects.get_or_create(name="Liderança de Produção")
        cls.grp_maint, _ = Group.objects.get_or_create(name="Tecnicos")
        cls.grp_tv_matriz, _ = Group.objects.get_or_create(name="Visualizador Matrizaria")

        # 6. Usuários
        cls.solicitante_user = User.objects.create_user(
            username="op_vulc", password="password123", first_name="Operador", last_name="Vulcanizador"
        )
        cls.solicitante_user.groups.add(cls.grp_vulc)

        cls.tecnico_matrizaria_1 = User.objects.create_user(
            username="tec_matriz_1", password="password123", first_name="Carlos", last_name="Matrizaria"
        )
        cls.tecnico_matrizaria_1.groups.add(cls.grp_matrizaria)

        cls.tecnico_matrizaria_2 = User.objects.create_user(
            username="tec_matriz_2", password="password123", first_name="Roberto", last_name="Ferramenteiro"
        )
        cls.tecnico_matrizaria_2.groups.add(cls.grp_matrizaria)

        cls.lider_user = User.objects.create_user(
            username="lider_prod", password="password123", first_name="Marcos", last_name="Líder"
        )
        cls.lider_user.groups.add(cls.grp_lider)
        perm_view = Permission.objects.get(
            codename="view_solicitacaoservicomatrizaria",
            content_type__app_label="matrizaria",
        )
        cls.lider_user.user_permissions.add(perm_view)

        cls.tecnico_maint_user = User.objects.create_user(
            username="tec_maint", password="password123", first_name="Pedro", last_name="Mecânico"
        )
        cls.tecnico_maint_user.groups.add(cls.grp_maint)

        cls.tv_user = User.objects.create_user(
            username="tv_matrizaria", password="password123"
        )
        cls.tv_user.groups.add(cls.grp_tv_matriz)


class MatrizariaModelAndSnapshotTestCase(MatrizariaBaseTestCase):
    def test_matriz_fisica_string_and_formatting(self):
        expected_name = f"{self.cat_hopper.nome_exibicao} #001"
        self.assertEqual(self.matriz_fisica_01.numero_sequencial_str, "001")
        self.assertEqual(self.matriz_fisica_01.nome_exibicao, expected_name)
        self.assertEqual(str(self.matriz_fisica_01), expected_name)

    def test_matriz_fisica_unique_constraints(self):
        # Mesma matriz e número sequencial duplicado deve falhar
        with self.assertRaises(Exception):
            MatrizFisica.objects.create(
                modelo=self.cat_hopper,
                identificador_estavel="MF-HOPPER-DUPLICADO",
                numero_sequencial=1,
            )

    def test_solicitacao_snapshots_preserve_history_when_catalogs_change(self):
        expected_matriz_name = f"{self.cat_hopper.nome_exibicao} #001"
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_troca,
            descricao_solicitacao="Trocar matriz com desgaste",
            solicitado_por=self.solicitante_user,
            matriz_fisica=self.matriz_fisica_01,
        )

        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")
        self.assertEqual(sol.tipo_servico_nome_snapshot, "Troca de Matriz Programada")
        self.assertEqual(sol.solicitado_por_nome, "Operador Vulcanizador")
        self.assertEqual(sol.matriz_identificador_snapshot, expected_matriz_name)
        self.assertTrue(sol.exige_matriz_fisica_snapshot)

        # Alteração posterior no cadastro da máquina ou tipo não deve alterar o snapshot
        self.prensa_01.nome = "PRENSA BOM 01 RENOMEADA"
        self.prensa_01.save()

        self.tipo_troca.nome = "Tipo de Serviço Renomeado"
        self.tipo_troca.exige_matriz_fisica = False
        self.tipo_troca.save()

        sol.refresh_from_db()
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")
        self.assertEqual(sol.tipo_servico_nome_snapshot, "Troca de Matriz Programada")
        self.assertTrue(sol.exige_matriz_fisica_snapshot)


class MatrizariaWorkflowAndCyclesTestCase(MatrizariaBaseTestCase):
    def test_complete_lifecycle_including_rework_and_cycles(self):
        # 1. Abertura
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_troca,
            descricao_solicitacao="Troca de molde",
            solicitado_por=self.solicitante_user,
            prioridade="NORMAL",
        )
        self.assertEqual(sol.status, "SOLICITADO")
        self.assertEqual(sol.versao, 1)
        self.assertEqual(sol.ciclos_execucao.count(), 0)

        # 2. Início do Atendimento (Abre Ciclo 1)
        sol = MatrizariaService.iniciar_atendimento(
            solicitacao_id=sol.id,
            usuario=self.tecnico_matrizaria_1,
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol.status, "EM_EXECUCAO")
        self.assertEqual(sol.responsavel_atribuido, self.tecnico_matrizaria_1)
        self.assertEqual(sol.ciclos_execucao.count(), 1)
        c1 = sol.ciclos_execucao.first()
        self.assertEqual(c1.numero_ciclo, 1)
        self.assertEqual(c1.usuario_inicio, self.tecnico_matrizaria_1)
        self.assertEqual(c1.forma_encerramento, "EM_ANDAMENTO")

        # 3. Finalização sem informar matriz quando exigida deve falhar
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.finalizar_execucao(
                solicitacao_id=sol.id,
                usuario=self.tecnico_matrizaria_1,
                descricao_servico_realizado="Instalada matriz",
                matriz_fisica=None,
                versao_esperada=sol.versao,
            )
        self.assertIn("exige a identificação da Matriz Física", str(ctx.exception))

        # 4. Finalização válida (Encerra Ciclo 1 com FINALIZADO_TECNICO)
        sol = MatrizariaService.finalizar_execucao(
            solicitacao_id=sol.id,
            usuario=self.tecnico_matrizaria_1,
            descricao_servico_realizado="Matriz instalada com alinhamento e vapor conectado",
            matriz_fisica=self.matriz_fisica_01,
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol.status, "AGUARDANDO_CONFERENCIA")
        c1.refresh_from_db()
        self.assertEqual(c1.forma_encerramento, "FINALIZADO_TECNICO")
        self.assertEqual(c1.usuario_fim, self.tecnico_matrizaria_1)
        self.assertIsNotNone(c1.data_fim)

        # 5. Conferência: Reprovação para Retrabalho por Líder
        sol = MatrizariaService.conferir_solicitacao(
            solicitacao_id=sol.id,
            usuario_conferente=self.lider_user,
            aprovado=False,
            versao_esperada=sol.versao,
            observacao="Vazamento de vapor detectado na conexão traseira",
        )
        self.assertEqual(sol.status, "AGUARDANDO_RETRABALHO")
        self.assertEqual(sol.quantidade_retrabalhos, 1)

        # 6. Início do Retrabalho (Abre Ciclo 2, preservando Ciclo 1)
        sol = MatrizariaService.iniciar_atendimento(
            solicitacao_id=sol.id,
            usuario=self.tecnico_matrizaria_2,
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol.status, "EM_EXECUCAO")
        self.assertEqual(sol.ciclos_execucao.count(), 2)

        ciclos = list(sol.ciclos_execucao.order_by("numero_ciclo"))
        self.assertEqual(ciclos[0].numero_ciclo, 1)
        self.assertEqual(ciclos[0].forma_encerramento, "FINALIZADO_TECNICO")
        self.assertEqual(ciclos[1].numero_ciclo, 2)
        self.assertEqual(ciclos[1].usuario_inicio, self.tecnico_matrizaria_2)
        self.assertEqual(ciclos[1].forma_encerramento, "EM_ANDAMENTO")

        # 7. Finalização do Retrabalho
        sol = MatrizariaService.finalizar_execucao(
            solicitacao_id=sol.id,
            usuario=self.tecnico_matrizaria_2,
            descricao_servico_realizado="Substituída vedação de vapor e reapertada flange",
            matriz_fisica=self.matriz_fisica_01,
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol.status, "AGUARDANDO_CONFERENCIA")

        # 8. Aprovação Final por Líder
        sol = MatrizariaService.conferir_solicitacao(
            solicitacao_id=sol.id,
            usuario_conferente=self.lider_user,
            aprovado=True,
            versao_esperada=sol.versao,
            observacao="Aprovado sem vazamentos, prensa liberada para produção",
        )
        self.assertEqual(sol.status, "CONCLUIDO")
        self.assertEqual(sol.conferido_por, self.lider_user)

        # Transição adicional após concluído deve ser bloqueada
        with self.assertRaises(ValidationError):
            MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)

    def test_cancellation_during_execution_interrupts_active_cycle(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_02,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Ajuste na prensa 02",
            solicitado_por=self.solicitante_user,
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)

        # Cancelamento durante execução técnica
        sol = MatrizariaService.cancelar_solicitacao(
            solicitacao_id=sol.id,
            usuario=self.lider_user,
            versao_esperada=sol.versao,
            motivo_cancelamento="Ordem de produção suspensa pela diretoria",
        )
        self.assertEqual(sol.status, "CANCELADO")
        c = sol.ciclos_execucao.first()
        self.assertEqual(c.forma_encerramento, "INTERROMPIDO_CANCELAMENTO")
        self.assertIn("Interrompido por cancelamento", c.descricao_servico_executado)
        self.assertIsNotNone(c.data_fim)


class MatrizariaSecurityAndAntiSelfInspectionTestCase(MatrizariaBaseTestCase):
    def test_anti_self_inspection_blocks_execution_participants(self):
        # Técnico que executou não pode aprovar
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Ajuste teste",
            solicitado_por=self.solicitante_user,
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)
        sol = MatrizariaService.finalizar_execucao(
            solicitacao_id=sol.id,
            usuario=self.tecnico_matrizaria_1,
            descricao_servico_executado="Ajuste técnico concluído com sucesso",
            matriz_fisica=None,
            versao_esperada=sol.versao,
        )

        # Técnico tenta aprovar o próprio serviço
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.conferir_solicitacao(
                solicitacao_id=sol.id,
                usuario_conferente=self.tecnico_matrizaria_1,
                aprovado=True,
                versao_esperada=sol.versao,
            )
        self.assertIn("Regra de Segregação de Funções", str(ctx.exception))

    def test_transfer_of_responsibility_audited(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Troca de turno",
            solicitado_por=self.solicitante_user,
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)

        # Transfere para técnico 2
        sol = MatrizariaService.transferir_responsabilidade(
            solicitacao_id=sol.id,
            usuario_origem=self.tecnico_matrizaria_1,
            novo_responsavel=self.tecnico_matrizaria_2,
            motivo="Fim de turno técnico 1",
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol.responsavel_atribuido, self.tecnico_matrizaria_2)
        # O autor do início original continua sendo o técnico 1
        self.assertEqual(sol.iniciado_por, self.tecnico_matrizaria_1)

        # Histórico gravado
        h = sol.historico_transicoes.filter(tipo_evento="TRANSFERENCIA").first()
        self.assertIsNotNone(h)
        self.assertIn("Transferência de responsabilidade", h.observacao)


class MatrizariaConcurrencyTestCase(MatrizariaBaseTestCase):
    def test_stale_version_triggers_concurrency_error(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Teste concorrencia",
            solicitado_por=self.solicitante_user,
        )
        versao_antiga = sol.versao

        # Primeiro usuário inicia com sucesso
        MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, versao_antiga)

        # Segundo usuário tenta agir com a versão antiga (stale)
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_2, versao_antiga)
        self.assertIn("Conflito de concorrência", str(ctx.exception))


class MatrizariaReportingAndTemporalTestCase(MatrizariaBaseTestCase):
    def test_temporal_scenario_execution_crossing_dates(self):
        """
        Cenário Temporal:
        - Abertura no dia 10
        - Execução no dia 11
        - Conferência no dia 12
        - Relatório de execução (COM_EXECUCAO) do dia 11 deve encontrar o chamado.
        - Relatório de execução do dia 10 NÃO deve encontrar (pois dia 10 teve apenas espera).
        """
        d10 = datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc)
        d11_ini = datetime(2026, 9, 11, 8, 0, tzinfo=timezone.utc)
        d11_fim = datetime(2026, 9, 11, 12, 0, tzinfo=timezone.utc)
        d12 = datetime(2026, 9, 12, 9, 0, tzinfo=timezone.utc)

        sol = SolicitacaoServicoMatrizaria.objects.create(
            prensa=self.prensa_01,
            prensa_nome_snapshot=self.prensa_01.nome,
            tipo_servico=self.tipo_ajuste,
            tipo_servico_nome_snapshot=self.tipo_ajuste.nome,
            descricao_solicitacao="Teste temporal",
            solicitado_por=self.solicitante_user,
            solicitado_por_nome="Solicitante",
            data_solicitacao=d10,
            status="CONCLUIDO",
            conferido_por=self.lider_user,
            data_conferencia=d12,
            data_inicio_execucao=d11_ini,
            data_fim_execucao=d11_fim,
        )

        CicloExecucaoMatrizaria.objects.create(
            solicitacao=sol,
            numero_ciclo=1,
            usuario_inicio=self.tecnico_matrizaria_1,
            usuario_inicio_nome="Tecnico 1",
            data_inicio=d11_ini,
            usuario_fim=self.tecnico_matrizaria_1,
            usuario_fim_nome="Tecnico 1",
            data_fim=d11_fim,
            forma_encerramento="FINALIZADO_TECNICO",
            descricao_servico_executado="Servico executado dia 11",
        )

        HistoricoTransicaoServicoMatrizaria.objects.create(
            solicitacao=sol,
            status_anterior="SOLICITADO",
            status_novo="EM_EXECUCAO",
            tipo_evento="INICIO",
            data_evento=d11_ini,
        )
        HistoricoTransicaoServicoMatrizaria.objects.create(
            solicitacao=sol,
            status_anterior="EM_EXECUCAO",
            status_novo="CONCLUIDO",
            tipo_evento="CONFERENCIA_APROVADA",
            data_evento=d12,
        )

        # Consulta COM_EXECUCAO para o dia 11 (deve encontrar)
        res_dia11 = MatrizariaService.consultar_relatorio(
            criterio_temporal="COM_EXECUCAO",
            data_inicio=date(2026, 9, 11),
            data_fim=date(2026, 9, 11),
        )
        ids_dia11 = [r["id"] for r in res_dia11]
        self.assertIn(sol.id, ids_dia11)

        # Consulta COM_EXECUCAO para o dia 10 (apenas espera, não teve execução)
        res_dia10 = MatrizariaService.consultar_relatorio(
            criterio_temporal="COM_EXECUCAO",
            data_inicio=date(2026, 9, 10),
            data_fim=date(2026, 9, 10),
        )
        ids_dia10 = [r["id"] for r in res_dia10]
        self.assertNotIn(sol.id, ids_dia10)

        # Reconstituição do status histórico no fim do dia 11:
        item_dia11 = [r for r in res_dia11 if r["id"] == sol.id][0]
        self.assertEqual(item_dia11["status_atual"], "CONCLUIDO")
        self.assertEqual(item_dia11["status_fim_periodo"], "EM_EXECUCAO")


class MatrizariaExcelAndExportTestCase(MatrizariaBaseTestCase):
    def test_excel_export_creates_two_sheets_and_sanitizes_formulas(self):
        client = Client()
        client.force_login(self.lider_user)

        # Cria solicitação com texto iniciando em sinal de fórmula
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="=1+1 perigo de formula",
            solicitado_por=self.solicitante_user,
        )

        res = client.get(reverse("matrizaria:exportar_excel"))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # Reabre a planilha para validação rigorosa
        import io
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Serviços", wb.sheetnames)
        self.assertIn("Histórico", wb.sheetnames)

        # Sanitização de fórmula
        self.assertEqual(sanitize_excel_cell("=1+1"), "'=1+1")
        self.assertEqual(sanitize_excel_cell("+5511"), "'+5511")
        self.assertEqual(sanitize_excel_cell("-100"), "'-100")
        self.assertEqual(sanitize_excel_cell("@usuario"), "'@usuario")
        self.assertEqual(sanitize_excel_cell("Texto Normal"), "Texto Normal")


class MatrizariaTVAndIsolationTestCase(MatrizariaBaseTestCase):
    def test_tv_view_and_scada_resilience(self):
        client = Client()
        client.force_login(self.tv_user)

        # Rota de TV carrega 200 OK
        res = client.get(reverse("matrizaria:tv"))
        self.assertEqual(res.status_code, 200)

        # API JSON da TV
        res_api = client.get(reverse("matrizaria:api_tv_data"))
        self.assertEqual(res_api.status_code, 200)
        data = res_api.json()
        self.assertIn("servicos", data)
        self.assertIn("bladder_alerts", data)

        # Falha de telemetria SCADA isolada
        with patch("production.services.BladderTrackingService.get_active_bladders_context", side_effect=Exception("Timeout no SCADA")):
            ctx = MatrizariaService.get_tv_dashboard_context()
            self.assertEqual(ctx["scada_status"], "OFFLINE")
            self.assertIn("Timeout no SCADA", ctx["scada_mensagem"])
            # Serviços continuam funcionando normalmente!
            self.assertIn("servicos", ctx)


class MatrizariaRouterAndMultiContextTestCase(MatrizariaBaseTestCase):
    def test_routing_and_portal_redirects(self):
        client = Client()

        # 1. Usuário exclusivo da TV Matrizaria -> /matrizaria/tv/
        client.force_login(self.tv_user)
        res_tv = client.get(reverse("home_redirect"))
        self.assertRedirects(res_tv, reverse("matrizaria:tv"))

        # 2. Usuário com 1 módulo (solicitante_user: apenas vulcanização -> matrizaria)
        client.force_login(self.solicitante_user)
        res_single = client.get(reverse("home_redirect"))
        self.assertRedirects(res_single, reverse("matrizaria:kanban"))

        # 3. Usuário com múltiplos módulos (lider_user -> tem produção e matrizaria)
        client.force_login(self.lider_user)
        res_multi = client.get(reverse("home_redirect"))
        self.assertRedirects(res_multi, reverse("portal_select"))


class MatrizariaFixesTestCase(MatrizariaBaseTestCase):
    """
    Testes específicos para validar as três correções solicitadas:
    1. Exclusão de CHECK-LIST e simplificação dos rótulos de prensas (sem criticidade entre parênteses).
    2. Rejeição de CHECK-LIST via formulário, serviço e modelo.
    3. Assinatura do desenvolvedor Paulo Sérgio nos templates compartilhados.
    """

    def test_prensa_label_and_checklist_exclusion(self):
        from matrizaria.forms import get_prensas_queryset, format_prensa_label, is_checklist_machine, SolicitacaoServicoForm

        # Cria registros de teste de CHECK-LIST
        chk_1 = Machine.objects.create(nome="CHECK-LIST", setor=self.sector_vulc, criticidade="BAIXA")
        chk_2 = Machine.objects.create(nome="check list", setor=self.sector_vulc, criticidade="BAIXA")
        chk_3 = Machine.objects.create(nome="Checklist", setor=self.sector_vulc, criticidade="BAIXA")

        # Verifica helper
        self.assertTrue(is_checklist_machine(chk_1))
        self.assertTrue(is_checklist_machine(chk_2))
        self.assertTrue(is_checklist_machine(chk_3))
        self.assertFalse(is_checklist_machine(self.prensa_01))

        # QuerySet de prensas não deve conter nenhum CHECK-LIST
        prensas_qs = get_prensas_queryset()
        self.assertNotIn(chk_1, prensas_qs)
        self.assertNotIn(chk_2, prensas_qs)
        self.assertNotIn(chk_3, prensas_qs)
        self.assertIn(self.prensa_01, prensas_qs)

        # Rótulo simplificado não deve conter parênteses residuais
        label_01 = format_prensa_label(self.prensa_01)
        self.assertEqual(label_01, "PRENSA BOM 01")
        self.assertNotIn("(", label_01)
        self.assertNotIn(")", label_01)

        # Máquina com hífen e nome composto
        prensa_robot = Machine.objects.create(nome="PRENSA BOM - ROBOT", setor=self.sector_vulc, criticidade="ALTA")
        label_robot = format_prensa_label(prensa_robot)
        self.assertEqual(label_robot, "PRENSA BOM - ROBOT")

        # Rejeição via formulário de nova solicitação
        form_data = {
            "prensa": chk_1.id,
            "tipo_servico": self.tipo_troca.id,
            "descricao_solicitacao": "Teste de tentativa de abertura com CHECK-LIST",
            "prioridade": "NORMAL",
        }
        form = SolicitacaoServicoForm(data=form_data)
        self.assertFalse(form.is_valid())
        self.assertIn("prensa", form.errors)

        # Rejeição via serviço
        with self.assertRaises(ValidationError):
            MatrizariaService.criar_solicitacao(
                prensa=chk_1,
                tipo_servico=self.tipo_troca,
                descricao_solicitacao="Tentativa de serviço",
                solicitado_por=self.solicitante_user,
            )

        # Rejeição via Model.clean()
        sol_invalid = SolicitacaoServicoMatrizaria(
            prensa=chk_1,
            tipo_servico=self.tipo_troca,
            descricao_solicitacao="Tentativa direta no model",
            solicitado_por=self.solicitante_user,
        )
        with self.assertRaises(ValidationError):
            sol_invalid.clean()

    def test_developer_footer_presence_across_templates(self):
        from django.template.loader import render_to_string

        templates_to_test = [
            ("maintenance/login.html", {}),
            ("maintenance/portal_select.html", {"accessible_modules": ["matrizaria"], "can_access_matrizaria": True}),
            ("maintenance/base.html", {}),
            ("production/base_production.html", {}),
            ("matrizaria/base_matrizaria.html", {}),
            ("matrizaria/kanban.html", {"fila_solicitados": [], "fila_execucao": [], "fila_conferencia": [], "fila_concluidos": []}),
            ("matrizaria/form_solicitacao.html", {}),
            ("matrizaria/relatorios.html", {"page_obj": []}),
            ("matrizaria/tv.html", MatrizariaService.get_tv_dashboard_context()),
            ("maintenance/tv_dashboard.html", {"technicians": []}),
            ("admin/base_site.html", {}),
        ]

        for template_name, ctx in templates_to_test:
            ctx["user"] = self.lider_user
            rendered = render_to_string(template_name, ctx)
            self.assertIn("Desenvolvido por", rendered, f"Assinatura ausente em {template_name}")
            self.assertIn("Paulo Sérgio", rendered, f"Nome do desenvolvedor ausente em {template_name}")
            # Garante que não foi duplicado
            self.assertEqual(rendered.count("Paulo Sérgio"), 1, f"Assinatura duplicada em {template_name}")


class MatrizariaCincoAjustesTestCase(MatrizariaBaseTestCase):
    """
    Testes automatizados cobrindo os 5 novos ajustes implementados:
    1. TV: renderização completa do HTML das colunas no polling, inclusão de retrabalhos em solicitados, fuso horário local.
    2. Cancelados: critério CANCELADAS filtrando por data de cancelamento, compatibilidade com histórico/execução, exportação Excel.
    3. Sessão e TV: middleware de inatividade humana (5 min), isenção para conta de TV exclusiva, endpoints keep-alive e status.
    4. Edição de Solicitação: apenas quando SOLICITADO e sem ciclos iniciados, verificação de concorrência, auditoria de diff.
    5. Transferência: listagem restrita a matrizeiros habilitados (grupo Matrizaria ativo, excluindo atual e contas TV).
    """

    def test_item1_tv_api_renders_column_html_and_local_timezone(self):
        # 1. Cria duas solicitações
        sol_1 = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_troca,
            descricao_solicitacao="Troca de matriz urgente para BOM 01",
            solicitado_por=self.solicitante_user,
            prioridade="URGENTE",
        )
        sol_2 = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_02,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Ajuste operacional para BOM 02",
            solicitado_por=self.solicitante_user,
            prioridade="NORMAL",
        )

        client = Client()
        client.force_login(self.tv_user)

        # 2. Chama endpoint de API que alimenta a TV
        res = client.get(reverse("matrizaria:api_tv_data"))
        self.assertEqual(res.status_code, 200)
        data = res.json()

        # Verifica contadores
        self.assertEqual(data["servicos"]["total_solicitados"], 2)

        # Verifica que o HTML das colunas é gerado e retornado
        self.assertIn("html_solicitados", data)
        self.assertIn("html_execucao", data)
        self.assertIn("html_conferencia", data)
        self.assertIn("html_bladders", data)

        # Ambos os chamados devem estar presentes no HTML da coluna solicitados
        self.assertIn("Troca de matriz urgente para BOM 01", data["html_solicitados"])
        self.assertIn("Ajuste operacional para BOM 02", data["html_solicitados"])
        self.assertIn("PRENSA BOM 01", data["html_solicitados"])
        self.assertIn("PRENSA BOM 02", data["html_solicitados"])

        # 3. Verifica consistência de timezone no contexto da TV
        ctx = MatrizariaService.get_tv_dashboard_context()
        now_local = timezone.localtime(timezone.now())
        self.assertEqual(ctx["hora_atual"], now_local.strftime("%H:%M"))

    def test_item2_chamados_cancelados_relatorio_e_excel(self):
        # 1. Chamado cancelado ANTES de qualquer início
        sol_pre = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Cancelado antes do inicio",
            solicitado_por=self.solicitante_user,
        )
        sol_pre = MatrizariaService.cancelar_solicitacao(
            solicitacao_id=sol_pre.id,
            usuario=self.lider_user,
            versao_esperada=sol_pre.versao,
            motivo_cancelamento="Planejamento cancelou ordem da prensa 01",
        )

        # 2. Chamado cancelado DEPOIS de iniciado
        sol_exec = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_02,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Cancelado durante execucao",
            solicitado_por=self.solicitante_user,
        )
        sol_exec = MatrizariaService.iniciar_atendimento(sol_exec.id, self.tecnico_matrizaria_1, sol_exec.versao)
        sol_exec = MatrizariaService.cancelar_solicitacao(
            solicitacao_id=sol_exec.id,
            usuario=self.lider_user,
            versao_esperada=sol_exec.versao,
            motivo_cancelamento="Prensa apresentou falha eletrica externa",
        )

        # 3. Consulta de relatório por 'CANCELADAS'
        hoje = timezone.localdate()
        rel = MatrizariaService.consultar_relatorio(
            criterio_temporal="CANCELADAS",
            data_inicio=hoje,
            data_fim=hoje,
        )
        ids_rel = [r["id"] for r in rel]
        self.assertIn(sol_pre.id, ids_rel)
        self.assertIn(sol_exec.id, ids_rel)

        # Verifica dados de cancelamento retornados no relatório
        item_pre = [r for r in rel if r["id"] == sol_pre.id][0]
        self.assertEqual(item_pre["cancelado_por_nome"], "Marcos Líder")
        self.assertEqual(item_pre["motivo_cancelamento"], "Planejamento cancelou ordem da prensa 01")
        self.assertIsNotNone(item_pre["data_cancelamento"])

        item_exec = [r for r in rel if r["id"] == sol_exec.id][0]
        self.assertEqual(item_exec["cancelado_por_nome"], "Marcos Líder")
        self.assertEqual(item_exec["motivo_cancelamento"], "Prensa apresentou falha eletrica externa")
        # Ciclo interrompido sem inventar término fictício de execução com sucesso
        ciclo_cancelado = sol_exec.ciclos_execucao.first()
        self.assertEqual(ciclo_cancelado.forma_encerramento, "INTERROMPIDO_CANCELAMENTO")

        # 4. Exportação Excel para chamados cancelados
        client = Client()
        client.force_login(self.lider_user)
        res_excel = client.get(reverse("matrizaria:exportar_excel") + "?criterio_temporal=CANCELADAS")
        self.assertEqual(res_excel.status_code, 200)
        self.assertEqual(
            res_excel["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_item3_logout_inatividade_middleware_e_perfil_tv(self):
        from maintenance.middleware import is_dedicated_tv_account

        # 1. Conta TV exclusiva é reconhecida
        self.assertTrue(is_dedicated_tv_account(self.tv_user))

        # 2. Usuário com grupo TV mas também grupo operacional NÃO é isento
        self.tv_user.groups.add(self.grp_matrizaria)
        self.assertFalse(is_dedicated_tv_account(self.tv_user))
        self.tv_user.groups.remove(self.grp_matrizaria)

        # 3. Superusuário ou Staff NÃO é conta TV exclusiva
        admin_tv = User.objects.create_superuser(username="admin_tv", password="password123", email="")
        admin_tv.groups.add(self.grp_tv_matriz)
        self.assertFalse(is_dedicated_tv_account(admin_tv))

        # 4. Endpoints de sessão
        client = Client()
        client.force_login(self.solicitante_user)

        # Status para usuário humano
        res_status = client.get(reverse("api_session_status"))
        self.assertEqual(res_status.status_code, 200)
        data_status = res_status.json()
        self.assertTrue(data_status["is_authenticated"])
        self.assertFalse(data_status["is_tv"])
        self.assertIn("remaining_seconds", data_status)
        self.assertEqual(data_status["timeout_seconds"], 300)

        # Keep-alive via POST
        res_alive = client.post(reverse("api_session_keep_alive"))
        self.assertEqual(res_alive.status_code, 200)
        self.assertEqual(res_alive.json()["status"], "ok")

        # Status para conta TV dedicada
        client.force_login(self.tv_user)
        res_tv_status = client.get(reverse("api_session_status"))
        self.assertEqual(res_tv_status.status_code, 200)
        self.assertTrue(res_tv_status.json()["is_tv"])

    def test_item4_editar_solicitacao_antes_do_primeiro_inicio(self):
        # 1. Cria solicitação pendente
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Descricao original antes de editar",
            solicitado_por=self.solicitante_user,
            prioridade="BAIXA",
        )
        self.assertEqual(sol.status, "SOLICITADO")
        self.assertEqual(sol.versao, 1)

        # 2. Solicitante edita a solicitação com sucesso
        sol_editada = MatrizariaService.editar_solicitacao(
            solicitacao_id=sol.id,
            usuario=self.solicitante_user,
            prensa=self.prensa_02,
            tipo_servico=self.tipo_troca,
            matriz_fisica=self.matriz_fisica_01,
            prioridade="ALTA",
            descricao_solicitacao="Descricao corrigida pelo solicitante",
            motivo_edicao="Prensa e matriz corrigidas conforme necessidade real",
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol_editada.prensa, self.prensa_02)
        self.assertEqual(sol_editada.prensa.nome, "PRENSA BOM 02")
        self.assertEqual(sol_editada.prensa_nome_snapshot, "PRENSA BOM 01")  # Preserva evidência da abertura
        self.assertEqual(sol_editada.tipo_servico, self.tipo_troca)
        self.assertEqual(sol_editada.matriz_fisica, self.matriz_fisica_01)
        self.assertEqual(sol_editada.prioridade, "ALTA")
        self.assertEqual(sol_editada.descricao_solicitacao, "Descricao corrigida pelo solicitante")
        self.assertEqual(sol_editada.versao, 2)

        # Verifica histórico de auditoria
        h = sol_editada.historico_transicoes.filter(tipo_evento="ALTERACAO_DADO").first()
        self.assertIsNotNone(h)
        self.assertIn("Prensa e matriz corrigidas", h.observacao)
        self.assertIn("PRENSA BOM 01", h.dados_modificados)
        self.assertIn("PRENSA BOM 02", h.dados_modificados)

        # 3. Usuário sem permissão (técnico de manutenção comum) não pode editar
        with self.assertRaises(PermissionDenied):
            MatrizariaService.editar_solicitacao(
                solicitacao_id=sol_editada.id,
                usuario=self.tecnico_maint_user,
                prensa=self.prensa_01,
                tipo_servico=self.tipo_ajuste,
                matriz_fisica=None,
                prioridade="NORMAL",
                descricao_solicitacao="Tentativa indevida",
                motivo_edicao="Sem permissao",
                versao_esperada=sol_editada.versao,
            )

        # 4. Versão desatualizada (conflito de concorrência)
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.editar_solicitacao(
                solicitacao_id=sol_editada.id,
                usuario=self.solicitante_user,
                prensa=self.prensa_01,
                tipo_servico=self.tipo_ajuste,
                matriz_fisica=None,
                prioridade="NORMAL",
                descricao_solicitacao="Tentativa de edicao desatualizada",
                motivo_edicao="Versao antiga",
                versao_esperada=1,  # versão agora é 2
            )
        self.assertIn("Conflito de concorrência", str(ctx.exception))

        # 5. Se o atendimento já foi iniciado, a edição é bloqueada
        sol_iniciada = MatrizariaService.iniciar_atendimento(
            solicitacao_id=sol_editada.id,
            usuario=self.tecnico_matrizaria_1,
            versao_esperada=sol_editada.versao,
        )
        with self.assertRaises(ValidationError) as ctx2:
            MatrizariaService.editar_solicitacao(
                solicitacao_id=sol_iniciada.id,
                usuario=self.solicitante_user,
                prensa=self.prensa_01,
                tipo_servico=self.tipo_ajuste,
                matriz_fisica=None,
                prioridade="NORMAL",
                descricao_solicitacao="Tentativa durante execucao",
                motivo_edicao="Tarde demais",
                versao_esperada=sol_iniciada.versao,
            )
        self.assertIn("Apenas solicitações com status 'SOLICITADO' podem ser editadas", str(ctx2.exception))

    def test_item5_transferencia_somente_matrizeiros_habilitados(self):
        # 1. Cria chamado e inicia com o tecnico 1
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Servico para teste de transferencia",
            solicitado_por=self.solicitante_user,
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)

        # 2. Verifica elegibilidade de matrizeiros
        candidatos = MatrizariaService.get_matrizeiros_habilitados(solicitacao=sol)
        # Deve conter tecnico_matrizaria_2
        self.assertIn(self.tecnico_matrizaria_2, candidatos)
        # NÃO deve conter o responsável atual (tecnico_matrizaria_1)
        self.assertNotIn(self.tecnico_matrizaria_1, candidatos)
        # NÃO deve conter solicitante, manutenção mecânica ou TV
        self.assertNotIn(self.solicitante_user, candidatos)
        self.assertNotIn(self.tecnico_maint_user, candidatos)
        self.assertNotIn(self.tv_user, candidatos)

        # 3. Líder com grupo Matrizaria também se torna habilitado
        self.lider_user.groups.add(self.grp_matrizaria)
        candidatos_com_lider = MatrizariaService.get_matrizeiros_habilitados(solicitacao=sol)
        self.assertIn(self.lider_user, candidatos_com_lider)
        self.lider_user.groups.remove(self.grp_matrizaria)

        # 4. Tentativa de transferir para usuário não habilitado é rejeitada
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.transferir_responsabilidade(
                solicitacao_id=sol.id,
                usuario_origem=self.tecnico_matrizaria_1,
                novo_responsavel=self.tecnico_maint_user,
                motivo="Tentativa para mecanico de manutencao",
                versao_esperada=sol.versao,
            )
        self.assertIn("não está habilitado para receber serviços da Matrizaria", str(ctx.exception))

        # 5. Transferência válida para tecnico 2
        sol_transf = MatrizariaService.transferir_responsabilidade(
            solicitacao_id=sol.id,
            usuario_origem=self.tecnico_matrizaria_1,
            novo_responsavel=self.tecnico_matrizaria_2,
            motivo="Transferencia por rendicao de turno",
            versao_esperada=sol.versao,
        )
        self.assertEqual(sol_transf.responsavel_atribuido, self.tecnico_matrizaria_2)


class MatrizariaSpec1AdjustsTestCase(MatrizariaBaseTestCase):
    """
    Testes automatizados cobrindo os requisitos da SPEC 1:
    1. Criação com prensa;
    2. Criação escolhendo Matrizaria (serviço interno);
    3. Máquina obrigatória quando destino exigir máquina;
    4. Relatórios suportam serviço interno;
    5. Excel suporta serviço interno;
    6. Edição existente suporta alteração de destino;
    7. Exclusão operacional funciona conforme permissão;
    8. Usuário comum não recebe privilégio de exclusão administrativa;
    9. Superuser consegue excluir registro relacionado usando fluxo administrativo especial;
    10. Registros protegidos continuam protegidos fora do fluxo do superuser;
    11. Nenhuma escrita em scada.
    """

    def setUp(self):
        super().setUp()
        self.superuser = User.objects.create_superuser(
            username="admin_supremo",
            email="admin@example.com",
            password="adminpassword123",
        )
        self.client = Client()

    def test_item1_criacao_com_prensa(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Ajuste normal em prensa",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        self.assertEqual(sol.destino, "MAQUINA")
        self.assertEqual(sol.prensa, self.prensa_01)
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")
        self.assertIn("PRENSA BOM 01", str(sol))

    def test_item2_criacao_escolhendo_matrizaria_sem_maquina(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Manutenção interna de bancada e gabarito",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        self.assertEqual(sol.destino, "MATRIZARIA")
        self.assertIsNone(sol.prensa)
        self.assertEqual(sol.prensa_nome_snapshot, "Matrizaria")
        self.assertIn("Matrizaria", str(sol))

    def test_item3_maquina_obrigatoria_quando_destino_maquina(self):
        # 1. MatrizariaService.criar_solicitacao sem prensa quando destino é MAQUINA
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.criar_solicitacao(
                prensa=None,
                tipo_servico=self.tipo_ajuste,
                descricao_solicitacao="Tentativa inválida sem máquina",
                solicitado_por=self.solicitante_user,
                destino="MAQUINA",
            )
        self.assertIn("A prensa é obrigatória", str(ctx.exception))

        # 2. Model clean() rejeita destino=MAQUINA com prensa=None
        sol_invalida_1 = SolicitacaoServicoMatrizaria(
            destino="MAQUINA",
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Teste model",
            solicitado_por=self.solicitante_user,
        )
        with self.assertRaises(ValidationError):
            sol_invalida_1.full_clean()

        # 3. Model clean() rejeita destino=MATRIZARIA com prensa associada
        sol_invalida_2 = SolicitacaoServicoMatrizaria(
            destino="MATRIZARIA",
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Teste model",
            solicitado_por=self.solicitante_user,
        )
        with self.assertRaises(ValidationError):
            sol_invalida_2.full_clean()

    def test_item4_relatorios_suportam_servico_interno(self):
        sol_int = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Serviço interno para relatório",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        sol_maq = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Serviço máquina para relatório",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )

        hoje = timezone.localdate()
        # Filtro com destino=MATRIZARIA
        res_matriz = MatrizariaService.consultar_relatorio(
            criterio_temporal="ABERTAS",
            data_inicio=hoje,
            data_fim=hoje,
            destino="MATRIZARIA",
        )
        ids_matriz = [r["id"] for r in res_matriz]
        self.assertIn(sol_int.id, ids_matriz)
        self.assertNotIn(sol_maq.id, ids_matriz)

        item = next(r for r in res_matriz if r["id"] == sol_int.id)
        self.assertEqual(item["prensa_nome"], "Matrizaria")

    def test_item5_excel_suporta_servico_interno(self):
        sol_int = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Serviço interno para Excel",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        self.client.force_login(self.superuser)
        url = reverse("matrizaria:exportar_excel") + "?criterio_temporal=ABERTAS&prensa=__MATRIZARIA__"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        import io
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws_servicos = wb["Serviços"]
        # Encontra a linha da solicitação interna
        linhas = list(ws_servicos.iter_rows(values_only=True))
        self.assertGreater(len(linhas), 1)
        header = linhas[0]
        col_prensa = header.index("Prensa")
        col_prot = header.index("Protocolo")
        
        achou = False
        for row in linhas[1:]:
            if row[col_prot] == f"SM #{sol_int.id}":
                self.assertEqual(row[col_prensa], "Matrizaria")
                achou = True
                break
        self.assertTrue(achou, "Solicitação interna não encontrada na planilha gerada")

    def test_item6_edicao_existente_suporta_mudanca_de_destino(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Criado com máquina por engano",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        sol_editada = MatrizariaService.editar_solicitacao(
            solicitacao_id=sol.id,
            usuario=self.solicitante_user,
            versao_esperada=sol.versao,
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            matriz_fisica=None,
            prioridade="NORMAL",
            descricao_solicitacao="Corrigido para serviço interno da Matrizaria",
            motivo_edicao="Não é serviço de prensa, é na bancada",
            destino="MATRIZARIA",
        )
        self.assertEqual(sol_editada.destino, "MATRIZARIA")
        self.assertIsNone(sol_editada.prensa)

        h = sol_editada.historico_transicoes.filter(tipo_evento="ALTERACAO_DADO").first()
        self.assertIsNotNone(h)
        self.assertIn("Destino", h.dados_modificados)

    def test_item7_exclusao_operacional_sucesso_e_restricoes(self):
        # 1. Cria chamado pendente
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado para teste de exclusão",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        sol_id = sol.id

        # 2. Usuário sem permissão (técnico de manutenção) tenta excluir
        with self.assertRaises(PermissionDenied):
            MatrizariaService.excluir_solicitacao_operacional(sol_id, self.tecnico_maint_user)

        # 3. Solicitante próprio exclui com sucesso
        MatrizariaService.excluir_solicitacao_operacional(sol_id, self.solicitante_user)
        self.assertFalse(SolicitacaoServicoMatrizaria.objects.filter(id=sol_id).exists())
        self.assertFalse(HistoricoTransicaoServicoMatrizaria.objects.filter(solicitacao_id=sol_id).exists())

        # 4. Chamado já iniciado NÃO pode ser excluído operacionalmente
        sol_iniciado = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado iniciado",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        MatrizariaService.iniciar_atendimento(sol_iniciado.id, self.tecnico_matrizaria_1, sol_iniciado.versao)
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.excluir_solicitacao_operacional(sol_iniciado.id, self.solicitante_user)
        self.assertIn("Apenas chamados pendentes podem ser excluídos", str(ctx.exception))

        # Se forçado para status SOLICITADO mas ainda com ciclo existente
        sol_iniciado.status = "SOLICITADO"
        sol_iniciado.save()
        with self.assertRaises(ValidationError) as ctx:
            MatrizariaService.excluir_solicitacao_operacional(sol_iniciado.id, self.solicitante_user)
        self.assertIn("atendimento técnico já foi iniciado", str(ctx.exception))

    def test_item8_usuario_comum_nao_tem_privilegio_exclusao_administrativa(self):
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Teste privilégio superuser",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        # Operador
        with self.assertRaises(PermissionDenied):
            AdminCascadeDeletionService.excluir_objeto(sol, self.solicitante_user)
        # Líder
        with self.assertRaises(PermissionDenied):
            AdminCascadeDeletionService.excluir_objeto(sol, self.lider_user)

    def test_item9_superuser_exclusao_administrativa_forcada_com_dependentes(self):
        # 1. Cria chamado, inicia atendimento (gerando Ciclo e Históricos)
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado com dependentes",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)
        self.assertGreater(sol.ciclos_execucao.count(), 0)
        self.assertGreater(sol.historico_transicoes.count(), 0)

        # 2. Preview de dependentes
        preview = AdminCascadeDeletionService.coletar_dependentes([sol])
        self.assertEqual(len(preview), 1)
        self.assertGreater(len(preview[0]["dependentes"]), 0)

        # 3. Exclusão forçada por superusuário
        sol_id = sol.id
        resultado = AdminCascadeDeletionService.excluir_objeto(sol, self.superuser)
        self.assertGreaterEqual(resultado["total_deletados"], 3)
        self.assertFalse(SolicitacaoServicoMatrizaria.objects.filter(id=sol_id).exists())
        self.assertFalse(CicloExecucaoMatrizaria.objects.filter(solicitacao_id=sol_id).exists())
        self.assertFalse(HistoricoTransicaoServicoMatrizaria.objects.filter(solicitacao_id=sol_id).exists())

    def test_item10_registros_protegidos_continuam_protegidos_fora_fluxo_superuser(self):
        # Cria chamado e ciclo
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado com PROTECT ativo",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        sol = MatrizariaService.iniciar_atendimento(sol.id, self.tecnico_matrizaria_1, sol.versao)
        
        # Tentar chamar delete() do Django diretamente no model DEVE falhar com ProtectedError
        with self.assertRaises(ProtectedError):
            sol.delete()

    def test_item11_nenhuma_escrita_em_scada(self):
        # Garante que os modelos residem no banco default
        sol = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Verificação de banco",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        self.assertEqual(sol._state.db, "default")

    def test_botao_excluir_aparece_exclusivamente_para_superuser(self):
        """
        Valida que o botão 'Excluir Solicitação' é exibido no HTML exclusivamente
        quando o usuário logado é SuperUser (is_superuser=True).
        Operadores, técnicos e líderes não visualizam o botão de exclusão.
        """
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Teste de visibilidade de botões",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        url = reverse("matrizaria:detalhe_servico", kwargs={"pk": sol.id})

        # 1. Solicitante comum acessa a tela de detalhes
        self.client.force_login(self.solicitante_user)
        res_user = self.client.get(url)
        self.assertEqual(res_user.status_code, 200)
        self.assertFalse(res_user.context["can_excluir"])
        self.assertNotContains(res_user, "Excluir Solicitação")
        self.assertContains(res_user, "Editar Solicitação")
        self.assertContains(res_user, "Cancelar Solicitação")
        self.assertContains(res_user, "bootstrap.bundle.min.js")
        self.assertContains(res_user, "openModalDirect")

        # 2. Usuário não-superuser tenta postar para exclusão -> Bloqueado
        res_delete_denied = self.client.post(reverse("matrizaria:excluir_solicitacao", kwargs={"pk": sol.id}))
        self.assertEqual(res_delete_denied.status_code, 302)
        self.assertTrue(SolicitacaoServicoMatrizaria.objects.filter(id=sol.id).exists())

        # 3. SuperUser acessa a tela de detalhes -> Botão Excluir presente
        self.client.force_login(self.superuser)
        res_super = self.client.get(url)
        self.assertEqual(res_super.status_code, 200)
        self.assertTrue(res_super.context["can_excluir"])
        self.assertContains(res_super, "Excluir Solicitação")
        self.assertContains(res_super, "Editar Solicitação")
        self.assertContains(res_super, "Cancelar Solicitação")

        # 4. SuperUser posta exclusão -> Excluído com sucesso
        res_delete_ok = self.client.post(reverse("matrizaria:excluir_solicitacao", kwargs={"pk": sol.id}))
        self.assertEqual(res_delete_ok.status_code, 302)
        self.assertFalse(SolicitacaoServicoMatrizaria.objects.filter(id=sol.id).exists())

    def test_abrir_chamado_matrizaria_via_post_sucesso(self):
        """
        1. Abrir chamado selecionando MATRIZARIA via POST:
           - HTTP esperado (302 -> redireciona para detalhe);
           - registro criado com prensa=None;
           - destino identificado como MATRIZARIA;
           - prensa_nome_snapshot preenchido como 'Matrizaria'.
        """
        self.client.force_login(self.solicitante_user)
        url = reverse("matrizaria:solicitar_servico")
        payload = {
            "prensa": "__MATRIZARIA__",
            "tipo_servico": self.tipo_ajuste.id,
            "descricao_solicitacao": "Serviço interno na bancada da Matrizaria",
            "prioridade": "NORMAL",
            "idempotency_key": "test-key-matrizaria-01",
        }
        res = self.client.post(url, data=payload)
        self.assertEqual(res.status_code, 302)

        sol = SolicitacaoServicoMatrizaria.objects.filter(descricao_solicitacao="Serviço interno na bancada da Matrizaria").first()
        self.assertIsNotNone(sol)
        self.assertRedirects(res, reverse("matrizaria:detalhe_servico", kwargs={"pk": sol.id}))
        self.assertEqual(sol.destino, "MATRIZARIA")
        self.assertIsNone(sol.prensa)
        self.assertEqual(sol.prensa_nome_snapshot, "Matrizaria")
        self.assertEqual(sol.equipamento_display, "Matrizaria")

    def test_abrir_chamado_prensa_real_via_post_sucesso(self):
        """
        2. Abrir chamado selecionando uma prensa real via POST:
           - registro criado;
           - prensa corretamente vinculada;
           - destino identificado como MAQUINA.
        """
        self.client.force_login(self.solicitante_user)
        url = reverse("matrizaria:solicitar_servico")
        payload = {
            "prensa": self.prensa_01.id,
            "tipo_servico": self.tipo_ajuste.id,
            "descricao_solicitacao": "Ajuste na prensa 01",
            "prioridade": "URGENTE",
            "idempotency_key": "test-key-prensa-01",
        }
        res = self.client.post(url, data=payload)
        self.assertEqual(res.status_code, 302)

        sol = SolicitacaoServicoMatrizaria.objects.filter(descricao_solicitacao="Ajuste na prensa 01").first()
        self.assertIsNotNone(sol)
        self.assertRedirects(res, reverse("matrizaria:detalhe_servico", kwargs={"pk": sol.id}))
        self.assertEqual(sol.destino, "MAQUINA")
        self.assertEqual(sol.prensa, self.prensa_01)
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")
        self.assertEqual(sol.equipamento_display, "PRENSA BOM 01")

    def test_abrir_chamado_sem_selecionar_destino_invalido(self):
        """
        3. Informar destino de máquina sem selecionar uma máquina (ou deixar campo vazio):
           - formulário inválido;
           - mensagem amigável;
           - nenhum registro criado.
        """
        self.client.force_login(self.solicitante_user)
        url = reverse("matrizaria:solicitar_servico")
        payload = {
            "prensa": "",
            "tipo_servico": self.tipo_ajuste.id,
            "descricao_solicitacao": "Tentativa sem selecionar destino",
            "prioridade": "NORMAL",
            "idempotency_key": "test-key-vazio",
        }
        res = self.client.post(url, data=payload)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Selecione o Destino do Serviço (Matrizaria ou uma Prensa / Máquina).")
        self.assertFalse(SolicitacaoServicoMatrizaria.objects.filter(descricao_solicitacao="Tentativa sem selecionar destino").exists())

    def test_servico_interno_no_kanban_exibe_matrizaria_sem_null(self):
        """
        4. Serviço interno aparece corretamente no Kanban:
           - exibe 'Matrizaria';
           - não exibe None, NULL, —, vazio ou erro de template.
        """
        sol = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado para validar Kanban visual",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        self.client.force_login(self.solicitante_user)
        res = self.client.get(reverse("matrizaria:kanban"))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Matrizaria")
        self.assertNotContains(res, "None")
        self.assertNotContains(res, "NULL")

    def test_servico_interno_nao_aparece_na_linha_do_tempo_de_nenhuma_prensa(self):
        """
        5. Serviço interno não aparece na linha do tempo de nenhuma prensa.
        """
        sol_int = MatrizariaService.criar_solicitacao(
            prensa=None,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado interno de bancada",
            solicitado_por=self.solicitante_user,
            destino="MATRIZARIA",
        )
        self.assertEqual(self.prensa_01.solicitacoes_matrizaria.count(), 0)
        self.assertEqual(self.prensa_02.solicitacoes_matrizaria.count(), 0)

        # Agora cria um chamado para a prensa 01
        sol_maq = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Chamado para a prensa 01",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        self.assertEqual(self.prensa_01.solicitacoes_matrizaria.count(), 1)
        self.assertIn(sol_maq, self.prensa_01.solicitacoes_matrizaria.all())
        self.assertNotIn(sol_int, self.prensa_01.solicitacoes_matrizaria.all())
        self.assertEqual(self.prensa_02.solicitacoes_matrizaria.count(), 0)

    def test_edicao_alternando_entre_maquina_e_matrizaria_via_post(self):
        """
        Testa edição via formulário / POST mudando de Máquina para Matrizaria e vice-versa.
        """
        sol = MatrizariaService.criar_solicitacao(
            prensa=self.prensa_01,
            tipo_servico=self.tipo_ajuste,
            descricao_solicitacao="Ajuste inicial em máquina",
            solicitado_por=self.solicitante_user,
            destino="MAQUINA",
        )
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")

        self.client.force_login(self.solicitante_user)
        # Edita para Matrizaria
        res_edit_1 = self.client.post(
            reverse("matrizaria:editar_solicitacao", kwargs={"pk": sol.id}),
            data={
                "versao": sol.versao,
                "prensa": "__MATRIZARIA__",
                "tipo_servico": self.tipo_ajuste.id,
                "prioridade": "NORMAL",
                "descricao_solicitacao": "Corrigido para bancada Matrizaria",
                "motivo_edicao": "Não era problema na prensa, e sim bancada",
            },
        )
        self.assertEqual(res_edit_1.status_code, 302)
        sol.refresh_from_db()
        self.assertEqual(sol.destino, "MATRIZARIA")
        self.assertIsNone(sol.prensa)
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")  # Preserva snapshot da abertura
        self.assertEqual(sol.equipamento_display, "Matrizaria")

        # Edita de volta para Prensa 02
        res_edit_2 = self.client.post(
            reverse("matrizaria:editar_solicitacao", kwargs={"pk": sol.id}),
            data={
                "versao": sol.versao,
                "prensa": self.prensa_02.id,
                "tipo_servico": self.tipo_ajuste.id,
                "prioridade": "NORMAL",
                "descricao_solicitacao": "Corrigido agora para prensa 02",
                "motivo_edicao": "Confirmado que é na prensa 02",
            },
        )
        self.assertEqual(res_edit_2.status_code, 302)
        sol.refresh_from_db()
        self.assertEqual(sol.destino, "MAQUINA")
        self.assertEqual(sol.prensa, self.prensa_02)
        self.assertEqual(sol.prensa_nome_snapshot, "PRENSA BOM 01")  # Preserva snapshot da abertura
        self.assertEqual(sol.equipamento_display, "PRENSA BOM 02")

