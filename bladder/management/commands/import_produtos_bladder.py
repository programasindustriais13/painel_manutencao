import os
import re
from decimal import Decimal
import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction

from bladder.models import ProdutoBladder, RecursoBladder
from production.models import ProductionBladder


def parse_decimal(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    s = str(val).strip()
    m = re.search(r'(\d+([,\.]\d+)?)', s)
    if m:
        num_str = m.group(1).replace(',', '.')
        return Decimal(num_str)
    return None


class Command(BaseCommand):
    help = "Importa de forma idempotente e segura o catálogo de produtos e especificações técnicas da planilha ET.029."

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default=None,
            help="Caminho do arquivo Excel ET.029"
        )

    def handle(self, *args, **options):
        file_path = options.get('file')
        if not file_path:
            # Prefere a nova versão atualizada se presente na raiz do projeto
            cand_nova = os.path.join(settings.BASE_DIR, "ET.029 - ESPECIFICAÇÃO TÉCNICA - BLA - TESTE (1).xlsx")
            cand_padrao = os.path.join(settings.BASE_DIR, "ET.029 - ESPECIFICAÇÃO TÉCNICA - BLA - TESTE.xlsx")
            if os.path.exists(cand_nova):
                file_path = cand_nova
            else:
                file_path = cand_padrao

        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"Arquivo não encontrado: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Lendo planilha: {file_path}..."))
        wb = openpyxl.load_workbook(file_path, data_only=True)

        ws_ext = wb.worksheets[0]
        ws_vulc = wb.worksheets[1] if len(wb.worksheets) > 1 else None

        # Se a planilha atual não tiver aba de vulcanização (ex: versão resumida só de extrusão),
        # tenta carregar vulcanização do arquivo de referência anterior para não perder histórico
        if ws_vulc is None:
            ref_vulc = os.path.join(settings.BASE_DIR, "ET.029 - ESPECIFICAÇÃO TÉCNICA - BLA - TESTE.xlsx")
            if os.path.exists(ref_vulc) and ref_vulc != file_path:
                try:
                    wb_vulc_ref = openpyxl.load_workbook(ref_vulc, data_only=True)
                    if len(wb_vulc_ref.worksheets) > 1:
                        ws_vulc = wb_vulc_ref.worksheets[1]
                        self.stdout.write(self.style.NOTICE(f"Carregando dados complementares de vulcanização de: {ref_vulc}"))
                except Exception as e:
                    self.stdout.write(self.style.WARNING(f"Não foi possível carregar vulcanização complementar: {e}"))

        # 1. Mapeamento de medidas e detalhes da vulcanização por código BLA
        vulc_data = {}
        if ws_vulc is not None:
            for r in range(5, ws_vulc.max_row + 1):
                pneu = ws_vulc.cell(r, 1).value
                bla = ws_vulc.cell(r, 2).value
                tempo = ws_vulc.cell(r, 3).value
                peso_vulc = parse_decimal(ws_vulc.cell(r, 4).value)
                circ = ws_vulc.cell(r, 5).value
                alt = ws_vulc.cell(r, 6).value
                status_vulc = ws_vulc.cell(r, 7).value

                if bla and str(bla).strip().startswith("BLA"):
                    bla_code = str(bla).strip().upper()
                    if bla_code not in vulc_data:
                        vulc_data[bla_code] = {
                            "pneus": [],
                            "peso_vulc": peso_vulc,
                            "tempo_min": 120,
                            "circ": str(circ).strip() if circ and circ != '-' else "",
                            "alt": str(alt).strip() if alt and alt != '-' else "",
                            "status": status_vulc or "ATIVO"
                        }
                    if pneu:
                        pneu_clean = str(pneu).strip()
                        if pneu_clean not in vulc_data[bla_code]["pneus"]:
                            vulc_data[bla_code]["pneus"].append(pneu_clean)

        # 2. Leitura da extrusão (Sheet 0) e gravação no banco
        criados = 0
        atualizados = 0

        col1_header = str(ws_ext.cell(5, 1).value or "").strip().lower()
        is_medida = "medida" in col1_header

        with transaction.atomic(using='default'):
            for r in range(6, ws_ext.max_row + 1):
                col1_val = ws_ext.cell(r, 1).value
                cod = ws_ext.cell(r, 2).value
                mat = ws_ext.cell(r, 3).value
                comp_ext = parse_decimal(ws_ext.cell(r, 4).value)
                comp_chanf = parse_decimal(ws_ext.cell(r, 5).value)
                peso_tarugo = parse_decimal(ws_ext.cell(r, 6).value)
                diam = parse_decimal(ws_ext.cell(r, 7).value)
                status_raw = ws_ext.cell(r, 8).value

                if not cod or not str(cod).strip().upper().startswith("BLA"):
                    continue

                cod_str = str(cod).strip().upper()
                col1_str = str(col1_val).strip() if col1_val else ""
                vd = vulc_data.get(cod_str, {})
                existing_prod = ProdutoBladder.objects.filter(codigo=cod_str).first()

                # Monta descrição compreensível
                pneus_list = vd.get("pneus", [])
                if is_medida and col1_str:
                    medida = col1_str
                    if pneus_list:
                        pneus_txt = ", ".join(pneus_list[:2])
                        if len(pneus_list) > 2:
                            pneus_txt += f" (+{len(pneus_list)-2})"
                        desc_resumida = f"{medida} — {pneus_txt}"
                    else:
                        desc_resumida = f"{medida}"
                elif pneus_list:
                    desc_resumida = ", ".join(pneus_list[:3])
                    if len(pneus_list) > 3:
                        desc_resumida += f" (+{len(pneus_list)-3} medidas)"
                else:
                    desc_resumida = f"Bladder {cod_str}"

                # Nomenclatura antiga: preserva histórica se coluna 1 for 'Medida', senão atualiza
                if is_medida:
                    if existing_prod and existing_prod.nomenclatura_antiga:
                        nomenclatura_antiga = existing_prod.nomenclatura_antiga
                    else:
                        nomenclatura_antiga = col1_str
                else:
                    nomenclatura_antiga = col1_str

                # Status
                st = "ATIVO"
                if status_raw and "DESENVOLVIMENTO" in str(status_raw).upper():
                    st = "EM DESENVOLVIMENTO"
                elif status_raw and "INATIVO" in str(status_raw).upper():
                    st = "INATIVO"

                # Link opcional com ProductionBladder
                prod_bla_ref = ProductionBladder.objects.filter(codigo_bladder=cod_str).first()

                # Matriz de extrusão
                mat_str = str(mat).strip() if mat and mat != '-' else ""
                if mat_str:
                    RecursoBladder.objects.get_or_create(
                        codigo=mat_str,
                        defaults={
                            "categoria": "MATRIZ",
                            "nome": f"Matriz de Extrusão {mat_str}",
                            "ativo": True,
                        }
                    )

                # Vulcanização segura
                peso_vulc = vd.get("peso_vulc")
                if peso_vulc is None and existing_prod:
                    peso_vulc = existing_prod.peso_vulcanizado_kg

                tempo_vulc = vd.get("tempo_min", 120)
                if existing_prod and existing_prod.tempo_vulcanizacao_min:
                    tempo_vulc = existing_prod.tempo_vulcanizacao_min

                circ = vd.get("circ", "")
                if not circ and existing_prod:
                    circ = existing_prod.circunferencia_centro_mm

                alt = vd.get("alt", "")
                if not alt and existing_prod:
                    alt = existing_prod.altura_cm

                defaults = {
                    "descricao": desc_resumida,
                    "nomenclatura_antiga": nomenclatura_antiga,
                    "matriz_extrusao": mat_str,
                    "comprimento_extrusao_cm": comp_ext,
                    "comprimento_chanfrado_cm": comp_chanf,
                    "peso_tarugo_kg": peso_tarugo,
                    "diametro_tarugo_mm": diam,
                    "peso_vulcanizado_kg": peso_vulc,
                    "tempo_vulcanizacao_min": tempo_vulc,
                    "circunferencia_centro_mm": circ,
                    "altura_cm": alt,
                    "production_bladder": prod_bla_ref,
                    "status": st,
                    "ativo": st != "INATIVO",
                }

                obj, created = ProdutoBladder.objects.update_or_create(
                    codigo=cod_str,
                    defaults=defaults
                )

                if created:
                    criados += 1
                    self.stdout.write(self.style.SUCCESS(f" [NOVO] {cod_str} - {desc_resumida} (Tarugo: {peso_tarugo} kg, Matriz: {mat_str})"))
                else:
                    atualizados += 1
                    self.stdout.write(self.style.NOTICE(f" [ATUALIZADO] {cod_str} - {desc_resumida} (Tarugo: {peso_tarugo} kg, Matriz: {mat_str})"))

        self.stdout.write(self.style.SUCCESS(
            f"\nImportação Concluída com Sucesso!\n"
            f"Produtos Criados: {criados}\n"
            f"Produtos Atualizados: {atualizados}\n"
            f"Total de Modelos no Catálogo: {ProdutoBladder.objects.count()}"
        ))
