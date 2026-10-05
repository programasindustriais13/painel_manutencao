import calendar
import datetime
from django.db import transaction
from django.db.models import Sum, Count, Q, Case, When, Value, IntegerField
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from .models import (
    ProcessoBladder,
    ConfiguracaoEscalaBladder,
    AjusteEscalaExcepcionalBladder,
    FuncionarioApoioBladder,
    ProdutoBladder,
    RecursoBladder,
    OrdemProducaoBladder,
    SaldoPendenteBladder,
    ApontamentoTurnoBladder,
    FechamentoTurnoBladder,
    ItemFechamentoTurnoBladder,
    CategoriaDesvioBladder,
    HistoricoApontamentoBladder,
    HistoricoProgramacaoBladder,
    MensagemPassagemTurnoBladder,
    AcaoMensagemTurnoBladder,
)
from .forms import (
    OrdemProducaoBladderForm,
    ReprogramarOrdemForm,
    CancelarOrdemForm,
    ApontamentoTurnoForm,
    CorrecaoApontamentoForm,
    MensagemPassagemTurnoForm,
)
from .decorators import (
    lider_bladder_required,
    operador_ou_lider_bladder_required,
    user_is_lider_bladder,
)
from .services import (
    obter_configuracao_escala_ativa,
    calcular_turma_do_dia,
    verificar_usuario_apoio_no_dia,
    criar_ordem_producao_com_saldos,
    encerrar_op_parcial_ou_total,
    reprogramar_ordem_producao,
    cancelar_ordem_producao,
    registrar_ou_atualizar_apontamento,
    corrigir_apontamento_operador_ou_lider,
    obter_saldos_pendentes_produto,
    executar_fechamento_turno,
    calcular_proximo_turno_operacional,
    criar_mensagem_passagem_turno,
    registrar_ciencia_mensagem_turno,
    resolver_mensagem_acompanhamento,
    repassar_mensagem_acompanhamento,
    obter_mensagens_recebidas_turno,
)


@login_required
@lider_bladder_required
def dashboard(request):
    """
    Painel Geral e Dashboard Operacional do Setor de Bladder.
    Exibe indicadores do dia, status das ordens, escala e saldos pendentes.
    """
    hoje = timezone.localdate()
    turma_hoje, is_ajuste, motivo_ajuste = calcular_turma_do_dia(hoje)
    config_escala = obter_configuracao_escala_ativa()

    # Ordens do dia
    ordens_dia = OrdemProducaoBladder.objects.filter(
        data_programada=hoje
    ).select_related('processo', 'produto').prefetch_related('recursos_alocados')

    total_ops_dia = ordens_dia.count()
    planejado_dia = ordens_dia.aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    realizado_dia = ordens_dia.aggregate(s=Sum('quantidade_realizada'))['s'] or 0

    pendentes_dia = ordens_dia.filter(status='PENDENTE').count()
    em_execucao_dia = ordens_dia.filter(status='EM_EXECUCAO').count()
    concluidas_dia = ordens_dia.filter(status='CONCLUIDA').count()
    parciais_dia = ordens_dia.filter(status='PARCIAL').count()
    canceladas_dia = ordens_dia.filter(status='CANCELADA').count()

    # Horário de encerramento da escala ativa
    hora_fim_config = config_escala.hora_fim if (config_escala and config_escala.hora_fim) else datetime.time(18, 0)
    agora = timezone.localtime(timezone.now())
    dt_limite_hoje = timezone.make_aware(
        datetime.datetime.combine(hoje, hora_fim_config),
        timezone.get_current_timezone()
    )
    turno_hoje_encerrado = agora > dt_limite_hoje

    # Cálculo dinâmico de ordens atrasadas
    if turno_hoje_encerrado:
        atrasadas_qs = OrdemProducaoBladder.objects.filter(
            status__in=['PENDENTE', 'EM_EXECUCAO'],
            data_programada__lte=hoje
        )
    else:
        atrasadas_qs = OrdemProducaoBladder.objects.filter(
            status__in=['PENDENTE', 'EM_EXECUCAO'],
            data_programada__lt=hoje
        )
    total_atrasadas = atrasadas_qs.count()

    # Saldos pendentes consolidados por produto
    saldos_pendentes = SaldoPendenteBladder.objects.filter(
        status='PENDENTE'
    ).values('produto__codigo', 'produto__descricao').annotate(
        total_saldo=Sum('quantidade'),
        qtd_ops=Count('id')
    ).order_by('-total_saldo')

    total_saldos_pendentes_un = SaldoPendenteBladder.objects.filter(
        status='PENDENTE'
    ).aggregate(s=Sum('quantidade'))['s'] or 0
    total_saldos_pendentes_ops = SaldoPendenteBladder.objects.filter(
        status='PENDENTE'
    ).count()

    excedente_dia = sum(op.excedente for op in ordens_dia)
    pendencia_dia = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ordens_dia if op.status in ['CONCLUIDA', 'PARCIAL'])
    diferenca_dia = realizado_dia - planejado_dia

    # Motivos das pendências (para todos os saldos PENDENTES no Ledger)
    motivos_map = {
        'PROBLEMA_EQUIPAMENTO': {
            'codigo': 'PROBLEMA_EQUIPAMENTO',
            'nome': 'Problema de equipamento',
            'icone': 'bi-gear-wide-connected',
            'cor': 'danger',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'FALTA_MATERIA_PRIMA': {
            'codigo': 'FALTA_MATERIA_PRIMA',
            'nome': 'Falta de matéria-prima',
            'icone': 'bi-boxes',
            'cor': 'warning',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'PROBLEMA_QUALIDADE': {
            'codigo': 'PROBLEMA_QUALIDADE',
            'nome': 'Problema de qualidade',
            'icone': 'bi-shield-x',
            'cor': 'danger',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'MANUTENCAO': {
            'codigo': 'MANUTENCAO',
            'nome': 'Manutenção',
            'icone': 'bi-tools',
            'cor': 'info',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'FALTA_OPERADOR': {
            'codigo': 'FALTA_OPERADOR',
            'nome': 'Falta de operador',
            'icone': 'bi-person-dash',
            'cor': 'primary',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'FALTA_COMPONENTE': {
            'codigo': 'FALTA_COMPONENTE',
            'nome': 'Falta de componente',
            'icone': 'bi-cpu',
            'cor': 'info',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'PROBLEMA_OPERACIONAL': {
            'codigo': 'PROBLEMA_OPERACIONAL',
            'nome': 'Problema operacional',
            'icone': 'bi-person-exclamation',
            'cor': 'primary',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'ALTERACAO_PROGRAMACAO': {
            'codigo': 'ALTERACAO_PROGRAMACAO',
            'nome': 'Alteração de programação',
            'icone': 'bi-calendar-event',
            'cor': 'secondary',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
        'OUTRO': {
            'codigo': 'OUTRO',
            'nome': 'Outro',
            'icone': 'bi-three-dots',
            'cor': 'secondary',
            'quantidade_pecas': 0,
            'quantidade_ops': 0,
        },
    }

    saldos_abertos = SaldoPendenteBladder.objects.filter(
        status='PENDENTE'
    ).select_related('op_origem', 'produto', 'fechamento').prefetch_related(
        'op_origem__itens_fechamento',
        'op_origem__itens_fechamento__categorias',
        'fechamento__itens',
        'fechamento__itens__categorias'
    )

    total_pecas_pendentes_motivos = 0
    for s in saldos_abertos:
        total_pecas_pendentes_motivos += s.quantidade
        item_fech = None
        if s.fechamento:
            for it in s.fechamento.itens.all():
                if it.ordem_id == s.op_origem_id:
                    item_fech = it
                    break
        if not item_fech:
            item_fech = s.op_origem.itens_fechamento.order_by('-id').first()

        cats_encontradas = []
        if item_fech:
            m2m_cats = list(item_fech.categorias.all())
            if m2m_cats:
                for c in m2m_cats:
                    cats_encontradas.append(c.codigo)
            elif item_fech.motivo:
                cats_encontradas.append(item_fech.motivo)

        if not cats_encontradas:
            motivo_txt = (s.op_origem.motivo_pendencia or "").upper()
            if 'EQUIPAMENTO' in motivo_txt or 'MÁQUINA' in motivo_txt or 'MAQUINA' in motivo_txt:
                cats_encontradas.append('PROBLEMA_EQUIPAMENTO')
            elif 'MATÉRIA' in motivo_txt or 'MATERIA' in motivo_txt or 'MP' in motivo_txt:
                cats_encontradas.append('FALTA_MATERIA_PRIMA')
            elif 'COMPONENTE' in motivo_txt:
                cats_encontradas.append('FALTA_COMPONENTE')
            elif 'OPERACIONAL' in motivo_txt:
                cats_encontradas.append('PROBLEMA_OPERACIONAL')
            elif 'PROGRAMAÇÃO' in motivo_txt or 'PROGRAMACAO' in motivo_txt:
                cats_encontradas.append('ALTERACAO_PROGRAMACAO')
            elif 'QUALIDADE' in motivo_txt:
                cats_encontradas.append('PROBLEMA_QUALIDADE')
            elif 'MANUTENÇÃO' in motivo_txt or 'MANUTENCAO' in motivo_txt:
                cats_encontradas.append('MANUTENCAO')
            elif 'OPERADOR' in motivo_txt:
                cats_encontradas.append('FALTA_OPERADOR')
            else:
                cats_encontradas.append('OUTRO')

        for cat in cats_encontradas:
            if cat not in motivos_map:
                motivos_map[cat] = {
                    'codigo': cat,
                    'nome': cat.replace('_', ' ').capitalize(),
                    'icone': 'bi-tag',
                    'cor': 'secondary',
                    'quantidade_pecas': 0,
                    'quantidade_ops': 0,
                }
            motivos_map[cat]['quantidade_pecas'] += s.quantidade
            motivos_map[cat]['quantidade_ops'] += 1

    motivos_pendencias_lista = []
    for k, v in motivos_map.items():
        v['percentual'] = round((v['quantidade_pecas'] / total_pecas_pendentes_motivos * 100), 1) if total_pecas_pendentes_motivos > 0 else 0
        motivos_pendencias_lista.append(v)

    # Consumo teórico do dia (kg)
    peso_teorico_planejado = sum(
        (op.consumo_teorico_planejado_kg or 0) for op in ordens_dia
    )
    peso_teorico_realizado = sum(
        (op.consumo_teorico_realizado_kg or 0) for op in ordens_dia
    )

    percentual_cumprimento_dia = round((realizado_dia / planejado_dia) * 100, 1) if planejado_dia > 0 else 0

    context = {
        'hoje': hoje,
        'turma_hoje': turma_hoje,
        'turma_hoje_display': "Turma A" if turma_hoje == 'TURMA_A' else ("Turma B" if turma_hoje == 'TURMA_B' else turma_hoje),
        'is_ajuste': is_ajuste,
        'motivo_ajuste': motivo_ajuste,
        'config_escala': config_escala,
        'hora_fim_config': hora_fim_config,
        'turno_encerrado': turno_hoje_encerrado,
        'ordens_dia': ordens_dia,
        'total_ops_dia': total_ops_dia,
        'planejado_dia': planejado_dia,
        'realizado_dia': realizado_dia,
        'excedente_dia': excedente_dia,
        'pendencia_dia': pendencia_dia,
        'diferenca_dia': diferenca_dia,
        'percentual_cumprimento_dia': percentual_cumprimento_dia,
        'pendentes_dia': pendentes_dia,
        'em_execucao_dia': em_execucao_dia,
        'concluidas_dia': concluidas_dia,
        'parciais_dia': parciais_dia,
        'canceladas_dia': canceladas_dia,
        'total_atrasadas': total_atrasadas,
        'saldos_pendentes': saldos_pendentes,
        'total_saldos_pendentes_un': total_saldos_pendentes_un,
        'total_saldos_pendentes_ops': total_saldos_pendentes_ops,
        'motivos_pendencias': motivos_pendencias_lista,
        'total_pecas_pendentes_motivos': total_pecas_pendentes_motivos,
        'peso_teorico_planejado': round(peso_teorico_planejado, 3),
        'peso_teorico_realizado': round(peso_teorico_realizado, 3),
        'is_lider': user_is_lider_bladder(request.user),
    }
    return render(request, 'bladder/dashboard.html', context)


@login_required
@operador_ou_lider_bladder_required
def operador_turno(request):
    """
    Quadro Operacional do Turno para Chão de Fábrica (Tablet Industrial).
    Visão limpa, compacta e orientada à consulta técnica das programações do turno.
    Não exige apontamentos contínuos durante a execução.
    """
    data_param = request.GET.get('data')
    if data_param:
        try:
            hoje = datetime.datetime.strptime(data_param, '%Y-%m-%d').date()
        except ValueError:
            hoje = timezone.localdate()
    else:
        hoje = timezone.localdate()
    turma_hoje, is_ajuste, motivo_ajuste = calcular_turma_do_dia(hoje)
    is_apoio, apoio_obj = verificar_usuario_apoio_no_dia(request.user, hoje)

    # Ordens programadas para o turno do dia
    ordens = OrdemProducaoBladder.objects.filter(
        data_programada=hoje
    ).exclude(
        status='CANCELADA'
    ).select_related(
        'processo', 'processo__maquina', 'produto'
    ).prefetch_related(
        'recursos_alocados'
    ).order_by(
        '-prioridade', 'numero_ordem'
    )

    perfil_operacional = getattr(request.user, 'perfil_operacional_bladder', None)
    if perfil_operacional is None and not request.user.is_anonymous:
        from .models import PerfilOperacionalBladder
        perfil_operacional = PerfilOperacionalBladder.objects.filter(usuario=request.user).first()

    # Verifica se já houve fechamento do turno hoje
    fechamento_hoje = FechamentoTurnoBladder.objects.filter(
        data_turno=hoje,
        turma=turma_hoje,
        status='CONCLUIDO'
    ).select_related('operador').first()

    # Mensagens de passagem de turno recebidas para o turno de hoje
    recados_recebidos = obter_mensagens_recebidas_turno(data_turno=hoje, turma=turma_hoje, usuario=request.user)
    total_recados = len(recados_recebidos)
    recados_pendentes_ciencia = sum(1 for r in recados_recebidos if not getattr(r, 'usuario_ciente', False))
    recados_urgentes = sum(1 for r in recados_recebidos if r.prioridade == 'URGENTE')
    form_recado = MensagemPassagemTurnoForm()

    context = {
        'hoje': hoje,
        'turma_hoje': turma_hoje,
        'turma_hoje_display': "Turma A" if turma_hoje == 'TURMA_A' else ("Turma B" if turma_hoje == 'TURMA_B' else turma_hoje),
        'perfil_operacional': perfil_operacional,
        'is_apoio': is_apoio,
        'apoio_obj': apoio_obj,
        'ordens': ordens,
        'fechamento_hoje': fechamento_hoje,
        'recados_recebidos': recados_recebidos,
        'total_recados': total_recados,
        'recados_pendentes_ciencia': recados_pendentes_ciencia,
        'recados_urgentes': recados_urgentes,
        'form_recado': form_recado,
        'is_lider': user_is_lider_bladder(request.user),
    }
    return render(request, 'bladder/operador_turno.html', context)


@login_required
@operador_ou_lider_bladder_required
def fechamento_turno(request):
    """
    Tela de Fechamento do Turno operacional pelo Operador.
    Solicita para cada OP programada no dia:
    - Quantidade realizada
    - Motivo se não cumpriu a meta (obrigatório se saldo > 0)
    - Descrição se motivo for Outro
    - Observação opcional
    Inclui seção informativa de Passagem para o Próximo Turno.
    """
    hoje = timezone.localdate()
    turma_hoje, _, _ = calcular_turma_do_dia(hoje)
    turma_hoje_display = "Turma A" if turma_hoje == 'TURMA_A' else ("Turma B" if turma_hoje == 'TURMA_B' else turma_hoje)

    # Verifica se o turno já foi encerrado
    fechamento_existente = FechamentoTurnoBladder.objects.filter(
        data_turno=hoje,
        turma=turma_hoje,
        status='CONCLUIDO'
    ).select_related('operador').prefetch_related('itens__ordem__produto', 'itens__ordem__processo').first()

    ordens = OrdemProducaoBladder.objects.filter(
        data_programada=hoje
    ).exclude(
        status='CANCELADA'
    ).select_related(
        'processo', 'processo__maquina', 'produto'
    ).order_by('numero_ordem')

    if request.method == 'POST':
        if fechamento_existente:
            messages.error(request, f"O turno de hoje já foi encerrado por {fechamento_existente.operador.get_full_name() or fechamento_existente.operador.username}.")
            return redirect('bladder:operador')

        itens_dados = []
        for op in ordens:
            qtd_str = request.POST.get(f'qtd_realizada_{op.id}', '0').strip()
            categorias_list = request.POST.getlist(f'categorias_{op.id}')
            descricao_desvio = (
                request.POST.get(f'descricao_desvio_{op.id}', '') or
                request.POST.get(f'motivo_outro_{op.id}', '') or
                request.POST.get(f'obs_{op.id}', '') or
                ""
            ).strip()
            motivo = request.POST.get(f'motivo_{op.id}', '').strip()
            motivo_outro = request.POST.get(f'motivo_outro_{op.id}', '').strip()
            obs = request.POST.get(f'obs_{op.id}', '').strip()

            try:
                qtd_int = int(qtd_str)
            except ValueError:
                qtd_int = -1

            itens_dados.append({
                'ordem_id': op.id,
                'quantidade_realizada': qtd_int,
                'categorias': categorias_list,
                'descricao_desvio': descricao_desvio,
                'motivo': motivo,
                'motivo_outro': motivo_outro,
                'observacao': obs
            })

        observacoes_gerais = request.POST.get('observacoes_gerais', '').strip()

        try:
            fechamento = executar_fechamento_turno(
                data_turno=hoje,
                operador=request.user,
                itens_dados=itens_dados,
                observacoes=observacoes_gerais
            )
            messages.success(request, f"Fechamento do turno ({turma_hoje_display}) realizado com sucesso! Todas as {len(ordens)} OPs foram auditadas.")
            return redirect('bladder:operador')
        except (ValueError, PermissionError) as e:
            messages.error(request, f"Erro no fechamento: {str(e)}")

    categorias_desvio = CategoriaDesvioBladder.objects.filter(ativo=True).order_by('ordem', 'nome')

    # Passagem para o próximo turno (seção informativa e acompanhamentos pendentes)
    recados_criados_turno = MensagemPassagemTurnoBladder.objects.filter(
        data_turno_origem=hoje,
        turma_origem=turma_hoje
    ).select_related('autor', 'ordem_producao', 'processo', 'maquina', 'produto').order_by('-created_at')

    acompanhamentos_abertos = MensagemPassagemTurnoBladder.objects.filter(
        Q(data_turno_destino=hoje, turma_destino=turma_hoje) |
        Q(data_turno_destino__lt=hoje, turma_destino=turma_hoje),
        tipo='ACOMPANHAMENTO',
        status='ABERTA'
    ).select_related('autor', 'ordem_producao', 'processo', 'maquina', 'produto').order_by('-prioridade', '-created_at')

    context = {
        'hoje': hoje,
        'turma_hoje': turma_hoje,
        'turma_hoje_display': turma_hoje_display,
        'ordens': ordens,
        'fechamento_existente': fechamento_existente,
        'motivo_choices': ItemFechamentoTurnoBladder.MOTIVO_CHOICES,
        'categorias_desvio': categorias_desvio,
        'recados_criados_turno': recados_criados_turno,
        'acompanhamentos_abertos': acompanhamentos_abertos,
        'form_recado': MensagemPassagemTurnoForm(),
    }
    return render(request, 'bladder/fechamento_turno.html', context)


@login_required
@lider_bladder_required
def cronograma_calendario(request):
    """
    Visão de Calendário Mensal para planejamento e controle pelo Líder.
    Exibe a alternância Turma A / Turma B e as programações por dia.
    """
    hoje = timezone.localdate()
    ano = int(request.GET.get('ano', hoje.year))
    mes = int(request.GET.get('mes', hoje.month))

    # Garante ano e mês válidos
    if mes < 1:
        mes = 12
        ano -= 1
    elif mes > 12:
        mes = 1
        ano += 1

    cal = calendar.Calendar(firstweekday=6)  # Domingo primeiro
    dias_matriz = cal.monthdatescalendar(ano, mes)

    # Carrega OPs do mês
    primeiro_dia = datetime.date(ano, mes, 1)
    ultimo_dia = datetime.date(ano, mes, calendar.monthrange(ano, mes)[1])

    ops_mes = OrdemProducaoBladder.objects.filter(
        data_programada__gte=dias_matriz[0][0],
        data_programada__lte=dias_matriz[-1][-1]
    ).select_related('processo', 'produto')

    ops_por_data = {}
    for op in ops_mes:
        dt_str = op.data_programada.strftime('%Y-%m-%d')
        if dt_str not in ops_por_data:
            ops_por_data[dt_str] = []
        ops_por_data[dt_str].append(op)

    # Monta matriz com informações de escala por dia
    calendario_mes = []
    for semana in dias_matriz:
        semana_dados = []
        for dia in semana:
            turma, is_ajuste, _ = calcular_turma_do_dia(dia)
            dt_str = dia.strftime('%Y-%m-%d')
            ops_dia = ops_por_data.get(dt_str, [])
            total_plan = sum(op.quantidade_planejada for op in ops_dia)
            total_real = sum(op.quantidade_realizada for op in ops_dia)

            semana_dados.append({
                'data': dia,
                'no_mes': dia.month == mes,
                'is_hoje': dia == hoje,
                'turma': turma,
                'is_ajuste': is_ajuste,
                'ops': ops_dia,
                'total_plan': total_plan,
                'total_real': total_real,
            })
        calendario_mes.append(semana_dados)

    context = {
        'ano': ano,
        'mes': mes,
        'mes_nome': calendar.month_name[mes],
        'ano_ant': ano if mes > 1 else ano - 1,
        'mes_ant': mes - 1 if mes > 1 else 12,
        'ano_prox': ano if mes < 12 else ano + 1,
        'mes_prox': mes + 1 if mes < 12 else 1,
        'calendario_mes': calendario_mes,
        'hoje': hoje,
    }
    return render(request, 'bladder/cronograma_calendario.html', context)


@login_required
@lider_bladder_required
def ordens_lista(request):
    """Listagem geral de Ordens de Produção com filtros avançados."""
    qs = OrdemProducaoBladder.objects.select_related(
        'processo', 'processo__maquina', 'produto', 'criado_por'
    ).prefetch_related(
        'saldos_gerados', 'apontamentos', 'historicos_programacao'
    ).order_by(
        '-data_programada', '-id'
    )

    # Filtros
    status = request.GET.get('status')
    processo_id = request.GET.get('processo')
    produto_id = request.GET.get('produto')
    turma = request.GET.get('turma')
    data_ini = request.GET.get('data_ini')
    data_fim = request.GET.get('data_fim')

    if status:
        qs = qs.filter(status=status)
    if processo_id:
        qs = qs.filter(processo_id=processo_id)
    if produto_id:
        qs = qs.filter(produto_id=produto_id)
    if turma:
        qs = qs.filter(turma_prevista=turma)
    if data_ini:
        qs = qs.filter(data_programada__gte=data_ini)
    if data_fim:
        qs = qs.filter(data_programada__lte=data_fim)

    processos = ProcessoBladder.objects.filter(ativo=True)
    produtos = ProdutoBladder.objects.filter(ativo=True)

    context = {
        'ordens': qs[:100],
        'total_registros': qs.count(),
        'processos': processos,
        'produtos': produtos,
        'status_selecionado': status,
        'processo_selecionado': int(processo_id) if processo_id else None,
        'produto_selecionado': int(produto_id) if produto_id else None,
        'turma_selecionada': turma,
        'data_ini': data_ini,
        'data_fim': data_fim,
        'is_lider': user_is_lider_bladder(request.user),
    }
    return render(request, 'bladder/ordens_lista.html', context)


@login_required
@lider_bladder_required
def ordem_nova(request):
    """Criação de nova Ordem de Produção pelo Líder com decisão explícita sobre pendências."""
    if request.method == 'POST':
        form = OrdemProducaoBladderForm(request.POST)
        if form.is_valid():
            try:
                decisao = request.POST.get('decisao_pendencias', request.POST.get('decisao_pendencia', 'IGNORAR')).strip().upper()
                saldos_selecionados_ids = []
                if decisao == 'INCORPORAR':
                    saldos_ids_raw = request.POST.getlist('saldos_selecionados')
                    if saldos_ids_raw:
                        saldos_selecionados_ids = [int(sid) for sid in saldos_ids_raw if str(sid).isdigit()]
                    else:
                        # Se o líder escolheu incorporar pendências mas não filtrou IDs individuais (incorporar todas do produto)
                        saldos_selecionados_ids = list(
                            obter_saldos_pendentes_produto(form.cleaned_data['produto'].id).values_list('id', flat=True)
                        )

                op = criar_ordem_producao_com_saldos(
                    processo=form.cleaned_data['processo'],
                    produto=form.cleaned_data['produto'],
                    data_programada=form.cleaned_data['data_programada'],
                    quantidade_nova=form.cleaned_data['quantidade_nova'],
                    prioridade=form.cleaned_data['prioridade'],
                    saldos_selecionados_ids=saldos_selecionados_ids if decisao == 'INCORPORAR' else None,
                    recursos_alocados=form.cleaned_data.get('recursos_alocados'),
                    recursos_observacoes=form.cleaned_data.get('recursos_observacoes', ''),
                    observacoes=form.cleaned_data.get('observacoes', ''),
                    usuario=request.user
                )
                msg = f"Ordem de Produção {op.numero_ordem} programada com sucesso! Total a produzir: {op.quantidade_planejada} un."
                if op.saldo_anterior_incorporado > 0:
                    msg += f" (Incluindo {op.saldo_anterior_incorporado} un de saldo pendente anterior incorporado)."
                messages.success(request, msg)
                return redirect('bladder:ordem_detalhe', pk=op.pk)
            except Exception as e:
                messages.error(request, f"Erro ao criar Ordem de Produção: {str(e)}")
    else:
        initial = {'data_programada': timezone.localdate()}
        data_param = request.GET.get('data')
        if data_param:
            try:
                initial['data_programada'] = datetime.datetime.strptime(data_param, '%Y-%m-%d').date()
            except ValueError:
                pass
        produto_param = request.GET.get('produto')
        if produto_param and produto_param.isdigit():
            initial['produto'] = int(produto_param)
        form = OrdemProducaoBladderForm(initial=initial)

    context = {
        'form': form,
        'titulo': "Nova Programação de Bladder",
    }
    return render(request, 'bladder/ordem_form.html', context)


@login_required
@lider_bladder_required
def ordem_detalhe(request, pk):
    """Ficha completa da Ordem de Produção com saldo, apontamentos e histórico."""
    op = get_object_or_404(
        OrdemProducaoBladder.objects.select_related(
            'processo', 'produto', 'criado_por', 'atualizado_por', 'cancelado_por'
        ).prefetch_related('recursos_alocados'),
        pk=pk
    )

    apontamentos = op.apontamentos.select_related('operador').order_by('-data_hora_inicio')
    historicos = op.historicos_programacao.select_related('usuario').order_by('-created_at')
    saldos_gerados = op.saldos_gerados.all()
    saldos_incorporados = op.saldos_incorporados.select_related('op_origem')

    reprogramar_form = ReprogramarOrdemForm(initial={'nova_data': op.data_programada})
    cancelar_form = CancelarOrdemForm()
    apontar_form = ApontamentoTurnoForm()

    context = {
        'op': op,
        'apontamentos': apontamentos,
        'historicos': historicos,
        'saldos_gerados': saldos_gerados,
        'saldos_incorporados': saldos_incorporados,
        'reprogramar_form': reprogramar_form,
        'cancelar_form': cancelar_form,
        'apontar_form': apontar_form,
        'is_lider': user_is_lider_bladder(request.user),
    }
    return render(request, 'bladder/ordem_detalhe.html', context)


@login_required
@lider_bladder_required
def ordem_reprogramar(request, pk):
    """Reprogramação de data da OP pelo líder com motivo obrigatório."""
    if request.method == 'POST':
        form = ReprogramarOrdemForm(request.POST)
        if form.is_valid():
            try:
                nova_data = form.cleaned_data['nova_data']
                motivo = form.cleaned_data['motivo']
                op = reprogramar_ordem_producao(pk, nova_data, motivo, request.user)
                messages.success(request, f"Ordem {op.numero_ordem} reprogramada para {nova_data.strftime('%d/%m/%Y')}!")
            except Exception as e:
                messages.error(request, f"Falha na reprogramação: {str(e)}")
        else:
            messages.error(request, "Dados inválidos para reprogramação.")
    return redirect('bladder:ordem_detalhe', pk=pk)


@login_required
@lider_bladder_required
def ordem_cancelar(request, pk):
    """Cancelamento lógico da OP pelo líder com motivo obrigatório."""
    if request.method == 'POST':
        form = CancelarOrdemForm(request.POST)
        if form.is_valid():
            try:
                motivo = form.cleaned_data['motivo']
                op = cancelar_ordem_producao(pk, motivo, request.user)
                messages.warning(request, f"Ordem {op.numero_ordem} cancelada com sucesso.")
            except Exception as e:
                messages.error(request, f"Falha ao cancelar: {str(e)}")
        else:
            messages.error(request, "Dados inválidos para cancelamento.")
    return redirect('bladder:ordem_detalhe', pk=pk)


@login_required
@lider_bladder_required
def encerrar_parcial_view(request, pk):
    """Encerramento manual ou de fim de turno pelo líder gerando saldo no Ledger."""
    if request.method == 'POST':
        motivo = request.POST.get('motivo_encerramento', '').strip()
        try:
            op = encerrar_op_parcial_ou_total(pk, request.user, motivo)
            if op.status == 'PARCIAL':
                saldo = op.quantidade_planejada - op.quantidade_realizada
                messages.info(request, f"Ordem {op.numero_ordem} encerrada como PARCIAL. Saldo de {saldo} un registrado no Ledger!")
            else:
                messages.success(request, f"Ordem {op.numero_ordem} concluída com sucesso!")
        except Exception as e:
            messages.error(request, f"Erro ao encerrar: {str(e)}")
    return redirect('bladder:ordem_detalhe', pk=pk)


@login_required
@operador_ou_lider_bladder_required
def operador_iniciar(request, pk):
    """Início de atividade pelo operador no chão de fábrica."""
    if request.method == 'POST':
        op = get_object_or_404(OrdemProducaoBladder, pk=pk)
        if op.status == 'PENDENTE':
            registrar_ou_atualizar_apontamento(
                ordem_id=op.pk,
                operador=request.user,
                quantidade=0,
                situacao='EM_ANDAMENTO',
                observacoes="Início de atividade apontado no chão de fábrica."
            )
            messages.success(request, f"Produção da OP {op.numero_ordem} INICIADA com sucesso!")
        elif op.status == 'EM_EXECUCAO':
            messages.info(request, f"A OP {op.numero_ordem} já está em execução.")
    return redirect('bladder:operador')


@login_required
@operador_ou_lider_bladder_required
def operador_apontar(request, pk):
    """Apontamento de quantidade produzida pelo operador."""
    if request.method == 'POST':
        form = ApontamentoTurnoForm(request.POST)
        if form.is_valid():
            try:
                op = get_object_or_404(OrdemProducaoBladder, pk=pk)
                qtd = form.cleaned_data['quantidade_realizada']
                situacao = form.cleaned_data['situacao']
                motivo = form.cleaned_data.get('motivo_desvio', '')
                obs = form.cleaned_data.get('observacoes', '')

                registrar_ou_atualizar_apontamento(
                    ordem_id=op.pk,
                    operador=request.user,
                    quantidade=qtd,
                    situacao=situacao,
                    motivo_desvio=motivo,
                    observacoes=obs
                )

                messages.success(request, f"Apontamento de {qtd} un registrado para a OP {op.numero_ordem}!")
            except Exception as e:
                messages.error(request, f"Erro ao registrar apontamento: {str(e)}")
        else:
            for erro in form.errors.values():
                messages.error(request, erro[0])

    if request.GET.get('origem') == 'detalhe':
        return redirect('bladder:ordem_detalhe', pk=pk)
    return redirect('bladder:operador')


@login_required
@operador_ou_lider_bladder_required
def corrigir_apontamento(request, pk):
    """Correção auditável de apontamento pelo operador ou líder."""
    apontamento = get_object_or_404(ApontamentoTurnoBladder, pk=pk)
    hoje = timezone.localdate()

    # Validação de permissão: operador só corrige apontamento próprio no mesmo turno do dia
    is_lider = user_is_lider_bladder(request.user)
    if not is_lider:
        if apontamento.operador != request.user:
            messages.error(request, "Você só pode corrigir os seus próprios apontamentos.")
            return redirect('bladder:operador')
        if apontamento.data_turno != hoje:
            messages.error(request, "O turno deste apontamento já encerrou. Correção permitida apenas pela Liderança.")
            return redirect('bladder:operador')

    if request.method == 'POST':
        form = CorrecaoApontamentoForm(request.POST)
        if form.is_valid():
            try:
                nova_qtd = form.cleaned_data['quantidade_realizada']
                motivo = form.cleaned_data['motivo_correcao']
                corrigir_apontamento_operador_ou_lider(
                    apontamento_id=apontamento.pk,
                    nova_quantidade=nova_qtd,
                    motivo_correcao=motivo,
                    usuario=request.user
                )
                messages.success(request, f"Apontamento corrigido com sucesso para {nova_qtd} un.")
                return redirect('bladder:ordem_detalhe', pk=apontamento.ordem_id)
            except Exception as e:
                messages.error(request, f"Erro ao corrigir apontamento: {str(e)}")
    else:
        form = CorrecaoApontamentoForm(initial={'quantidade_realizada': apontamento.quantidade_realizada})

    context = {
        'apontamento': apontamento,
        'form': form,
    }
    return render(request, 'bladder/corrigir_apontamento.html', context)


@login_required
@operador_ou_lider_bladder_required
def api_saldo_produto(request, produto_id):
    """API para consultar saldo pendente disponível no Ledger para um produto com detalhes para decisão do Líder."""
    saldos = obter_saldos_pendentes_produto(produto_id).select_related(
        'op_origem', 'op_origem__processo', 'op_origem__processo__maquina', 'produto'
    )
    total_saldo = sum(s.quantidade for s in saldos)
    pendencias = []
    for s in saldos:
        op = s.op_origem
        pendencias.append({
            'id': s.id,
            'op_origem_numero': op.numero_ordem,
            'data': op.data_programada.strftime('%d/%m/%Y'),
            'modelo': s.produto.codigo,
            'processo_equipamento': str(op.processo),
            'quantidade_programada': op.quantidade_planejada,
            'quantidade_realizada': op.quantidade_realizada,
            'saldo': s.quantidade,
            'turma_origem': op.turma_responsavel_fechamento,
            'motivo': op.motivo_pendencia,
        })
    return JsonResponse({
        'produto_id': produto_id,
        'total_saldo_pendente': total_saldo,
        'quantidade_registros': len(pendencias),
        'pendencias': pendencias
    })


@login_required
@operador_ou_lider_bladder_required
def api_escala_dia(request):
    """API para consultar a escala e a turma prevista para uma determinada data."""
    data_param = request.GET.get('data')
    if data_param:
        try:
            data = datetime.datetime.strptime(data_param, '%Y-%m-%d').date()
        except ValueError:
            data = timezone.localdate()
    else:
        data = timezone.localdate()

    turma, is_ajuste, motivo = calcular_turma_do_dia(data)
    turma_display = "Turma A" if turma == 'TURMA_A' else ("Turma B" if turma == 'TURMA_B' else turma)
    return JsonResponse({
        'data': data.strftime('%Y-%m-%d'),
        'turma': turma,
        'turma_display': turma_display,
        'is_ajuste': is_ajuste,
        'motivo': motivo,
    })


def formatar_percentual(valor):
    """
    Formata percentual de cumprimento conforme regra obrigatória da SPEC:
    percentual_cumprimento = realizado / programado * 100
    - 40 / 40 -> 100%
    - 27 / 30 -> 90%
    - 59 / 113 -> 52,21%
    """
    if valor is None:
        return "-"
    if round(valor, 2) == int(round(valor, 2)):
        return f"{int(round(valor, 2))}%"
    return f"{valor:.2f}%".replace('.', ',')


@login_required
@lider_bladder_required
def relatorios(request):
    """
    Relatório mensal e indicadores consolidados de produção por fechamento de turno.
    Permite filtros por Mês, Ano, Turno (Todos, Turma A, Turma B) e Modelo de Bladder.
    Fonte de dados realizada: Fechamentos de Turno (auditados).
    """
    hoje = timezone.localdate()
    try:
        mes = int(request.GET.get('mes', hoje.month))
    except (ValueError, TypeError):
        mes = hoje.month
    try:
        ano = int(request.GET.get('ano', hoje.year))
    except (ValueError, TypeError):
        ano = hoje.year

    turno_filtro = request.GET.get('turno', 'TODOS').strip().upper()
    if turno_filtro not in ('TURMA_A', 'TURMA_B'):
        turno_filtro = 'TODOS'

    modelo_filtro = request.GET.get('modelo', 'TODOS').strip()

    primeiro_dia = datetime.date(ano, mes, 1)
    ultimo_dia = datetime.date(ano, mes, calendar.monthrange(ano, mes)[1])

    modelos_disponiveis = ProdutoBladder.objects.filter(ativo=True).order_by('codigo')

    ordens = OrdemProducaoBladder.objects.filter(
        data_programada__range=[primeiro_dia, ultimo_dia]
    ).exclude(status='CANCELADA').select_related('processo', 'produto', 'criado_por')

    if turno_filtro in ('TURMA_A', 'TURMA_B'):
        ordens = ordens.filter(turma_prevista=turno_filtro)

    if modelo_filtro and modelo_filtro != 'TODOS':
        if modelo_filtro.isdigit():
            ordens = ordens.filter(produto_id=int(modelo_filtro))
        else:
            ordens = ordens.filter(produto__codigo=modelo_filtro)

    total_ops = ordens.count()
    total_demanda_nova = ordens.aggregate(s=Sum('quantidade_nova'))['s'] or 0
    total_saldo_incorporado = ordens.aggregate(s=Sum('saldo_anterior_incorporado'))['s'] or 0
    total_programado = ordens.aggregate(s=Sum('quantidade_planejada'))['s'] or 0

    # Fechamentos de turno no período
    fechamentos_mes = FechamentoTurnoBladder.objects.filter(
        data_turno__range=[primeiro_dia, ultimo_dia],
        status='CONCLUIDO'
    )
    if turno_filtro in ('TURMA_A', 'TURMA_B'):
        fechamentos_mes = fechamentos_mes.filter(turma=turno_filtro)

    itens_fech_mes = ItemFechamentoTurnoBladder.objects.filter(
        fechamento__in=fechamentos_mes
    ).select_related('fechamento', 'ordem', 'ordem__produto', 'fechamento__operador')

    if modelo_filtro and modelo_filtro != 'TODOS':
        if modelo_filtro.isdigit():
            itens_fech_mes = itens_fech_mes.filter(ordem__produto_id=int(modelo_filtro))
        else:
            itens_fech_mes = itens_fech_mes.filter(ordem__produto__codigo=modelo_filtro)

    # 2.2 RESUMO DO MÊS
    if itens_fech_mes.exists():
        total_realizado = itens_fech_mes.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_total = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_fech_mes)
        pendencia_total = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_fech_mes)
    else:
        total_realizado = ordens.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_total = sum(op.excedente for op in ordens)
        pendencia_total = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ordens if op.status in ['CONCLUIDA', 'PARCIAL'])

    saldo_pendente_total = pendencia_total
    diferenca_liquida = total_realizado - total_programado
    percentual_geral = (total_realizado / total_programado * 100) if total_programado > 0 else None
    percentual_geral_display = formatar_percentual(percentual_geral)

    # 2.3 PRODUÇÃO POR TURNO
    # Turma A
    prog_turma_a = ordens.filter(turma_prevista='TURMA_A').aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    if itens_fech_mes.filter(fechamento__turma='TURMA_A').exists():
        itens_a = itens_fech_mes.filter(fechamento__turma='TURMA_A')
        prod_turma_a = itens_a.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_turma_a = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_a)
        pendencia_turma_a = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_a)
    else:
        ops_a = ordens.filter(turma_prevista='TURMA_A')
        prod_turma_a = ops_a.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_turma_a = sum(op.excedente for op in ops_a)
        pendencia_turma_a = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_a if op.status in ['CONCLUIDA', 'PARCIAL'])
    saldo_turma_a = pendencia_turma_a
    diferenca_turma_a = prod_turma_a - prog_turma_a
    pct_turma_a = (prod_turma_a / prog_turma_a * 100) if prog_turma_a > 0 else None
    pct_turma_a_display = formatar_percentual(pct_turma_a)

    # Turma B
    prog_turma_b = ordens.filter(turma_prevista='TURMA_B').aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    if itens_fech_mes.filter(fechamento__turma='TURMA_B').exists():
        itens_b = itens_fech_mes.filter(fechamento__turma='TURMA_B')
        prod_turma_b = itens_b.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_turma_b = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_b)
        pendencia_turma_b = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_b)
    else:
        ops_b = ordens.filter(turma_prevista='TURMA_B')
        prod_turma_b = ops_b.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_turma_b = sum(op.excedente for op in ops_b)
        pendencia_turma_b = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_b if op.status in ['CONCLUIDA', 'PARCIAL'])
    saldo_turma_b = pendencia_turma_b
    diferenca_turma_b = prod_turma_b - prog_turma_b
    pct_turma_b = (prod_turma_b / prog_turma_b * 100) if prog_turma_b > 0 else None
    pct_turma_b_display = formatar_percentual(pct_turma_b)

    # Total Consolidado das Turmas
    prog_total_turmas = prog_turma_a + prog_turma_b
    prod_total_turmas = prod_turma_a + prod_turma_b
    excedente_total_turmas = excedente_turma_a + excedente_turma_b
    pendencia_total_turmas = pendencia_turma_a + pendencia_turma_b
    saldo_total_turmas = pendencia_total_turmas
    diferenca_total_turmas = prod_total_turmas - prog_total_turmas
    pct_total_turmas = (prod_total_turmas / prog_total_turmas * 100) if prog_total_turmas > 0 else None
    pct_total_turmas_display = formatar_percentual(pct_total_turmas)

    # 2.4 PRODUÇÃO POR DIA
    dias_prod = {}
    if itens_fech_mes.exists():
        for item in itens_fech_mes:
            dt = item.fechamento.data_turno
            turma = item.fechamento.turma
            qtd = item.quantidade_realizada
            if dt not in dias_prod:
                dias_prod[dt] = {'data': dt, 'turma_a': 0, 'turma_b': 0, 'total': 0}
            if turma == 'TURMA_A':
                dias_prod[dt]['turma_a'] += qtd
            elif turma == 'TURMA_B':
                dias_prod[dt]['turma_b'] += qtd
            dias_prod[dt]['total'] += qtd
    else:
        for op in ordens:
            if op.quantidade_realizada > 0:
                dt = op.data_programada
                turma = op.turma_prevista
                qtd = op.quantidade_realizada
                if dt not in dias_prod:
                    dias_prod[dt] = {'data': dt, 'turma_a': 0, 'turma_b': 0, 'total': 0}
                if turma == 'TURMA_A':
                    dias_prod[dt]['turma_a'] += qtd
                elif turma == 'TURMA_B':
                    dias_prod[dt]['turma_b'] += qtd
                dias_prod[dt]['total'] += qtd

    producao_diaria = sorted(dias_prod.values(), key=lambda x: x['data'])
    acumulado_periodo = sum(d['total'] for d in producao_diaria)

    # 2.5 PRODUÇÃO POR MODELO
    produtos_agrupados = ordens.values(
        'produto__id', 'produto__codigo', 'produto__descricao'
    ).annotate(
        qtd_ops=Count('id'),
        programado=Sum('quantidade_planejada'),
        demanda_nova=Sum('quantidade_nova'),
        saldo_inc=Sum('saldo_anterior_incorporado'),
    ).order_by('produto__codigo')

    prod_por_produto = []
    for p in produtos_agrupados:
        p_id = p['produto__id']
        prog = p['programado'] or 0
        ops_p = ordens.filter(produto_id=p_id)
        itens_p = itens_fech_mes.filter(ordem__produto_id=p_id) if itens_fech_mes.exists() else None

        if itens_p and itens_p.exists():
            real = itens_p.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
            excedente = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_p)
            pendencia = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_p)
        else:
            real = ops_p.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
            excedente = sum(op.excedente for op in ops_p)
            pendencia = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_p if op.status in ['CONCLUIDA', 'PARCIAL'])

        saldo = pendencia
        diferenca = real - prog
        pct = (real / prog * 100) if prog > 0 else None

        prod_obj = ProdutoBladder.objects.filter(id=p_id).first()
        medida = prod_obj.medida if prod_obj else ""

        prod_por_produto.append({
            'codigo': p['produto__codigo'],
            'medida': medida,
            'codigo_com_medida': f"{p['produto__codigo']} — {medida}" if medida else p['produto__codigo'],
            'descricao': p['produto__descricao'],
            'quantidade_ops': p['qtd_ops'],
            'programado': prog,
            'demanda_nova': p['demanda_nova'] or 0,
            'saldo_incorporado': p['saldo_inc'] or 0,
            'realizado': real,
            'excedente': excedente,
            'pendencia': pendencia,
            'saldo': saldo,
            'diferenca': diferenca,
            'percentual': round(pct, 2) if pct is not None else 0.0,
            'percentual_display': formatar_percentual(pct),
        })
    prod_por_produto.sort(key=lambda x: x['realizado'], reverse=True)

    # 2.6 PENDÊNCIAS DO PERÍODO
    pendencias_periodo = []
    if itens_fech_mes.exists():
        for item in itens_fech_mes.filter(saldo_gerado__gt=0).select_related(
            'fechamento', 'ordem', 'ordem__produto', 'ordem__processo', 'ordem__processo__maquina', 'fechamento__operador'
        ).prefetch_related('categorias').order_by('-fechamento__data_turno', 'ordem__numero_ordem'):
            pendencias_periodo.append({
                'data_fechamento': item.fechamento.data_turno,
                'data_hora_fechamento': item.fechamento.data_hora_fechamento,
                'turno': item.fechamento.get_turma_display(),
                'turma': item.fechamento.get_turma_display(),
                'operador': item.fechamento.operador.get_full_name() or item.fechamento.operador.username,
                'operador_fechamento': item.fechamento.operador.get_full_name() or item.fechamento.operador.username,
                'numero_ordem': item.ordem.numero_ordem,
                'modelo': item.ordem.produto.codigo,
                'medida': item.ordem.produto.medida,
                'codigo_com_medida': item.ordem.produto.codigo_com_medida,
                'modelo_descricao': item.ordem.produto.descricao,
                'processo': item.ordem.processo.nome if item.ordem.processo else "-",
                'maquina': item.ordem.processo.maquina.nome if (item.ordem.processo and item.ordem.processo.maquina) else "-",
                'quantidade_programada': item.quantidade_programada,
                'quantidade_realizada': item.quantidade_realizada,
                'saldo': item.saldo_gerado,
                'pendencia': item.saldo_gerado,
                'excedente': item.excedente,
                'diferenca': item.diferenca,
                'categorias': item.categorias_display,
                'descricao_desvio': item.descricao_desvio or item.motivo_outro or "",
                'motivo': item.motivo_completo or "Sem justificativa informada",
            })
    else:
        saldos_fallback = SaldoPendenteBladder.objects.filter(
            op_origem__in=ordens,
            quantidade__gt=0
        ).select_related('op_origem', 'produto', 'fechamento', 'fechamento__operador', 'op_origem__processo', 'op_origem__processo__maquina')
        for s in saldos_fallback:
            pendencias_periodo.append({
                'data_fechamento': s.fechamento.data_turno if s.fechamento else s.op_origem.data_programada,
                'data_hora_fechamento': s.fechamento.data_hora_fechamento if s.fechamento else s.created_at,
                'turno': s.fechamento.get_turma_display() if s.fechamento else s.op_origem.turma_responsavel_fechamento,
                'turma': s.fechamento.get_turma_display() if s.fechamento else s.op_origem.turma_responsavel_fechamento,
                'operador': (s.fechamento.operador.get_full_name() or s.fechamento.operador.username) if (s.fechamento and s.fechamento.operador) else "-",
                'operador_fechamento': (s.fechamento.operador.get_full_name() or s.fechamento.operador.username) if (s.fechamento and s.fechamento.operador) else "-",
                'numero_ordem': s.op_origem.numero_ordem,
                'modelo': s.produto.codigo,
                'medida': s.produto.medida,
                'codigo_com_medida': s.produto.codigo_com_medida,
                'modelo_descricao': s.produto.descricao,
                'processo': s.op_origem.processo.nome if s.op_origem.processo else "-",
                'maquina': s.op_origem.processo.maquina.nome if (s.op_origem.processo and s.op_origem.processo.maquina) else "-",
                'quantidade_programada': s.op_origem.quantidade_planejada,
                'quantidade_realizada': s.op_origem.quantidade_realizada,
                'saldo': s.quantidade,
                'pendencia': s.quantidade,
                'excedente': s.op_origem.excedente,
                'diferenca': s.op_origem.diferenca,
                'categorias': s.op_origem.motivo_pendencia or "-",
                'descricao_desvio': s.op_origem.motivo_pendencia or "-",
                'motivo': s.op_origem.motivo_pendencia,
            })

    context = {
        'mes': mes,
        'ano': ano,
        'mes_nome': calendar.month_name[mes],
        'turno_filtro': turno_filtro,
        'modelo_filtro': modelo_filtro,
        'modelos_disponiveis': modelos_disponiveis,
        'total_ops': total_ops,
        'total_demanda_nova': total_demanda_nova,
        'total_saldo_incorporado': total_saldo_incorporado,
        'total_programado': total_programado,
        'total_realizado': total_realizado,
        'excedente_total': excedente_total,
        'pendencia_total': pendencia_total,
        'diferenca_liquida': diferenca_liquida,
        'saldo_pendente_total': saldo_pendente_total,
        'percentual_geral': round(percentual_geral, 2) if percentual_geral is not None else 0.0,
        'percentual_geral_display': percentual_geral_display,
        # Turno A
        'prog_turma_a': prog_turma_a,
        'prod_turma_a': prod_turma_a,
        'excedente_turma_a': excedente_turma_a,
        'pendencia_turma_a': pendencia_turma_a,
        'saldo_turma_a': saldo_turma_a,
        'diferenca_turma_a': diferenca_turma_a,
        'pct_turma_a_display': pct_turma_a_display,
        # Turno B
        'prog_turma_b': prog_turma_b,
        'prod_turma_b': prod_turma_b,
        'excedente_turma_b': excedente_turma_b,
        'pendencia_turma_b': pendencia_turma_b,
        'saldo_turma_b': saldo_turma_b,
        'diferenca_turma_b': diferenca_turma_b,
        'pct_turma_b_display': pct_turma_b_display,
        # Total das Turmas
        'prog_total_turmas': prog_total_turmas,
        'prod_total_turmas': prod_total_turmas,
        'excedente_total_turmas': excedente_total_turmas,
        'pendencia_total_turmas': pendencia_total_turmas,
        'saldo_total_turmas': saldo_total_turmas,
        'diferenca_total_turmas': diferenca_total_turmas,
        'pct_total_turmas_display': pct_total_turmas_display,
        # Produção Diária
        'producao_diaria': producao_diaria,
        'acumulado_periodo': acumulado_periodo,
        # Produção por Modelo
        'prod_por_produto': prod_por_produto,
        # Pendências do Período
        'pendencias_periodo': pendencias_periodo,
    }
    return render(request, 'bladder/relatorios.html', context)


@login_required
@lider_bladder_required
def relatorios_exportar_excel(request):
    """Exportação sanitizada em Excel do relatório mensal via openpyxl alinhada ao relatório visual."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    hoje = timezone.localdate()
    try:
        mes = int(request.GET.get('mes', hoje.month))
    except (ValueError, TypeError):
        mes = hoje.month
    try:
        ano = int(request.GET.get('ano', hoje.year))
    except (ValueError, TypeError):
        ano = hoje.year

    turno_filtro = request.GET.get('turno', 'TODOS').strip().upper()
    if turno_filtro not in ('TURMA_A', 'TURMA_B'):
        turno_filtro = 'TODOS'

    modelo_filtro = request.GET.get('modelo', 'TODOS').strip()

    primeiro_dia = datetime.date(ano, mes, 1)
    ultimo_dia = datetime.date(ano, mes, calendar.monthrange(ano, mes)[1])

    ordens = OrdemProducaoBladder.objects.filter(
        data_programada__range=[primeiro_dia, ultimo_dia]
    ).exclude(status='CANCELADA').select_related('processo', 'produto', 'criado_por').order_by('data_programada', 'numero_ordem')

    if turno_filtro in ('TURMA_A', 'TURMA_B'):
        ordens = ordens.filter(turma_prevista=turno_filtro)

    if modelo_filtro and modelo_filtro != 'TODOS':
        if modelo_filtro.isdigit():
            ordens = ordens.filter(produto_id=int(modelo_filtro))
        else:
            ordens = ordens.filter(produto__codigo=modelo_filtro)

    fechamentos_mes = FechamentoTurnoBladder.objects.filter(
        data_turno__range=[primeiro_dia, ultimo_dia],
        status='CONCLUIDO'
    )
    if turno_filtro in ('TURMA_A', 'TURMA_B'):
        fechamentos_mes = fechamentos_mes.filter(turma=turno_filtro)

    itens_fech_mes = ItemFechamentoTurnoBladder.objects.filter(
        fechamento__in=fechamentos_mes
    ).select_related('fechamento', 'ordem', 'ordem__produto', 'fechamento__operador')

    if modelo_filtro and modelo_filtro != 'TODOS':
        if modelo_filtro.isdigit():
            itens_fech_mes = itens_fech_mes.filter(ordem__produto_id=int(modelo_filtro))
        else:
            itens_fech_mes = itens_fech_mes.filter(ordem__produto__codigo=modelo_filtro)

    wb = openpyxl.Workbook()

    # Sanitização contra formula injection
    def sanitize(v):
        if isinstance(v, str) and v.startswith(('=', '+', '-', '@')):
            return "'" + v
        return v

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Aba 1: Detalhe das OPs
    ws1 = wb.active
    ws1.title = f"Ordens {mes:02d}-{ano}"

    headers_ws1 = [
        "Número OP", "Data", "Turma", "Processo", "Código Bladder", "Medida", "Descrição",
        "Demanda Nova", "Saldo Inc.", "Total Programado", "Realizado", "Excedente", "Pendência", "Diferença",
        "% Cumprimento", "Status"
    ]

    for col_num, h_text in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=1, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    for r_idx, op in enumerate(ordens, 2):
        pct = (op.quantidade_realizada / op.quantidade_planejada * 100) if op.quantidade_planejada > 0 else None
        row_values = [
            sanitize(op.numero_ordem),
            op.data_programada.strftime('%d/%m/%Y'),
            sanitize(op.get_turma_prevista_display()),
            sanitize(op.processo.nome if op.processo else ""),
            sanitize(op.produto.codigo if op.produto else ""),
            sanitize(op.produto.medida if op.produto else ""),
            sanitize(op.produto.descricao if op.produto else ""),
            op.quantidade_nova,
            op.saldo_anterior_incorporado,
            op.quantidade_planejada,
            op.quantidade_realizada,
            op.excedente,
            op.saldo_gerado,
            op.diferenca,
            formatar_percentual(pct),
            sanitize(op.status_visual),
        ]

        for col_num, val in enumerate(row_values, 1):
            cell = ws1.cell(row=r_idx, column=col_num, value=val)
            if isinstance(val, int):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Aba 2: Produção por Modelo
    ws2 = wb.create_sheet(title="Por Modelo")
    headers_ws2 = [
        "Código", "Medida", "Descrição", "Qtd OPs", "Demanda Nova", "Saldo Inc.",
        "Total Programado", "Realizado", "Excedente", "Pendência", "Diferença", "% Cumprimento"
    ]

    for col_num, h_text in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=1, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    produtos_agrupados = ordens.values(
        'produto__id', 'produto__codigo', 'produto__descricao'
    ).annotate(
        qtd_ops=Count('id'),
        programado=Sum('quantidade_planejada'),
        demanda_nova=Sum('quantidade_nova'),
        saldo_inc=Sum('saldo_anterior_incorporado'),
    ).order_by('produto__codigo')

    for r_idx, p in enumerate(produtos_agrupados, 2):
        p_id = p['produto__id']
        prog = p['programado'] or 0
        ops_p = ordens.filter(produto_id=p_id)
        itens_p = itens_fech_mes.filter(ordem__produto_id=p_id) if itens_fech_mes.exists() else None

        if itens_p and itens_p.exists():
            real = itens_p.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
            excedente = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_p)
            pendencia = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_p)
        else:
            real = ops_p.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
            excedente = sum(op.excedente for op in ops_p)
            pendencia = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_p if op.status in ['CONCLUIDA', 'PARCIAL'])

        diferenca = real - prog
        pct = (real / prog * 100) if prog > 0 else None
        prod_obj = ProdutoBladder.objects.filter(id=p_id).first()
        medida = prod_obj.medida if prod_obj else ""

        row_p = [
            sanitize(p['produto__codigo']),
            sanitize(medida),
            sanitize(p['produto__descricao']),
            p['qtd_ops'],
            p['demanda_nova'] or 0,
            p['saldo_inc'] or 0,
            prog,
            real,
            excedente,
            pendencia,
            diferenca,
            formatar_percentual(pct),
        ]
        for col_num, val in enumerate(row_p, 1):
            cell = ws2.cell(row=r_idx, column=col_num, value=val)
            if isinstance(val, int):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Aba 3: Resumo Mensal e Desempenho por Turno / Turmas
    ws3 = wb.create_sheet(title="Resumo e Turmas")
    ws3.cell(row=1, column=1, value="Indicador").font = header_font
    ws3.cell(row=1, column=1).fill = header_fill
    ws3.cell(row=1, column=2, value="Valor").font = header_font
    ws3.cell(row=1, column=2).fill = header_fill

    prog_a = ordens.filter(turma_prevista='TURMA_A').aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    prog_b = ordens.filter(turma_prevista='TURMA_B').aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    if itens_fech_mes.filter(fechamento__turma='TURMA_A').exists():
        itens_a = itens_fech_mes.filter(fechamento__turma='TURMA_A')
        prod_a = itens_a.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_a = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_a)
        pendencia_a = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_a)
    else:
        ops_a = ordens.filter(turma_prevista='TURMA_A')
        prod_a = ops_a.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_a = sum(op.excedente for op in ops_a)
        pendencia_a = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_a if op.status in ['CONCLUIDA', 'PARCIAL'])

    if itens_fech_mes.filter(fechamento__turma='TURMA_B').exists():
        itens_b = itens_fech_mes.filter(fechamento__turma='TURMA_B')
        prod_b = itens_b.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_b = sum(max(0, it.quantidade_realizada - it.quantidade_programada) for it in itens_b)
        pendencia_b = sum(max(0, it.quantidade_programada - it.quantidade_realizada) for it in itens_b)
    else:
        ops_b = ordens.filter(turma_prevista='TURMA_B')
        prod_b = ops_b.aggregate(s=Sum('quantidade_realizada'))['s'] or 0
        excedente_b = sum(op.excedente for op in ops_b)
        pendencia_b = sum(max(0, op.quantidade_planejada - op.quantidade_realizada) for op in ops_b if op.status in ['CONCLUIDA', 'PARCIAL'])

    tot_prog = ordens.aggregate(s=Sum('quantidade_planejada'))['s'] or 0
    tot_real = prod_a + prod_b if itens_fech_mes.exists() else (ordens.aggregate(s=Sum('quantidade_realizada'))['s'] or 0)
    tot_excedente = excedente_a + excedente_b
    tot_pendencia = pendencia_a + pendencia_b
    tot_diferenca = tot_real - tot_prog
    pct_geral = (tot_real / tot_prog * 100) if tot_prog > 0 else None

    resumo_linhas = [
        ("Mês / Ano", f"{mes:02d}/{ano}"),
        ("Filtro Turno", turno_filtro),
        ("Filtro Modelo", modelo_filtro),
        ("Total de OPs no Período", ordens.count()),
        ("Total Programado", f"{tot_prog} un"),
        ("Total Realizado", f"{tot_real} un"),
        ("Excedente Total", f"{tot_excedente} un"),
        ("Pendência Total", f"{tot_pendencia} un"),
        ("Diferença Líquida Geral", f"{tot_diferenca} un"),
        ("Percentual de Cumprimento Geral", formatar_percentual(pct_geral)),
        ("Turma A - Programado", f"{prog_a} un"),
        ("Turma A - Realizado", f"{prod_a} un"),
        ("Turma A - Excedente", f"{excedente_a} un"),
        ("Turma A - Pendência", f"{pendencia_a} un"),
        ("Turma A - Diferença", f"{prod_a - prog_a} un"),
        ("Turma A - % Cumprimento", formatar_percentual((prod_a / prog_a * 100) if prog_a > 0 else None)),
        ("Turma B - Programado", f"{prog_b} un"),
        ("Turma B - Realizado", f"{prod_b} un"),
        ("Turma B - Excedente", f"{excedente_b} un"),
        ("Turma B - Pendência", f"{pendencia_b} un"),
        ("Turma B - Diferença", f"{prod_b - prog_b} un"),
        ("Turma B - % Cumprimento", formatar_percentual((prod_b / prog_b * 100) if prog_b > 0 else None)),
        ("Total Consolidado - Programado", f"{prog_a + prog_b} un"),
        ("Total Consolidado - Realizado", f"{prod_a + prod_b} un"),
        ("Total Consolidado - Excedente", f"{tot_excedente} un"),
        ("Total Consolidado - Pendência", f"{tot_pendencia} un"),
        ("Total Consolidado - Diferença", f"{tot_diferenca} un"),
        ("Total Consolidado - % Cumprimento", formatar_percentual(((prod_a + prod_b) / (prog_a + prog_b) * 100) if (prog_a + prog_b) > 0 else None)),
    ]

    for r_idx, (ind, val) in enumerate(resumo_linhas, 2):
        ws3.cell(row=r_idx, column=1, value=ind).alignment = align_left
        ws3.cell(row=r_idx, column=2, value=val).alignment = align_center

    ws3.column_dimensions['A'].width = 36
    ws3.column_dimensions['B'].width = 24

    # Aba 4: Produção Diária
    ws4 = wb.create_sheet(title="Produção Diária")
    headers_ws4 = ["Data", "Turma A", "Turma B", "Total do Dia"]
    for col_num, h_text in enumerate(headers_ws4, 1):
        cell = ws4.cell(row=1, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    dias_prod = {}
    if itens_fech_mes.exists():
        for item in itens_fech_mes:
            dt = item.fechamento.data_turno
            turma = item.fechamento.turma
            qtd = item.quantidade_realizada
            if dt not in dias_prod:
                dias_prod[dt] = {'data': dt, 'turma_a': 0, 'turma_b': 0, 'total': 0}
            if turma == 'TURMA_A':
                dias_prod[dt]['turma_a'] += qtd
            elif turma == 'TURMA_B':
                dias_prod[dt]['turma_b'] += qtd
            dias_prod[dt]['total'] += qtd
    else:
        for op in ordens:
            if op.quantidade_realizada > 0:
                dt = op.data_programada
                turma = op.turma_prevista
                qtd = op.quantidade_realizada
                if dt not in dias_prod:
                    dias_prod[dt] = {'data': dt, 'turma_a': 0, 'turma_b': 0, 'total': 0}
                if turma == 'TURMA_A':
                    dias_prod[dt]['turma_a'] += qtd
                elif turma == 'TURMA_B':
                    dias_prod[dt]['turma_b'] += qtd
                dias_prod[dt]['total'] += qtd

    diaria_sorted = sorted(dias_prod.values(), key=lambda x: x['data'])
    acumulado_total = sum(d['total'] for d in diaria_sorted)

    for r_idx, d in enumerate(diaria_sorted, 2):
        ws4.cell(row=r_idx, column=1, value=d['data'].strftime('%d/%m/%Y')).alignment = align_center
        ws4.cell(row=r_idx, column=2, value=d['turma_a'] if d['turma_a'] > 0 else "-").alignment = align_center
        ws4.cell(row=r_idx, column=3, value=d['turma_b'] if d['turma_b'] > 0 else "-").alignment = align_center
        ws4.cell(row=r_idx, column=4, value=d['total']).alignment = align_center

    row_acum = len(diaria_sorted) + 2
    ws4.cell(row=row_acum, column=1, value="Acumulado do Período").font = header_font
    ws4.cell(row=row_acum, column=4, value=acumulado_total).font = header_font
    ws4.cell(row=row_acum, column=4).alignment = align_center

    for col in ws4.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = max(max_len + 3, 16)

    # Aba 5: Pendências do Período
    ws5 = wb.create_sheet(title="Pendências do Período")
    headers_ws5 = [
        "Data Fechamento", "Turno", "Número OP", "Processo / Equipamento", "Código Bladder", "Medida",
        "Programado", "Realizado", "Pendência", "Categorias do Desvio", "O que aconteceu? (Descrição)"
    ]
    for col_num, h_text in enumerate(headers_ws5, 1):
        cell = ws5.cell(row=1, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    pendencias = []
    if itens_fech_mes.exists():
        for item in itens_fech_mes.filter(saldo_gerado__gt=0).select_related(
            'fechamento', 'ordem', 'ordem__produto', 'ordem__processo', 'ordem__processo__maquina'
        ).prefetch_related('categorias').order_by('fechamento__data_turno', 'ordem__numero_ordem'):
            pendencias.append([
                item.fechamento.data_turno.strftime('%d/%m/%Y'),
                sanitize(item.fechamento.get_turma_display()),
                sanitize(item.ordem.numero_ordem),
                sanitize(item.ordem.processo.nome if item.ordem.processo else ""),
                sanitize(item.ordem.produto.codigo if item.ordem.produto else ""),
                sanitize(item.ordem.produto.medida if item.ordem.produto else ""),
                item.quantidade_programada,
                item.quantidade_realizada,
                item.saldo_gerado,
                sanitize(item.categorias_display),
                sanitize(item.descricao_desvio or item.motivo_outro or ""),
            ])
    else:
        saldos_fallback = SaldoPendenteBladder.objects.filter(
            op_origem__in=ordens,
            quantidade__gt=0
        ).select_related('op_origem', 'produto', 'fechamento', 'op_origem__processo')
        for s in saldos_fallback:
            dt_f = s.fechamento.data_turno if s.fechamento else s.op_origem.data_programada
            turma_f = s.fechamento.get_turma_display() if s.fechamento else s.op_origem.turma_responsavel_fechamento
            proc_f = s.op_origem.processo.nome if (s.op_origem and s.op_origem.processo) else ""
            pendencias.append([
                dt_f.strftime('%d/%m/%Y'),
                sanitize(turma_f),
                sanitize(s.op_origem.numero_ordem),
                sanitize(proc_f),
                sanitize(s.produto.codigo if s.produto else ""),
                sanitize(s.produto.medida if s.produto else ""),
                s.op_origem.quantidade_planejada,
                s.op_origem.quantidade_realizada,
                s.quantidade,
                sanitize(s.op_origem.motivo_pendencia or "-"),
                sanitize(s.op_origem.motivo_pendencia or "-"),
            ])

    for r_idx, row_pend in enumerate(pendencias, 2):
        for col_num, val in enumerate(row_pend, 1):
            cell = ws5.cell(row=r_idx, column=col_num, value=val)
            if isinstance(val, int):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    for col in ws5.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws5.column_dimensions[col_letter].width = max(max_len + 3, 16)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="Fechamento_Bladder_{mes:02d}_{ano}.xlsx"'
    wb.save(response)
    return response


# ==============================================================================
# PASSAGEM DE TURNO — COMUNICAÇÃO OPERACIONAL ENTRE TURMAS A/B
# ==============================================================================

@login_required
@operador_ou_lider_bladder_required
def passagem_turno_criar(request):
    """
    Criação de mensagem de passagem de turno pelo operador ou líder.
    Suporta criação geral ou vinculada a uma OP / Processo.
    """
    if request.method != 'POST':
        return redirect('bladder:operador')

    form = MensagemPassagemTurnoForm(request.POST)
    next_url = request.POST.get('next_url') or request.META.get('HTTP_REFERER') or '/bladder/operador/'

    if form.is_valid():
        tipo = form.cleaned_data['tipo']
        categoria = form.cleaned_data['categoria']
        prioridade = form.cleaned_data['prioridade']
        msg_texto = form.cleaned_data['mensagem']

        op_id = form.cleaned_data.get('ordem_producao_id')
        proc_id = form.cleaned_data.get('processo_id')
        maq_id = form.cleaned_data.get('maquina_id')
        prod_id = form.cleaned_data.get('produto_id')

        op = OrdemProducaoBladder.objects.filter(pk=op_id).first() if op_id else None
        proc = ProcessoBladder.objects.filter(pk=proc_id).first() if proc_id else None
        from maintenance.models import Machine
        maq = Machine.objects.filter(pk=maq_id).first() if maq_id else None
        prod = ProdutoBladder.objects.filter(pk=prod_id).first() if prod_id else None

        try:
            msg = criar_mensagem_passagem_turno(
                autor=request.user,
                mensagem=msg_texto,
                tipo=tipo,
                categoria=categoria,
                prioridade=prioridade,
                ordem_producao=op,
                processo=proc,
                maquina=maq,
                produto=prod,
            )
            messages.success(
                request,
                f"Recado registrado com sucesso para o próximo turno ({msg.get_turma_destino_display()} - {msg.data_turno_destino.strftime('%d/%m')})!"
            )
        except (ValueError, PermissionError) as e:
            messages.error(request, f"Erro ao registrar recado: {str(e)}")
    else:
        for err_field, err_msgs in form.errors.items():
            for m in err_msgs:
                messages.error(request, f"{m}")

    return redirect(next_url)


@login_required
@operador_ou_lider_bladder_required
def passagem_turno_acao(request, pk):
    """
    Executa ações sobre uma mensagem de passagem de turno:
    - CIENTE (leitura confirmada pelo usuário)
    - RESOLVIDO (para mensagens de acompanhamento)
    - REPASSAR (cria nova mensagem de acompanhamento para o próximo turno)
    """
    if request.method != 'POST':
        return redirect('bladder:operador')

    msg = get_object_or_404(MensagemPassagemTurnoBladder, pk=pk)
    acao = request.POST.get('acao', '').strip().upper()
    observacao = request.POST.get('observacao', '').strip()
    next_url = request.POST.get('next_url') or request.META.get('HTTP_REFERER') or '/bladder/operador/'

    is_ajax = request.headers.get('x-requested-with') == 'XMLHttpRequest'

    try:
        if acao == 'CIENTE':
            acao_obj, created = registrar_ciencia_mensagem_turno(msg, request.user)
            if created:
                msg_txt = "Ciência registrada com sucesso!"
            else:
                msg_txt = "Você já havia registrado ciência nesta mensagem."
            if is_ajax:
                return JsonResponse({'success': True, 'acao': 'CIENTE', 'message': msg_txt})
            messages.success(request, msg_txt)

        elif acao == 'RESOLVIDO':
            resolver_mensagem_acompanhamento(msg, request.user, observacao)
            msg_txt = "Acompanhamento marcado como RESOLVIDO com sucesso!"
            if is_ajax:
                return JsonResponse({'success': True, 'acao': 'RESOLVIDO', 'message': msg_txt})
            messages.success(request, msg_txt)

        elif acao == 'REPASSAR':
            nova_msg = repassar_mensagem_acompanhamento(msg, request.user, observacao)
            msg_txt = (
                f"Recado repassado com sucesso para a "
                f"{nova_msg.get_turma_destino_display()} ({nova_msg.data_turno_destino.strftime('%d/%m')})!"
            )
            if is_ajax:
                return JsonResponse({'success': True, 'acao': 'REPASSAR', 'message': msg_txt, 'nova_id': nova_msg.id})
            messages.success(request, msg_txt)

        else:
            if is_ajax:
                return JsonResponse({'success': False, 'error': 'Ação inválida.'}, status=400)
            messages.error(request, "Ação inválida solicitada.")

    except (ValueError, PermissionError) as e:
        if is_ajax:
            return JsonResponse({'success': False, 'error': str(e)}, status=400)
        messages.error(request, f"Erro: {str(e)}")

    return redirect(next_url)


@login_required
@lider_bladder_required
def passagem_turno_lista(request):
    """
    Painel de Gestão e Histórico de Passagens de Turno para o Líder Bladder.
    Permite busca, filtros avançados por período, tipo, prioridade, status, categoria
    e inspeção da cadeia completa de repasses.
    """
    tipo_filtro = request.GET.get('tipo', '').strip()
    categoria_filtro = request.GET.get('categoria', '').strip()
    prioridade_filtro = request.GET.get('prioridade', '').strip()
    status_filtro = request.GET.get('status', '').strip()
    turma_origem_filtro = request.GET.get('turma_origem', '').strip()
    turma_destino_filtro = request.GET.get('turma_destino', '').strip()
    dt_inicio_str = request.GET.get('data_inicio', '').strip()
    dt_fim_str = request.GET.get('data_fim', '').strip()
    q = request.GET.get('q', '').strip()

    qs = MensagemPassagemTurnoBladder.objects.all().select_related(
        'autor', 'ordem_producao', 'processo', 'maquina', 'produto',
        'resolvido_por', 'repassado_por', 'mensagem_origem'
    ).prefetch_related(
        'acoes__usuario'
    )

    if tipo_filtro:
        qs = qs.filter(tipo=tipo_filtro)
    if categoria_filtro:
        qs = qs.filter(categoria=categoria_filtro)
    if prioridade_filtro:
        qs = qs.filter(prioridade=prioridade_filtro)
    if status_filtro:
        qs = qs.filter(status=status_filtro)
    if turma_origem_filtro:
        qs = qs.filter(turma_origem=turma_origem_filtro)
    if turma_destino_filtro:
        qs = qs.filter(turma_destino=turma_destino_filtro)

    if dt_inicio_str:
        try:
            dt_inicio = datetime.datetime.strptime(dt_inicio_str, '%Y-%m-%d').date()
            qs = qs.filter(data_turno_origem__gte=dt_inicio)
        except ValueError:
            pass

    if dt_fim_str:
        try:
            dt_fim = datetime.datetime.strptime(dt_fim_str, '%Y-%m-%d').date()
            qs = qs.filter(data_turno_origem__lte=dt_fim)
        except ValueError:
            pass

    if q:
        qs = qs.filter(
            Q(mensagem__icontains=q) |
            Q(autor__first_name__icontains=q) |
            Q(autor__username__icontains=q) |
            Q(ordem_producao__numero_ordem__icontains=q)
        )

    # Ordenação por prioridade e data
    qs = qs.annotate(
        peso_prioridade=Case(
            When(prioridade='URGENTE', then=Value(1)),
            When(prioridade='IMPORTANTE', then=Value(2)),
            default=Value(3),
            output_field=IntegerField()
        )
    ).order_by('peso_prioridade', '-created_at')

    # Métricas agregadas
    total_geral = qs.count()
    total_informativos = qs.filter(tipo='INFORMATIVO').count()
    total_acompanhamentos = qs.filter(tipo='ACOMPANHAMENTO').count()
    total_abertas = qs.filter(status='ABERTA').count()
    total_resolvidas = qs.filter(status='RESOLVIDA').count()
    total_repassadas = qs.filter(status='REPASSADA').count()
    total_urgentes = qs.filter(prioridade='URGENTE').count()

    context = {
        'mensagens': qs,
        'total_geral': total_geral,
        'total_informativos': total_informativos,
        'total_acompanhamentos': total_acompanhamentos,
        'total_abertas': total_abertas,
        'total_resolvidas': total_resolvidas,
        'total_repassadas': total_repassadas,
        'total_urgentes': total_urgentes,
        'tipo_filtro': tipo_filtro,
        'categoria_filtro': categoria_filtro,
        'prioridade_filtro': prioridade_filtro,
        'status_filtro': status_filtro,
        'turma_origem_filtro': turma_origem_filtro,
        'turma_destino_filtro': turma_destino_filtro,
        'data_inicio': dt_inicio_str,
        'data_fim': dt_fim_str,
        'q': q,
        'tipos_choices': MensagemPassagemTurnoBladder.TIPO_CHOICES,
        'categorias_choices': MensagemPassagemTurnoBladder.CATEGORIA_CHOICES,
        'prioridades_choices': MensagemPassagemTurnoBladder.PRIORIDADE_CHOICES,
        'status_choices': MensagemPassagemTurnoBladder.STATUS_CHOICES,
        'turma_choices': MensagemPassagemTurnoBladder.TURMA_CHOICES,
    }
    return render(request, 'bladder/passagem_turno_lista.html', context)


@login_required
@operador_ou_lider_bladder_required
def recados_criados_lista(request):
    """
    Tela de Listagem e Acompanhamento de Recados Criados para o Próximo Turno.
    Acessível por Operadores, Apoio e Liderança do Setor de Bladder.
    Permite:
    - Visualizar todos os recados deixados pela equipe (ou histórico geral);
    - Acompanhar confirmações de leitura (ciência) dos operadores do próximo turno;
    - Consultar ocorrências resolvidas ou repassadas;
    - Criar novos recados para o próximo turno diretamente pela tela.
    """
    hoje = timezone.localdate()
    turma_atual, _, _ = calcular_turma_do_dia(hoje)
    prox_data, prox_turma, _, _ = calcular_proximo_turno_operacional(hoje)

    visao = request.GET.get('visao', 'meu_turno').strip()
    tipo_filtro = request.GET.get('tipo', '').strip()
    categoria_filtro = request.GET.get('categoria', '').strip()
    prioridade_filtro = request.GET.get('prioridade', '').strip()
    status_filtro = request.GET.get('status', '').strip()
    turma_origem_filtro = request.GET.get('turma_origem', '').strip()
    dt_inicio_str = request.GET.get('data_inicio', '').strip()
    dt_fim_str = request.GET.get('data_fim', '').strip()
    q = request.GET.get('q', '').strip()

    qs = MensagemPassagemTurnoBladder.objects.all().select_related(
        'autor', 'ordem_producao', 'ordem_producao__produto', 'processo', 'maquina', 'produto',
        'resolvido_por', 'repassado_por', 'mensagem_origem'
    ).prefetch_related(
        'acoes__usuario'
    )

    if visao == 'meu_turno':
        qs = qs.filter(data_turno_origem=hoje)
        if turma_atual:
            qs = qs.filter(turma_origem=turma_atual)
    elif visao == 'meus':
        qs = qs.filter(autor=request.user)
    elif visao == 'abertas':
        qs = qs.filter(status='ABERTA', tipo='ACOMPANHAMENTO')
    elif visao == 'resolvidas':
        qs = qs.filter(status='RESOLVIDA')

    if tipo_filtro:
        qs = qs.filter(tipo=tipo_filtro)
    if categoria_filtro:
        qs = qs.filter(categoria=categoria_filtro)
    if prioridade_filtro:
        qs = qs.filter(prioridade=prioridade_filtro)
    if status_filtro:
        qs = qs.filter(status=status_filtro)
    if turma_origem_filtro:
        qs = qs.filter(turma_origem=turma_origem_filtro)

    if dt_inicio_str:
        try:
            dt_inicio = datetime.datetime.strptime(dt_inicio_str, '%Y-%m-%d').date()
            qs = qs.filter(data_turno_origem__gte=dt_inicio)
        except ValueError:
            pass

    if dt_fim_str:
        try:
            dt_fim = datetime.datetime.strptime(dt_fim_str, '%Y-%m-%d').date()
            qs = qs.filter(data_turno_origem__lte=dt_fim)
        except ValueError:
            pass

    if q:
        qs = qs.filter(
            Q(mensagem__icontains=q) |
            Q(autor__first_name__icontains=q) |
            Q(autor__last_name__icontains=q) |
            Q(autor__username__icontains=q) |
            Q(ordem_producao__numero_ordem__icontains=q)
        )

    qs = qs.annotate(
        peso_prioridade=Case(
            When(prioridade='URGENTE', then=Value(1)),
            When(prioridade='IMPORTANTE', then=Value(2)),
            default=Value(3),
            output_field=IntegerField()
        )
    ).order_by('peso_prioridade', '-created_at')

    # Contadores para os cards informativos
    recados_hoje_count = MensagemPassagemTurnoBladder.objects.filter(data_turno_origem=hoje).count()
    abertos_count = MensagemPassagemTurnoBladder.objects.filter(status='ABERTA', tipo='ACOMPANHAMENTO').count()
    resolvidos_count = MensagemPassagemTurnoBladder.objects.filter(status='RESOLVIDA').count()
    urgentes_count = MensagemPassagemTurnoBladder.objects.filter(prioridade='URGENTE').count()
    total_criados = MensagemPassagemTurnoBladder.objects.count()

    turma_dict = dict(MensagemPassagemTurnoBladder.TURMA_CHOICES)

    context = {
        'mensagens': qs,
        'hoje': hoje,
        'turma_atual': turma_atual,
        'turma_atual_display': turma_dict.get(turma_atual, turma_atual) if turma_atual else 'Sem Escala',
        'prox_data': prox_data,
        'prox_turma': prox_turma,
        'prox_turma_display': turma_dict.get(prox_turma, prox_turma) if prox_turma else '',
        'visao': visao,
        'tipo_filtro': tipo_filtro,
        'categoria_filtro': categoria_filtro,
        'prioridade_filtro': prioridade_filtro,
        'status_filtro': status_filtro,
        'turma_origem_filtro': turma_origem_filtro,
        'data_inicio': dt_inicio_str,
        'data_fim': dt_fim_str,
        'q': q,
        'recados_hoje_count': recados_hoje_count,
        'abertos_count': abertos_count,
        'resolvidos_count': resolvidos_count,
        'urgentes_count': urgentes_count,
        'total_criados': total_criados,
        'tipos_choices': MensagemPassagemTurnoBladder.TIPO_CHOICES,
        'categorias_choices': MensagemPassagemTurnoBladder.CATEGORIA_CHOICES,
        'prioridades_choices': MensagemPassagemTurnoBladder.PRIORIDADE_CHOICES,
        'status_choices': MensagemPassagemTurnoBladder.STATUS_CHOICES,
        'turma_choices': MensagemPassagemTurnoBladder.TURMA_CHOICES,
        'form_mensagem': MensagemPassagemTurnoForm(),
    }
    return render(request, 'bladder/recados_criados_lista.html', context)


